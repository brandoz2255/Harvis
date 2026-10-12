"""Turn a notebook source (file, page, recording, video) into plain text.

Every extractor either returns text or raises IngestionError with a sentence
the user can act on. That message is what the Sources panel shows next to the
red status, so it must say what to do, not what Python did.
"""
from __future__ import annotations

import asyncio
import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
from typing import List, Optional

from .models import SourceType
from .source_types import AUDIO_EXTS, HTML_EXTS, extension_of, html_to_text, looks_like_html, read_text_file

logger = logging.getLogger(__name__)

AUDIO_TIMEOUT_S = 900
_BROWSER_UA = "Mozilla/5.0 (compatible; HarvisNotebooks/1.0)"


class IngestionError(Exception):
    """A source could not be read. ``str(exc)`` is shown to the user as is."""


def _missing(path: Optional[str]) -> bool:
    return not path or not os.path.exists(path)


async def extract_pdf(path: str) -> str:
    if _missing(path):
        raise IngestionError("The uploaded file is no longer on the server. Remove this source and upload it again.")
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise IngestionError("PDF reading is not installed on this server (pypdf).") from exc

    def _read() -> str:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise IngestionError("This PDF is password protected. Remove the password and upload it again.")
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception as exc:
                logger.warning("Skipped an unreadable PDF page: %s", exc)
        return "\n\n".join(pages)

    try:
        text = await asyncio.to_thread(_read)
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(f"This PDF could not be opened ({str(exc)[:120]}). It may be damaged.") from exc
    if len(text.strip()) < 10:
        raise IngestionError(
            "This PDF has no selectable text, so it is probably a scan. Upload its pages as images "
            "(PNG or JPG) and Notebooks will read them with OCR."
        )
    return text


async def extract_doc(path: str) -> str:
    if _missing(path):
        raise IngestionError("The uploaded file is no longer on the server. Remove this source and upload it again.")
    try:
        from docx import Document
    except ImportError as exc:
        raise IngestionError("Word reading is not installed on this server (python-docx).") from exc

    def _read() -> str:
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n\n".join(parts)

    try:
        return await asyncio.to_thread(_read)
    except Exception as exc:
        raise IngestionError(
            "This Word file could not be opened. Old .doc files are not supported; save it as .docx or PDF "
            f"and upload that ({str(exc)[:80]})."
        ) from exc


async def extract_text_file(path: str, ext: str = "") -> str:
    if _missing(path):
        raise IngestionError("The uploaded file is no longer on the server. Remove this source and upload it again.")
    text = await asyncio.to_thread(read_text_file, path)
    ext = ext or extension_of(path)
    if ext in HTML_EXTS or looks_like_html(text):
        text = await asyncio.to_thread(html_to_text, text)
    return text


async def extract_image(path: str) -> str:
    """A vision-model description plus OCR text. Either alone is enough."""
    if _missing(path):
        raise IngestionError("The uploaded image is no longer on the server. Remove this source and upload it again.")
    parts: List[str] = []
    caption_note = ""

    try:
        from vison_models.llm_connector import VISION_ERROR_PREFIX, query_vision

        caption = await asyncio.to_thread(
            query_vision,
            path,
            "Describe this image in detail, including any visible text, objects, people, charts, and setting.",
        )
        cap = (caption or "").strip()
        if cap and not cap.startswith(VISION_ERROR_PREFIX):
            parts.append(f"Image description:\n{cap}")
        else:
            caption_note = "no vision model answered"
    except Exception as exc:
        caption_note = f"vision step failed: {str(exc)[:80]}"
        logger.warning("Image caption failed for %s: %s", path, exc)

    if shutil.which("tesseract"):
        def _ocr() -> str:
            run = subprocess.run(["tesseract", path, "stdout"], capture_output=True, text=True, timeout=90)
            return (run.stdout or "").strip()

        try:
            ocr = await asyncio.to_thread(_ocr)
            if len(ocr) >= 3:
                parts.append(f"Text detected in image (OCR):\n{ocr}")
        except Exception as exc:
            logger.warning("Image OCR failed for %s: %s", path, exc)

    if not parts:
        hint = f" ({caption_note})" if caption_note else ""
        raise IngestionError(
            "No text or description could be read from this image" + hint + ". Install a vision model "
            "(for example ollama pull llava) or upload a sharper image that contains readable text."
        )
    return "\n\n".join(parts)


async def extract_audio(path: str) -> str:
    """Speech to text through the same STT provider the voice features use."""
    if _missing(path):
        raise IngestionError("The uploaded recording is no longer on the server. Remove this source and upload it again.")
    import transcription

    state = transcription.status()
    if not state.get("server_transcribes"):
        raise IngestionError(
            "This server has no speech-to-text turned on, so audio cannot be transcribed. "
            "Ask the host to enable it (HARVIS_STT_PROVIDER), or upload a text transcript instead."
        )

    work = tempfile.mkdtemp(prefix="nb-audio-")
    wav = os.path.join(work, "in.wav")
    target = path
    try:
        if shutil.which("ffmpeg"):
            def _convert() -> bool:
                run = subprocess.run(
                    ["ffmpeg", "-y", "-loglevel", "error", "-i", path, "-vn", "-ac", "1", "-ar", "16000", wav],
                    capture_output=True,
                    timeout=AUDIO_TIMEOUT_S,
                )
                return run.returncode == 0 and os.path.exists(wav)

            ok = await asyncio.to_thread(_convert)
            if not ok:
                raise IngestionError(
                    "This audio file could not be decoded. Try MP3, WAV, M4A or OGG, or re-export the recording."
                )
            target = wav
        result = await asyncio.to_thread(transcription.transcribe, target)
    except IngestionError:
        raise
    except transcription.TranscriptionUnavailable as exc:
        raise IngestionError(f"Speech-to-text is not available: {str(exc)[:240]}") from exc
    except subprocess.TimeoutExpired as exc:
        raise IngestionError("This recording is too long to transcribe in one go. Split it into shorter files.") from exc
    except Exception as exc:
        raise IngestionError(f"The recording could not be transcribed ({str(exc)[:160]}).") from exc
    finally:
        shutil.rmtree(work, ignore_errors=True)

    text = ((result or {}).get("text") or "").strip()
    if len(text) < 10:
        raise IngestionError("No speech was found in this recording.")
    return text


def _decode(resp) -> str:
    try:
        return resp.text
    except Exception:
        return resp.content.decode("utf-8", errors="replace")


async def extract_url(url: str) -> str:
    """A public web page as readable text. Fetched through the SSRF-safe reach client."""
    from agent_reach.tools import ReachBlocked, _safe_get

    import httpx

    try:
        resp = await _safe_get(url, headers={"User-Agent": _BROWSER_UA, "Accept": "text/html,*/*;q=0.8"}, timeout=25.0)
    except ReachBlocked as exc:
        raise IngestionError(
            f"That address cannot be added: {str(exc)[:160]}. Only public web pages are allowed, "
            "not private or internal addresses."
        ) from exc
    except httpx.HTTPStatusError as exc:
        raise IngestionError(
            f"The site answered HTTP {exc.response.status_code}, so the page could not be read. "
            "If it needs a login, copy the text and add it as a pasted text source."
        ) from exc
    except httpx.TimeoutException as exc:
        raise IngestionError("The site took too long to answer. Try again in a moment.") from exc
    except Exception as exc:
        raise IngestionError(f"The page could not be fetched ({str(exc)[:120]}).") from exc

    ctype = (resp.headers.get("content-type") or "").lower()
    if "pdf" in ctype or url.lower().split("?")[0].endswith(".pdf"):
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(resp.content)
        try:
            return await extract_pdf(tmp.name)
        finally:
            os.unlink(tmp.name)
    if ctype and not any(k in ctype for k in ("html", "text", "xml", "json")):
        raise IngestionError(f"That link is a {ctype.split(';')[0]} file, not a web page. Download it and upload it instead.")

    body = _decode(resp)
    text = await asyncio.to_thread(html_to_text, body) if (looks_like_html(body) or "html" in ctype or "<p" in body[:5000]) else body
    if len(text.strip()) < 50:
        raise IngestionError(
            "No readable text was found on that page. It may need JavaScript or a login; "
            "copy the text and add it as a pasted text source."
        )
    return text


_YT_PATTERNS = (
    r"(?:v=|/v/|youtu\.be/)([a-zA-Z0-9_-]{11})",
    r"(?:embed/|shorts/)([a-zA-Z0-9_-]{11})",
)


def youtube_video_id(url: str) -> Optional[str]:
    for pattern in _YT_PATTERNS:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return url if re.match(r"^[a-zA-Z0-9_-]{11}$", url or "") else None


async def extract_youtube(url: str) -> str:
    video_id = youtube_video_id(url)
    if not video_id:
        raise IngestionError("That does not look like a YouTube video link.")
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
    except ImportError as exc:
        raise IngestionError("YouTube transcripts are not installed on this server.") from exc

    def _fetch() -> str:
        api = YouTubeTranscriptApi()
        if hasattr(api, "fetch"):
            items = api.fetch(video_id)
            return " ".join((getattr(i, "text", None) or i.get("text", "")).strip() for i in items)
        items = YouTubeTranscriptApi.get_transcript(video_id)
        return " ".join(i.get("text", "").strip() for i in items)

    try:
        text = await asyncio.to_thread(_fetch)
    except Exception as exc:
        raise IngestionError(
            "No transcript could be fetched for this video (captions may be off, or YouTube blocked the "
            f"request: {type(exc).__name__}). Paste the transcript as a text source instead."
        ) from exc
    if len(text.strip()) < 10:
        raise IngestionError("This video has an empty transcript.")
    return text


async def extract_for(source) -> str:
    """Dispatch on the source type. ``source.storage_path`` is a file path or the URL."""
    kind, path = source.type, source.storage_path
    ext = ""
    meta = getattr(source, "metadata", None) or {}
    if isinstance(meta, dict):
        ext = (meta.get("extension") or "").lower()
    if kind == SourceType.PDF:
        return await extract_pdf(path)
    if kind == SourceType.DOC:
        return await extract_doc(path)
    if kind == SourceType.URL:
        return await extract_url(path)
    if kind == SourceType.YOUTUBE:
        return await extract_youtube(path)
    if kind == SourceType.IMAGE:
        return await extract_image(path)
    if kind == SourceType.AUDIO:
        return await extract_audio(path)
    if kind in (SourceType.TEXT, SourceType.MARKDOWN, SourceType.TRANSCRIPT):
        return await extract_text_file(path, ext)
    raise IngestionError(f"Sources of type {getattr(kind, 'value', kind)} are not supported yet.")


__all__ = ["IngestionError", "extract_for", "youtube_video_id", "AUDIO_EXTS"]
