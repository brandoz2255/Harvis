"""Spreadsheets that show their values in any viewer, and formula errors that are caught before delivery.

The sandbox is never called: a stand-in plays LibreOffice by writing each formula's value into the
workbook the way a recalculation stores it. The real /recalc is exercised against the real image
(docker exec harvis-sandbox ...), as with /render."""

import base64
import io
import os
import re
import shutil
import sys
import urllib.error
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from plugins.hermes_ui import docs_check, docs_export, docs_render_client  # noqa: E402
from plugins.hermes_ui.docs_check_sheet import _Grid, check_xlsx  # noqa: E402
from plugins.hermes_ui.docs_export import export_doc  # noqa: E402
from plugins.hermes_ui.docs_formula import (BLANK, RECALC_ONLY, SAFE_FUNCTIONS, SUPPORTED, FormulaError,  # noqa: E402
                                            Unknowable, evaluate, evaluate_any)
from sandbox_runner import jobs, render  # noqa: E402

PROBE = """---
kind: sheet
title: Probe
outline:
  - Orders
---
## Orders
| Item | Qty | Price | Line | Size |
| --- | --- | --- | --- | --- |
| Pen | 2 | $3.50 | =B2*C2 | =IF(B2>1,"bulk","single") |
| Ink | 4 | $10.00 | =B3*C3 | =IF(B3>1,"bulk","single") |
| **Total** | =SUM(B2:B3) | | =SUM(D2:D3) | |
"""


# ---- a stand-in for the sandbox's recalculation -----------------------------------------------------

def _stored(xlsx: bytes, errors=None) -> bytes:
    """``xlsx`` with a value stored beside every formula, as a recalculation does. Values come from this
    repo's evaluator; ``errors`` maps a cell to an error to store there instead."""
    errors = errors or {}
    wb = load_workbook(io.BytesIO(xlsx))
    values: dict[tuple[str, str], object] = {}
    for index, ws in enumerate(wb.worksheets, 1):
        grid = _Grid(ws)
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                if cell.coordinate in errors:
                    values[(f"sheet{index}", cell.coordinate)] = errors[cell.coordinate]
                    continue
                try:
                    values[(f"sheet{index}", cell.coordinate)] = grid.cell(cell.row - 1, cell.column - 1)
                except FormulaError:
                    pass
    src, out = zipfile.ZipFile(io.BytesIO(xlsx)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            m = re.match(r"xl/worksheets/(sheet\d+)\.xml$", item.filename)
            if m:
                xml = data.decode()
                for (sheet, coord), value in values.items():
                    if sheet != m.group(1):
                        continue
                    if isinstance(value, bool):
                        kind, text = ' t="b"', str(int(value))
                    elif isinstance(value, str):
                        kind, text = (' t="e"' if value.startswith("#") else ' t="str"'), value
                    else:
                        kind, text = "", repr(float(value))
                    xml = re.sub(rf'<c r="{coord}"((?: s="\d+")?)><f>(.*?)</f><v></v></c>',
                                 lambda g: f'<c r="{coord}"{g.group(1)}{kind}><f>{g.group(2)}</f><v>{text}</v></c>', xml)
                data = xml.encode()
            dst.writestr(item, data)
    return out.getvalue()


@pytest.fixture(autouse=True)
def no_sandbox(monkeypatch):
    """Every test starts with the sandbox out of reach and the client's memory empty."""
    monkeypatch.setattr(docs_render_client, "_cache", docs_render_client.OrderedDict())
    monkeypatch.setattr(docs_render_client, "_bad", docs_render_client.OrderedDict())
    monkeypatch.setattr(docs_render_client, "_skip_until", 0.0)

    def unreachable(req, timeout):
        raise urllib.error.URLError(ConnectionRefusedError("no sandbox in this test"))

    monkeypatch.setattr(docs_render_client.urllib.request, "urlopen", unreachable)
    monkeypatch.delenv("DOCS_XLSX_RECALC", raising=False)


@pytest.fixture
def recalculating(monkeypatch):
    """The export goes through a sandbox that stores every formula's value."""
    calls = []

    def fake(data):
        calls.append(len(data))
        return _stored(data)

    monkeypatch.setattr(docs_render_client, "recalc_xlsx", fake)
    return calls


# ---- values in any viewer ---------------------------------------------------------------------------

def test_a_recalculated_export_shows_numbers_with_data_only(recalculating):
    data = export_doc(PROBE, "xlsx").data
    assert recalculating, "the export did not ask for a recalculation"
    ws = load_workbook(io.BytesIO(data), data_only=True)["Orders"]
    assert ws["D2"].value == 7 and ws["D3"].value == 40 and ws["D4"].value == 47 and ws["B4"].value == 6
    assert ws["E2"].value == "bulk"
    live = load_workbook(io.BytesIO(data))["Orders"]
    assert live["D2"].value == "=B2*C2" and live["D2"].data_type == "f"


def test_a_workbook_with_no_formula_never_asks_the_sandbox(recalculating):
    export_doc("---\nkind: sheet\ntitle: T\n---\n## S\n| A | B |\n| --- | --- |\n| 1 | 2 |\n", "xlsx")
    assert recalculating == []


def test_the_export_succeeds_when_the_sandbox_is_down():
    data = export_doc(PROBE, "xlsx").data                                 # urlopen refuses: connection error
    ws = load_workbook(io.BytesIO(data), data_only=True)["Orders"]
    assert data[:2] == b"PK" and ws["D2"].value is None                   # no cached values, but a good file
    assert load_workbook(io.BytesIO(data))["Orders"]["D2"].value == "=B2*C2"


def test_an_older_sandbox_without_recalc_is_remembered_not_asked_again(monkeypatch):
    seen = []

    def missing(req, timeout):
        seen.append(req.full_url)
        raise urllib.error.HTTPError(req.full_url, 404, "not found", {}, None)

    monkeypatch.setattr(docs_render_client.urllib.request, "urlopen", missing)
    first = export_doc(PROBE, "xlsx").data
    export_doc(PROBE.replace("Probe", "Probe two"), "xlsx")
    assert first[:2] == b"PK" and len(seen) == 1 and seen[0].endswith("/recalc")


def test_a_slow_sandbox_falls_back_and_the_call_is_bounded(monkeypatch):
    waited = []

    def slow(req, timeout):
        waited.append(timeout)
        raise urllib.error.URLError(TimeoutError("timed out"))

    monkeypatch.setattr(docs_render_client.urllib.request, "urlopen", slow)
    assert docs_render_client.recalc_xlsx(b"PK-workbook") is None
    assert waited == [docs_render_client.RECALC_TIMEOUT_S] and docs_render_client.RECALC_TIMEOUT_S < 60


def test_the_same_workbook_is_recalculated_once(monkeypatch):
    seen = []

    def sandbox(req, timeout):
        seen.append(1)

        class Reply:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def read(self):
                import json
                return json.dumps({"ok": True, "b64": base64.b64encode(b"PK-calculated").decode()}).encode()

        return Reply()

    monkeypatch.setattr(docs_render_client.urllib.request, "urlopen", sandbox)
    assert docs_render_client.recalc_xlsx(b"PK-one") == b"PK-calculated"
    assert docs_render_client.recalc_xlsx(b"PK-one") == b"PK-calculated"
    assert len(seen) == 1
    assert docs_render_client.recalc_xlsx(b"PK-two") == b"PK-calculated" and len(seen) == 2


def test_a_refusal_from_the_sandbox_is_not_retried_for_that_workbook(monkeypatch):
    seen = []

    class Reply:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self):
            return b'{"ok": false, "error": "could not recalculate the file"}'

    monkeypatch.setattr(docs_render_client.urllib.request, "urlopen", lambda req, timeout: seen.append(1) or Reply())
    assert docs_render_client.recalc_xlsx(b"PK-broken") is None
    assert docs_render_client.recalc_xlsx(b"PK-broken") is None and len(seen) == 1


def test_a_recalculation_that_renames_a_sheet_is_not_used(monkeypatch):
    def other_sheets(data):
        wb = load_workbook(io.BytesIO(data))
        wb.active.title = "Renamed"
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    monkeypatch.setattr(docs_render_client, "recalc_xlsx", other_sheets)
    assert load_workbook(io.BytesIO(export_doc(PROBE, "xlsx").data)).sheetnames == ["Orders"]


def test_the_opt_out_keeps_the_sandbox_out_of_the_export(monkeypatch, recalculating):
    monkeypatch.setenv("DOCS_XLSX_RECALC", "0")
    export_doc(PROBE, "xlsx")
    assert recalculating == []


# ---- every formula error is caught ------------------------------------------------------------------

DIVIDE = """---
kind: sheet
title: Ratios
outline:
  - Ratios
---
## Ratios
| Item | Won | Lost | Rate |
| --- | --- | --- | --- |
| A | 4 | 2 | =B2/C2 |
| B | 3 | 0 | =B3/C3 |
"""


def test_a_division_by_zero_is_reported_with_its_cell_and_formula(recalculating):
    issues, _files, _doc = docs_check.check_source(DIVIDE)
    hit = [i for i in issues if i["severity"] == "error" and "D3" in i["message"]]
    assert hit and "=B3/C3" in hit[0]["message"] and "#DIV/0!" in hit[0]["message"]
    assert hit[0]["where"] == "sheet Ratios"
    assert not [i for i in issues if "D2" in i["message"]]                # the good row is not blamed


def test_an_error_a_recalculation_stored_is_reported_for_a_formula_the_evaluator_cannot_run():
    plain = export_doc("---\nkind: sheet\ntitle: T\n---\n## S\n| K | V |\n| --- | --- |\n| a | 1 |\n| b | 2 |\n"
                       "| Pick | =INDEX(B2:B3,MATCH(\"z\",A2:A3,0)) |\n", "xlsx").data
    assert not [i for i in check_xlsx(plain) if "B4" in i["message"]]     # nothing stored, nothing to say
    for stored in ("#N/A", "#REF!", "Err:502"):
        issues = check_xlsx(_stored(plain, {"B4": stored}))
        hit = [i for i in issues if "B4" in i["message"]]
        assert hit and hit[0]["severity"] == "error" and "INDEX(B2:B3" in hit[0]["message"] and stored in hit[0]["message"]


def test_a_clean_recalculated_sheet_has_no_issues(recalculating):
    issues, files, _doc = docs_check.check_source(PROBE)
    assert issues == [] and set(files) == {"xlsx", "csv"}


def test_a_stored_value_stands_in_for_a_function_the_evaluator_lacks():
    plain = export_doc("---\nkind: sheet\ntitle: T\n---\n## S\n| K | V |\n| --- | --- |\n| a | 3 |\n| b | 9 |\n"
                       "| Mid | =MEDIAN(B2:B3) |\n| Twice | =B4*2 |\n", "xlsx").data
    bare = _Grid(load_workbook(io.BytesIO(plain))["S"])
    assert bare.safe(3, 1) is None and bare.safe(4, 1) is None            # unknown without a stored value
    ws = load_workbook(io.BytesIO(plain))["S"]
    stored = _Grid(ws, load_workbook(io.BytesIO(_with(plain, "B4", 6.0)), data_only=True)["S"])
    assert stored.safe(3, 1) == 6.0 and stored.safe(4, 1) == 12.0


def _with(xlsx: bytes, coord: str, value: float) -> bytes:
    """The workbook with one stored value, for a function the evaluator does not know."""
    src, out = zipfile.ZipFile(io.BytesIO(xlsx)), io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rf'(<c r="{coord}"[^>]*><f>.*?</f>)<v></v>', rf"\g<1><v>{value}</v>", data.decode()).encode()
            dst.writestr(item, data)
    return out.getvalue()


# ---- number formats ---------------------------------------------------------------------------------

def test_a_currency_columns_total_gets_the_currency_format():
    ws = load_workbook(io.BytesIO(export_doc(PROBE, "xlsx").data))["Orders"]
    assert ws["C2"].number_format == "$#,##0.00"
    assert ws["D4"].number_format == "General"                            # column D holds only formulas: nothing to copy
    src = PROBE.replace("| =B2*C2 |", "| $7.00 |")
    ws = load_workbook(io.BytesIO(export_doc(src, "xlsx").data))["Orders"]
    assert ws["D4"].number_format == "$#,##0.00"


def test_the_nearest_typed_number_decides_and_counts_stay_plain():
    src = ("---\nkind: sheet\ntitle: T\n---\n## S\n| Item | Cost | Share |\n| --- | --- | --- |\n"
           "| a | $100 | 10% |\n| b | 200 | 20% |\n| Sum | =SUM(B2:B3) | =AVERAGE(C2:C3) |\n"
           "| How many | =COUNT(B2:B3) | =C3 |\n")
    ws = load_workbook(io.BytesIO(export_doc(src, "xlsx").data))["S"]
    assert ws["B4"].number_format == "General"                            # the nearest number above, 200, is unformatted
    assert ws["C4"].number_format == "0%" and ws["C5"].number_format == "0%"
    assert ws["B5"].number_format == "General"                            # a count is a plain number


def test_a_formula_with_no_number_above_takes_the_nearest_below():
    src = ("---\nkind: sheet\ntitle: T\n---\n## S\n| Item | Cost |\n| --- | --- |\n| x | =SUM(B3:B4) |\n"
           "| y | $5 |\n| z | $6 |\n")
    assert load_workbook(io.BytesIO(export_doc(src, "xlsx").data))["S"]["B2"].number_format == "$#,##0"


# ---- the evaluator ----------------------------------------------------------------------------------

GRID = {(0, 0): "Item", (0, 1): "Cost", (1, 0): "Food", (1, 1): 10.0, (2, 0): "Rent", (2, 1): 20.0,
        (3, 0): "food", (3, 1): 5.0}


def grid(r, c):
    return GRID.get((r, c))


@pytest.mark.parametrize("formula,expected", [
    ('=IF(B2>5,"hi","lo")', "hi"), ('=IF(B2<5,"hi","lo")', "lo"), ("=IF(B2<5,1)", False),
    ("=IF(B2=0,0,20/B2)", 2.0), ('=IF(A5="","empty",A5)', "empty"),
    ("=IFERROR(B2/0,-1)", -1.0), ('=IFERROR(A2+1,"text")', "text"), ("=IFERROR(B2,0)", 10.0),
    ('=SUMIF(A2:A4,"food",B2:B4)', 15.0), ('=SUMIF(B2:B4,">=10")', 30.0), ('=SUMIF(A2:A4,"R*",B2:B4)', 20.0),
    ('=COUNTIF(B2:B4,">7")', 2.0), ('=COUNTIF(A2:A4,"<>Rent")', 2.0), ("=COUNTIF(B2:B4,5)", 1.0),
    ('=VLOOKUP("rent",A2:B4,2,FALSE)', 20.0), ("=VLOOKUP(15,B2:B4,1)", 5.0), ("=VLOOKUP(10,B2:B4,1,FALSE)", 10.0),
    ("=AND(B2>1,B3>1)", True), ("=OR(B2>50,B3>50)", False), ("=NOT(B2>50)", True),
    ("=SUM(A1:B4)", 35.0), ("=COUNT(A1:B4)", 3.0), ("=10%*50+2^3-B2", 3.0), ("=1<=2", True), ('="a"="A"', True),
])
def test_the_evaluator_reads_if_iferror_sumif_countif_and_vlookup(formula, expected):
    assert evaluate_any(formula, grid) == expected


def test_only_the_branch_if_takes_is_calculated():
    assert evaluate("=IF(B2>5,B2,B2/0)", grid) == 10.0
    assert evaluate("=IF(B2<5,B2/0,B2)", grid) == 10.0
    with pytest.raises(FormulaError):
        evaluate("=IF(B2>5,B2/0,B2)", grid)


@pytest.mark.parametrize("bad", ['=VLOOKUP("zzz",A2:B4,2,FALSE)', '=VLOOKUP("Food",A2:B4,3,FALSE)', "=IF(B2)", '=A2+1',
                                 "=SUMIF(B2,1)", "=IF(B2>5,1,", "=HYPERLINK(1)"])
def test_the_evaluator_refuses_what_a_spreadsheet_would_show_as_an_error(bad):
    with pytest.raises(FormulaError):
        evaluate_any(bad, grid)


def test_a_text_result_is_not_a_number():
    with pytest.raises(FormulaError):
        evaluate('=IF(B2>5,"hi","lo")', grid)
    assert BLANK == "" and BLANK == 0 and BLANK != "x"


def test_every_safe_function_is_either_evaluated_here_or_left_to_the_recalculation():
    assert SUPPORTED.isdisjoint(RECALC_ONLY) and SUPPORTED | RECALC_ONLY == SAFE_FUNCTIONS
    assert {"IF", "IFERROR", "SUMIF", "COUNTIF", "VLOOKUP"} <= SUPPORTED


def test_the_csv_computes_if_and_lookups_instead_of_showing_an_error():
    src = ("---\nkind: sheet\ntitle: T\n---\n## S\n| Item | Cost | Flag | Sum |\n| --- | --- | --- | --- |\n"
           "| Food | 10 | =IF(B2>5,\"big\",\"small\") | =SUMIF(A2:A3,\"Food\",B2:B3) |\n"
           "| Rent | 20 | =IF(B3>5,\"big\",\"small\") | =VLOOKUP(\"Rent\",A2:B3,2,FALSE) |\n")
    rows = export_doc(src, "csv").data.decode("utf-8-sig").splitlines()
    assert rows[1] == "Food,10,big,10" and rows[2] == "Rent,20,big,20"


_HEAD = "---\nkind: sheet\ntitle: T\n---\n## S\n"


def _csv_rows(body: str, head: str) -> list[str]:
    return export_doc(_HEAD + head + "| --- | --- | --- | --- |\n" + body, "csv").data.decode("utf-8-sig").splitlines()


def test_iferror_does_not_hide_a_value_that_cannot_be_known():
    # MEDIAN is a function only a real spreadsheet program calculates, so C2 is #VALUE! here and D2 is unknowable.
    rows = _csv_rows("| A | 10 | =MEDIAN(B2:B4) | =IFERROR(C2*2,0) |\n| B | 20 | | |\n| C | 30 | | |\n",
                     "| K | V | Mid | Twice |\n")
    assert rows[1].split(",")[2:] == ["#VALUE!", "#VALUE!"]               # never the fallback 0
    with pytest.raises(Unknowable):
        evaluate_any("=IFERROR(X1+1,0)", lambda r, c: (_ for _ in ()).throw(Unknowable("unknown")))
    with pytest.raises(Unknowable):
        evaluate_any("=IFERROR(MEDIAN(A1:A2),0)", grid)                   # a function outside the evaluator
    assert evaluate_any("=IFERROR(B2/0,7)", grid) == 7.0                  # a real error value is still caught
    assert evaluate_any('=IFERROR(A2+1,"t")', grid) == "t"


def _sheet(body: str) -> bytes:
    return export_doc(_HEAD + "| K | V | W |\n| --- | --- | --- |\n" + body, "xlsx").data


def test_iferror_does_not_hide_a_circular_reference():
    plain = [i["message"] for i in check_xlsx(_sheet("| A | =B2+1 | |\n| Total | =SUM(B2:B2) | |\n")) if "B2" in i["message"]]
    masked = [i["message"] for i in check_xlsx(_sheet("| A | =IFERROR(B2+1,0) | |\n| Total | =SUM(B2:B2) | |\n"))
              if "B2" in i["message"]]
    assert plain and "circular" in plain[0]
    assert masked and "circular" in masked[0]


def test_iferror_over_unknown_cells_is_treated_like_the_plain_total():
    src = ("---\nkind: sheet\ntitle: Budget of $4\noutline:\n  - Costs\n---\n## Costs\n| K | V | W |\n| --- | --- | --- |\n"
           "| a | =ROUNDUP(3.2,0) | |\n| b | =ROUNDUP(4.2,0) | |\n| Total | =%s | |\n")
    masked = check_xlsx(export_doc(src % "IFERROR(SUM(B2:B3),0)", "xlsx").data)
    plain = check_xlsx(export_doc(src % "SUM(B2:B3)", "xlsx").data)
    assert masked == plain                                                # unknown either way, never a silent 0
    grid_ = _Grid(load_workbook(io.BytesIO(export_doc(src % "IFERROR(SUM(B2:B3),0)", "xlsx").data))["Costs"])
    assert grid_.safe(3, 1) is None                                       # not 0


@pytest.mark.parametrize("formula,expected", [
    ('=COUNTIF(A2:A5,"<>")', 2.0), ('=COUNTIF(A2:A5,"")', 2.0), ('=COUNTIF(A2:A5,"=")', 2.0),
    ('=SUMIF(A2:A5,"<>",B2:B5)', 15.0), ('=COUNTIF(A2:A5,"<>x")', 3.0),
])
def test_countif_with_a_bare_not_equal_means_not_empty(formula, expected):
    cells = {(1, 0): None, (2, 0): "x", (3, 0): 3.0, (4, 0): "", (1, 1): 1.0, (2, 1): 5.0, (3, 1): 10.0, (4, 1): 100.0}
    assert evaluate_any(formula, lambda r, c: cells.get((r, c))) == expected


def test_the_csv_counts_only_filled_notes():
    rows = _csv_rows("| a | x | | |\n| b | | | |\n| N | =COUNTIF(B2:B3,\"<>\") | | |\n", "| K | Note | | |\n")
    assert rows[3].split(",")[1] == "1"


def test_aggregates_skip_text_held_in_a_single_cell_reference():
    cells = {(0, 0): 1.0, (1, 0): "x", (2, 0): 2.0, (3, 0): True}
    read = lambda r, c: cells.get((r, c))                                 # noqa: E731
    assert evaluate_any("=SUM(A1,A2,A3)", read) == 3.0
    assert evaluate_any("=SUM(A1,A2,A3,A4)", read) == 3.0                 # a bool in a cell is skipped too
    assert evaluate_any("=AVERAGE(A1,A2,A3)", read) == 1.5
    assert evaluate_any("=MAX(A2,A3)", read) == 2.0
    assert evaluate_any("=COUNT(A1,A2,A3)", read) == 2.0
    assert evaluate_any("=PRODUCT(A1,A2,A3)", read) == 2.0
    assert evaluate_any("=SUM(A5,A1)", read) == 1.0                       # an empty cell is skipped
    assert evaluate_any("=SUM(A1,3)", read) == 4.0                        # literals still count
    assert evaluate_any('=SUM(A1,"4")', read) == 5.0                      # literal text that is a number still counts
    with pytest.raises(FormulaError):
        evaluate_any('=SUM(A1,"x")', read)                                # literal text that is not a number is an error
    with pytest.raises(FormulaError):
        evaluate_any("=A2+1", read)                                       # arithmetic on text still fails
    with pytest.raises(FormulaError):
        evaluate_any("=ABS(A2)", read)


def test_the_csv_ignores_a_text_cell_in_a_single_reference_sum():
    rows = _csv_rows("| a | 10 | n/a | |\n| T | =SUM(B2,C2) | | |\n", "| K | V | W | |\n")
    assert rows[2].split(",")[1] == "10"


# ---- the sandbox side -------------------------------------------------------------------------------

def _xlsx_b64() -> str:
    return base64.b64encode(export_doc(PROBE, "xlsx").data).decode()


def test_recalc_converts_to_xlsx_in_a_separate_output_directory(monkeypatch):
    seen = {}

    def fake_supervise(cmd, cwd, env, timeout, tmp, so, se):
        seen["cmd"] = cmd
        outdir = cmd[cmd.index("--outdir") + 1]
        with open(os.path.join(outdir, "input.xlsx"), "wb") as fh:
            fh.write(b"PK-with-values")
        return 0, None

    monkeypatch.setattr(render, "supervise", fake_supervise)
    out = render.recalc(_xlsx_b64())
    cmd = seen["cmd"]
    assert out == {"ok": True, "b64": base64.b64encode(b"PK-with-values").decode()}
    assert cmd[cmd.index("--convert-to") + 1] == "xlsx"
    assert cmd[cmd.index("--outdir") + 1] != os.path.dirname(cmd[-1])    # the source is never overwritten
    assert render.CHILD in cmd and "soffice" in cmd


def test_recalc_refuses_what_is_not_a_workbook_and_reports_a_failed_conversion(monkeypatch):
    with pytest.raises(jobs.JobError, match="base64"):
        render.recalc("%%%")
    with pytest.raises(jobs.JobError, match="not a valid .xlsx"):
        render.recalc(base64.b64encode(b"plain text").decode())
    monkeypatch.setattr(render, "supervise", lambda *a: (1, None))
    with pytest.raises(jobs.JobError, match="could not recalculate"):
        render.recalc(_xlsx_b64())
    monkeypatch.setattr(render, "supervise", lambda *a: (render.HARDEN_FAILED, None))
    with pytest.raises(jobs.JobError, match="hardening"):
        render.recalc(_xlsx_b64())


def test_to_pdf_still_converts_beside_its_input(monkeypatch):
    seen = {}

    def fake_supervise(cmd, cwd, env, timeout, tmp, so, se):
        seen["cmd"] = cmd
        with open(os.path.join(cmd[cmd.index("--outdir") + 1], "input.pdf"), "wb") as fh:
            fh.write(b"%PDF-converted")
        return 0, None

    monkeypatch.setattr(render, "supervise", fake_supervise)
    tmp = render.tempfile.mkdtemp()
    try:
        assert render.to_pdf(b"PK", "xlsx", tmp) == b"%PDF-converted"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    assert seen["cmd"][seen["cmd"].index("--convert-to") + 1] == "pdf"


@pytest.mark.skipif(not shutil.which("soffice"), reason="needs LibreOffice, which only the sandbox image has")
def test_real_libreoffice_stores_the_values():
    out = render.recalc(_xlsx_b64())
    ws = load_workbook(io.BytesIO(base64.b64decode(out["b64"])), data_only=True)["Orders"]
    assert ws["D2"].value == 7 and ws["D4"].value == 47 and ws["E2"].value == "bulk"


def test_exports_still_name_the_recalculating_builder():
    assert docs_export.to_xlsx(PROBE)[:2] == b"PK"
