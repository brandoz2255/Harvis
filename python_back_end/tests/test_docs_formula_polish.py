"""Evaluator results that must match Excel and LibreOffice: ROUND halves, empty AND/OR, and the lookup and
average errors the sheet check reports without the sandbox."""

import inspect
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook  # noqa: E402

from plugins.hermes_ui import docs_formula  # noqa: E402
from plugins.hermes_ui.docs_check_sheet import check_xlsx  # noqa: E402
from plugins.hermes_ui.docs_formula import FormulaError, evaluate, evaluate_any  # noqa: E402


def _none(r, c):
    return None


@pytest.mark.parametrize("formula, expected", [
    ("=ROUND(2.5,0)", 3.0),
    ("=ROUND(3.5,0)", 4.0),
    ("=ROUND(-2.5,0)", -3.0),
    ("=ROUND(0.125,2)", 0.13),
    ("=ROUND(1.005,2)", 1.01),
    ("=ROUND(2.675,2)", 2.68),
    ("=ROUND(1250,-2)", 1300.0),
    ("=ROUND(-1250,-2)", -1300.0),
    ("=ROUND(1234.5678,2)", 1234.57),
    ("=ROUND(7,2)", 7.0),
    ("=ROUND(-0.4,0)", 0.0),
])
def test_round_goes_half_away_from_zero(formula, expected):
    assert evaluate_any(formula, _none) == expected


def test_round_leaves_huge_and_deep_values_alone():
    assert evaluate("=ROUND(1E+300,2)", _none) == 1e300
    assert evaluate("=ROUND(0.5,400)", _none) == 0.5
    assert evaluate("=ROUND(123456,-400)", _none) == 0.0


def test_round_of_a_negative_zero_is_plain_zero():
    assert str(evaluate("=ROUND(-0.4,0)", _none)) == "0.0"


@pytest.mark.parametrize("formula", ["=AND()", "=OR()", "=AND(A1:A2)", "=OR(A1:A2)"])
def test_and_or_with_nothing_to_test_is_an_error(formula):
    with pytest.raises(FormulaError):
        evaluate_any(formula, _none)


def test_and_or_still_answer_when_there_are_values():
    assert evaluate_any("=AND(1>0,2>1)", _none) is True
    assert evaluate_any("=AND(1>0,2<1)", _none) is False
    assert evaluate_any("=OR(1<0,2>1)", _none) is True
    assert evaluate_any("=OR(1<0,2<1)", _none) is False
    assert evaluate_any("=IFERROR(AND(),7)", _none) == 7.0


def test_average_of_nothing_is_a_div0_error_the_check_can_name():
    with pytest.raises(FormulaError, match="div0"):
        evaluate_any("=AVERAGE(A1:A2)", _none)
    assert evaluate_any("=IFERROR(AVERAGE(A1:A2),9)", _none) == 9.0


def test_evaluate_functions_have_no_dead_depth_parameter():
    assert "depth" not in inspect.signature(evaluate_any).parameters
    assert "depth" not in inspect.signature(evaluate).parameters
    assert docs_formula.evaluate_any("=1+1", _none) == 2.0


def _workbook(formulas):
    """A table in A1:B3 (key, value), two empty rows under it, and one formula per row in column C."""
    wb = Workbook()
    ws = wb.active
    ws.append(["key", "value", "result"])
    ws.append(["a", 1])
    ws.append(["b", 2])
    ws.append(["x"])
    ws.append(["y"])
    for formula in formulas:
        ws.append([None, None, formula])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _messages(formulas):
    return [i["message"] for i in check_xlsx(_workbook(formulas)) if i["severity"] == "error"]


def test_check_reports_a_lookup_with_no_match():
    messages = _messages(['=VLOOKUP("zzz",A2:B3,2,FALSE)'])
    assert any("C6" in m and "#N/A" in m for m in messages), messages


def test_check_reports_a_lookup_past_its_last_column():
    messages = _messages(['=VLOOKUP("a",A2:B3,3,FALSE)'])
    assert any("C6" in m and "#REF!" in m for m in messages), messages


def test_check_reports_an_average_of_an_empty_range():
    messages = _messages(["=AVERAGE(B4:B5)"])
    assert any("C6" in m and "#DIV/0!" in m for m in messages), messages


def test_check_stays_quiet_when_the_errors_are_caught_or_absent():
    formulas = ['=IFERROR(VLOOKUP("zzz",A2:B3,2,FALSE),0)', '=VLOOKUP("b",A2:B3,2,FALSE)', "=AVERAGE(B2:B3)",
                '=IFERROR(VLOOKUP("a",A2:B3,3,FALSE),0)']
    assert _messages(formulas) == []


# ---- computed values round at 15 significant digits, empty MIN/MAX give 0, deep chains survive ---------

@pytest.mark.parametrize("formula, expected", [
    ("=ROUND(0.7*1.15,2)", 0.81),             # 0.8049999999999999 as a double, 0.805 as a spreadsheet holds it
    ("=ROUND(1.15*3,1)", 3.5),
    ("=ROUND(1.5*0.15,2)", 0.23),
    ("=ROUND(2.3*1.15,2)", 2.65),
])
def test_round_of_a_computed_value_uses_its_15_digit_form(formula, expected):
    assert evaluate(formula, _none) == expected


@pytest.mark.parametrize("formula", ["=MAX(A1:A3)", "=MIN(A1:A3)", "=MAX(A1)", "=MIN(A1)"])
def test_min_and_max_of_nothing_numeric_are_zero(formula):
    assert evaluate(formula, _none) == 0.0
    assert evaluate_any('=IFERROR(' + formula[1:] + ',"none")', _none) == 0.0


def test_min_and_max_of_text_only_are_zero():
    cells = {(0, 0): "foo", (1, 0): "bar"}
    assert evaluate("=MAX(A1:A2)", lambda r, c: cells.get((r, c))) == 0.0
    assert evaluate("=MIN(A1:A2,5)", lambda r, c: cells.get((r, c))) == 5.0


def _running_total(rows: int):
    from plugins.hermes_ui.docs_parse import Sheet
    body = [["Final", "Amount", "Running"], [f"=C{rows}", "1", "=B2"]]
    body += [["x", "1", f"=C{r - 1}+B{r}"] for r in range(3, rows + 1)]
    return Sheet(name="s", rows=body)


def test_a_summary_reading_the_bottom_of_a_long_running_total_does_not_poison_the_column():
    from plugins.hermes_ui.docs_xlsx import build_csv
    lines = [ln for ln in build_csv(_running_total(251)).decode("utf-8-sig").split("\r\n") if ln]
    assert len(lines) == 251
    assert lines[-1] == "x,1,250"
    assert lines[2] == "x,1,2" and lines[100] == "x,1,100"
    assert sum("#VALUE!" in ln for ln in lines) <= 1            # at most the summary cell that overflowed the stack


def test_a_recursion_failure_is_not_remembered_by_the_check_grid():
    from plugins.hermes_ui.docs_check_sheet import _Grid
    rows = _running_total(251).rows
    wb = Workbook()
    ws = wb.active
    for r, row in enumerate(rows, 1):
        for c, text in enumerate(row, 1):
            ws.cell(row=r, column=c).value = text if text.startswith("=") else (float(text) if r > 1 and c == 2 else text)
    grid = _Grid(ws)
    with pytest.raises(FormulaError):
        grid.cell(1, 0)                                          # the summary itself overflows the stack
    for r in range(1, 251):                                      # but the column below it was not poisoned
        assert grid.cell(r, 2) == float(r)                       # (read top-down, as a person reads it)
