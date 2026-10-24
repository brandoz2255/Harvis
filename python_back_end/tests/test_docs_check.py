"""Code checks on exported documents, and the vision pass with Ollama mocked."""

import asyncio
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from plugins.hermes_ui import docs_check, docs_vision  # noqa: E402
from plugins.hermes_ui.docs_check_rich import dropped_lines, is_empty, scan_rich  # noqa: E402
from plugins.hermes_ui.docs_check_sheet import limits_in  # noqa: E402
from plugins.hermes_ui.docs_parse import parse_harvis_doc  # noqa: E402

SHEET = """---
kind: sheet
title: Household Costs
outline:
  - Items
---
## Items
| Item | Qty | Cost |
| --- | --- | --- |
| Switch | 2 | 300 |
| Cable | 10 | 50 |
| Rack | 1 | 1800 |
| Total | | =SUM(C2:C4) |
"""

DECK = """---
kind: deck
title: Status
outline:
  - Intro
  - Plan
---
# Status
A short deck
---
## Plan
- Ship it
- Test it
"""

DOC = """---
kind: doc
title: Memo
outline:
  - Summary
---
## Summary
Plain words about the week.
"""


def errors(issues):
    return [i for i in issues if i["severity"] == "error"]


def messages(issues):
    return " | ".join(f"{i['where']}: {i['message']}" for i in issues)


# ---- sheets -----------------------------------------------------------------------------------

def test_a_clean_sheet_has_no_issues():
    issues, files, _doc = docs_check.check_source(SHEET)
    assert issues == [] and set(files) == {"xlsx", "csv"}


def test_a_formula_over_the_wrong_column_is_an_error():
    issues, _f, _d = docs_check.check_source(SHEET.replace("=SUM(C2:C4)", "=SUM(A2:A4)"))
    hit = [i for i in errors(issues) if "A2:A4" in i["message"] and "no numbers" in i["message"]]
    assert hit and hit[0]["where"] == "sheet Items"


def test_a_formula_outside_the_table_and_one_that_divides_by_zero():
    issues, _f, _d = docs_check.check_source(SHEET.replace("=SUM(C2:C4)", "=SUM(F2:F4)"))
    assert any("outside the table" in i["message"] for i in errors(issues)), messages(issues)
    issues, _f, _d = docs_check.check_source(SHEET.replace("=SUM(C2:C4)", "=C2/0"))
    assert any("#DIV/0!" in i["message"] for i in errors(issues)), messages(issues)


def test_a_range_that_runs_past_the_last_row_is_only_a_warning():
    issues, _f, _d = docs_check.check_source(SHEET.replace("| Total | | =SUM(C2:C4) |",
                                                           "| Total | | =SUM(C2:C4) |\n| Other | | =SUM(B2:B30) |"))
    assert any(i["severity"] == "warn" and "past the end" in i["message"] for i in issues), messages(issues)


def test_arithmetic_on_a_text_cell_is_an_error_but_text_functions_are_left_alone():
    issues, _f, _d = docs_check.check_source(SHEET.replace("=SUM(C2:C4)", "=A2*C2"))
    assert any("holds text" in i["message"] for i in errors(issues)), messages(issues)
    issues, _f, _d = docs_check.check_source(SHEET.replace("=SUM(C2:C4)", '=IF(A2="Switch",C2,0)'))
    assert issues == []


def test_formulas_the_checker_cannot_evaluate_are_never_called_wrong():
    # IF, IFERROR, ROUNDUP... are allowed in an export but not evaluated here; their cells are unknown, not text.
    cases = [
        "| Item | Qty |\n| --- | --- |\n| a | 5 |\n| b | =IF(B2>3,10,0) |\n| c | =B3*2 |\n",
        "| Item | Cost |\n| --- | --- |\n| a | =IFERROR(5*150,0) |\n| b | =ROUNDUP(3*9.5,0) |\n"
        "| Total | =SUM(B2:B3) |\n",
        "| Item | Cost |\n| --- | --- |\n| a | 4 |\n| Total | =SUM(B2:B3)+VLOOKUP(\"a\",A2:B2,2,0) |\n",
    ]
    for table in cases:
        issues, _f, _d = docs_check.check_source(f"---\nkind: sheet\ntitle: T\noutline:\n  - X\n---\n## X\n{table}")
        assert [i for i in issues if "text" in i["message"] or "no numbers" in i["message"]] == [], messages(issues)


def test_a_long_forward_chain_of_formulas_is_not_called_wrong():
    rows = "".join(f"| r{n} | =B{n + 3}+1 |\n" for n in range(200)) + "| end | 1 |\n"
    doc = f"---\nkind: sheet\ntitle: T\noutline:\n  - X\n---\n## X\n| Item | N |\n| --- | --- |\n{rows}"
    issues, _f, _d = docs_check.check_source(doc)
    assert errors(issues) == [], messages(issues)[:300]


def test_a_known_mistake_is_still_found_next_to_an_unevaluable_cell():
    table = ("| Item | Cost |\n| --- | --- |\n| a | =IF(1>0,5,0) |\n| Total | =SUM(A2:A3) |\n")
    issues, _f, _d = docs_check.check_source(f"---\nkind: sheet\ntitle: T\noutline:\n  - X\n---\n## X\n{table}")
    assert any("no numbers" in i["message"] for i in errors(issues)), messages(issues)


def test_a_limit_in_the_title_that_no_formula_checks():
    titled = SHEET.replace("Household Costs", "Lab Budget $2000")
    issues, _f, _d = docs_check.check_source(titled)
    assert any("over the 2,000 limit" in i["message"] for i in errors(issues)), messages(issues)
    under = titled.replace("1800", "1000")
    assert docs_check.check_source(under)[0] == []                                  # the total fits
    no_total = titled.replace("| Total | | =SUM(C2:C4) |", "| Sum | | =C2 |")
    assert any("no formula checks" in i["message"] for i in docs_check.check_source(no_total)[0])
    gap = titled.replace("| Total | | =SUM(C2:C4) |", "| Total | | =SUM(C2:C4) |\n| Left | | =2000-C5 |")
    assert docs_check.check_source(gap)[0] == []                                    # a formula shows the gap


def test_limits_are_read_from_titles_and_outlines():
    assert limits_in(["Lab budget $2,000"]) == [2000.0]
    assert limits_in(["$2k spending cap"]) == [2000.0]
    assert limits_in(["Quarterly numbers $5"]) == []                                # no budget word, no limit


def test_error_text_on_a_rendered_sheet_page():
    doc = parse_harvis_doc(SHEET)
    out = docs_check.check_pages(doc, ["Item Qty\n#REF!"], "xlsx")
    assert out and out[0]["severity"] == "error" and "#REF!" in out[0]["message"] and out[0]["where"] == "page 1"


# ---- decks ------------------------------------------------------------------------------------

def _deck_file(mutate=None, source=DECK):
    from pptx import Presentation

    issues, files, doc = docs_check.check_source(source)
    assert issues == [], messages(issues)
    prs = Presentation(io.BytesIO(files["pptx"]))
    if mutate:
        mutate(prs)
    out = io.BytesIO()
    prs.save(out)
    return doc, out.getvalue()


def test_a_clean_deck_has_no_issues():
    assert docs_check.check_source(DECK)[0] == []


def test_a_shape_past_the_slide_and_a_negative_size_are_errors():
    from pptx.util import Inches

    def mutate(prs):
        shape = prs.slides[1].shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(1))
        shape.left = prs.slide_width - Inches(1)
        shape.name = "Runaway"
        neg = prs.slides[1].shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(1))
        neg._element.xpath('.//a:ext')[0].set('cx', '-5')
        neg.name = "Negative"

    doc, data = _deck_file(mutate)
    out = docs_check.check_pptx(data, doc)
    assert any("Runaway" in i["message"] and "edge of the slide" in i["message"] for i in errors(out)), messages(out)
    assert any("Negative" in i["message"] and "negative size" in i["message"] for i in errors(out))
    assert all(i["where"] == "slide 2" for i in out)


def test_text_that_cannot_fit_its_box_is_flagged():
    from pptx.util import Inches, Pt

    def mutate(prs):
        box = prs.slides[1].shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(0.5))
        box.name = "Tiny"
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "This sentence is far too long to sit inside a two inch wide half inch tall box " * 3
        run.font.size = Pt(20)

    doc, data = _deck_file(mutate)
    out = docs_check.check_pptx(data, doc)
    assert any("'Tiny'" in i["message"] and "too long" in i["message"] for i in out), messages(out)


def test_a_slide_without_a_title_and_too_many_bullets():
    source = DECK + "---\n- a\n- b\n" + "".join(f"- bullet {i}\n" for i in range(11))
    issues, _f, _d = docs_check.check_source(source)
    assert any(i["where"] == "slide 3" and "no title" in i["message"] for i in issues), messages(issues)
    assert any(i["where"] == "slide 3" and "bullets" in i["message"] for i in issues), messages(issues)


def test_a_diagram_that_became_a_table_is_reported():
    flow = "".join(f"s{i} -> s{i + 1}\n" for i in range(40))
    source = DECK + f"---\n## Process\n```flow\n{flow}```\n"
    issues, _f, _d = docs_check.check_source(source)
    assert any("drawn as a table" in i["message"] for i in issues), messages(issues)


# ---- rich blocks and pages --------------------------------------------------------------------

def test_dropped_lines_come_from_the_real_parsers():
    assert dropped_lines("timeline", "title: Plan\nKickoff | Oct 5\n| orphan") == ["| orphan"]
    assert dropped_lines("flow", "A -> B\nlonely") == ["lonely"]
    assert dropped_lines("cards", "One | 5\n | no title\ntitle: Servers") == ["| no title"]
    assert dropped_lines("chart", "type: radar\nNorth: 4, x, 6") == ["type: radar", "North: 4, x, 6"]
    assert dropped_lines("chart", "type: bar\nNorth: 4, 6") == []
    rows = "\n".join(f"Q{i}: a, b" for i in range(5))
    assert dropped_lines("quadrant", rows) == ["Q4: a, b"]
    assert dropped_lines("csv", "a,b\n1,2") == []
    assert dropped_lines("csv", "a,b\n1,2,3") == ["(csv rows differ in width)"]
    assert dropped_lines("timeline", "") == ["(empty block)"]
    assert is_empty("flow", "just words") is False and is_empty("chart", "nothing here")


def test_rich_blocks_are_found_and_unclosed_ones_reported():
    blocks, problems = scan_rich("text\n```timeline\nA | Oct 1\n```\n```python\nx = 1\n```\n```cards\nOne | 2\n")
    assert blocks == [("timeline", "A | Oct 1")] and problems == ["cards block never closed"]


def test_check_rich_reports_empty_and_skipped_lines():
    source = DOC + "```flow\nA -> B\nlonely\n```\n```chart\nnothing\n```\n"
    issues, _f, _d = docs_check.check_source(source)
    assert any("skipped" in i["message"] and "lonely" in i["message"] for i in issues), messages(issues)
    assert any("chart block is empty" in i["message"] for i in errors(issues))


def test_a_heading_left_at_the_bottom_of_a_page():
    doc = parse_harvis_doc(DOC.replace("## Summary", "## Summary\nText.\n## Findings\nMore text."))
    pages = ["Memo\nSummary\nText text text text text text text text text.\nFindings",
             "More text goes here with plenty of words to be a normal full page of content."]
    out = docs_check.check_pages(doc, pages, "docx")
    assert [i["where"] for i in out] == ["page 1"] and "findings" in out[0]["message"].lower()
    full = ["Memo\nSummary\nText " + "word " * 40, pages[1]]
    assert docs_check.check_pages(doc, full, "docx") == []


def test_an_empty_rendered_slide_and_the_last_page_exemption():
    doc = parse_harvis_doc(DECK)
    out = docs_check.check_pages(doc, ["Status", ""], "pptx")
    assert [i["where"] for i in out] == ["slide 2"]
    assert docs_check.check_pages(parse_harvis_doc(DOC), ["Memo\nSummary\nshort", "end"], "docx")[0]["where"] == "page 1"


def test_an_export_that_fails_is_reported_once_per_reason():
    issues, files, _d = docs_check.check_source("---\nkind: deck\ntitle: Half")
    assert files == {} and len(issues) == 1 and "header is not finished" in issues[0]["message"]
    assert issues[0]["message"].startswith("Cannot export as pptx or pdf")


# ---- vision -----------------------------------------------------------------------------------

def fake_ollama(monkeypatch, *, installed, loaded=(), vision=(), chat=None, down=False):
    calls = []

    async def ask(path, payload=None, timeout=10):
        calls.append((path, payload))
        if down:
            raise OSError("refused")
        if path == "/api/tags":
            return {"models": [{"name": n} for n in installed]}
        if path == "/api/ps":
            return {"models": [{"name": n} for n in loaded]}
        if path == "/api/show":
            return {"capabilities": ["completion", "vision"] if payload["model"] in vision else ["completion"]}
        if path == "/api/chat":
            return chat(payload) if callable(chat) else {"message": {"content": chat}}
        raise AssertionError(path)

    monkeypatch.setattr(docs_vision, "_ask", ask)
    return calls


def run(coro):
    return asyncio.run(coro)


def test_the_model_named_in_the_request_wins_if_it_can_see(monkeypatch):
    fake_ollama(monkeypatch, installed=["gemma3:4b", "llava:7b"], loaded=["gemma3:4b"], vision=["llava:7b"])
    assert run(docs_vision.pick_model("llava:7b")) == "llava:7b"


def test_a_requested_model_that_cannot_see_falls_back_to_the_loaded_vision_model(monkeypatch):
    fake_ollama(monkeypatch, installed=["qwen3:8b", "gemma4:e2b", "gemma4:e4b", "gemma3:4b"],
                loaded=["gemma4:e4b"], vision=[])
    assert run(docs_vision.pick_model("qwen3:8b")) == "gemma4:e4b"
    assert run(docs_vision.pick_model(None)) == "gemma4:e4b"                       # loaded beats list order


def test_with_nothing_loaded_the_first_installed_fallback_is_chosen(monkeypatch):
    fake_ollama(monkeypatch, installed=["gemma4:e4b", "gemma4:e2b"], loaded=[])
    assert run(docs_vision.pick_model(None)) == "gemma4:e2b"


def test_a_requested_model_ollama_does_not_know_falls_back(monkeypatch):
    calls = fake_ollama(monkeypatch, installed=["gemma3:4b"], loaded=["gemma3:4b"])
    real = docs_vision._ask

    async def unknown(path, payload=None, timeout=10):
        if path == "/api/show" and payload["model"] == "no-such-model:1b":
            raise OSError("404 not found")
        return await real(path, payload, timeout)

    monkeypatch.setattr(docs_vision, "_ask", unknown)
    assert run(docs_vision.pick_model("no-such-model:1b")) == "gemma3:4b"
    assert calls


def test_only_a_few_checks_look_at_once(monkeypatch):
    fake_ollama(monkeypatch, installed=["gemma3:4b"], loaded=["gemma3:4b"], chat="[]")
    monkeypatch.setattr(docs_vision, "_active", docs_vision.MAX_CHECKS)
    result = run(docs_vision.vision_check([png()], lambda i: "page"))
    assert result["busy"] is True and result["done"] == 0 and result["model"] is None
    monkeypatch.setattr(docs_vision, "_active", 0)
    assert run(docs_vision.vision_check([png()], lambda i: "page"))["busy"] is False
    assert docs_vision._active == 0                                                  # the slot is given back


def test_no_vision_model_or_no_ollama_means_none(monkeypatch):
    fake_ollama(monkeypatch, installed=["qwen3:8b"])
    assert run(docs_vision.pick_model(None)) is None
    fake_ollama(monkeypatch, installed=[], down=True)
    assert run(docs_vision.pick_model("anything")) is None


def test_replies_are_read_tolerantly():
    p = docs_vision.parse_problems
    assert p("[]") == [] and p("No problems here.") == []
    assert p('Sure! ["Text is cut off on the right", "Chart is empty"] Hope that helps') == \
        ["Text is cut off on the right", "Chart is empty"]
    assert p('```json\n["Overlapping boxes"]\n```') == ["Overlapping boxes"]
    assert p('["Low contrast text", "Overlapping sha') == ["Low contrast text"]
    assert p('[{"problem": "Slide is empty"}]') == ["Slide is empty"]
    assert p('["a", "b", "c", "d", "none", "Four problem words", "Another long problem", "Yet another one"]') == \
        ["Four problem words", "Another long problem", "Yet another one"]
    assert p('["No issues"]') == []


def test_yes_no_answers_become_problems_and_a_list_reply_still_works():
    r = docs_vision.parse_reply
    assert r('{"text_runs_off_the_page_edge": false, "two_shapes_overlap": false, "text_too_faint_to_read": false}') == []
    assert r('{"text_runs_off_the_page_edge": true, "two_shapes_overlap": false, "text_too_faint_to_read": true}') == \
        ["text runs off the edge of the page or out of its box", "text too faint to read"]
    assert r('<think>hm [x]</think>{"two_shapes_overlap": true}') == ["shapes drawn on top of each other"]
    assert r('{"two_shapes_overlap": "yes"}') == [] and r("{cut off") == []
    assert r('["Overlapping boxes"]') == ["Overlapping boxes"]


def test_pages_are_downscaled_and_chosen_evenly():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (2000, 1000), "white").save(buf, "PNG")
    small = Image.open(io.BytesIO(docs_vision.downscale(buf.getvalue())))
    assert small.size == (896, 448)
    assert docs_vision.pick_pages(3) == [0, 1, 2]
    assert docs_vision.pick_pages(20) == [0, 4, 8, 11, 15, 19] and len(docs_vision.pick_pages(100)) == 6


def png(w=1000, h=600, blank=False):
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (w, h), "white")
    if not blank:
        ImageDraw.Draw(im).rectangle((100, 100, 500, 300), fill="navy")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def test_the_vision_check_sends_one_small_image_per_request(monkeypatch):
    calls = fake_ollama(monkeypatch, installed=["gemma3:4b"], loaded=["gemma3:4b"],
                        chat='{"text_runs_off_the_page_edge": true, "two_shapes_overlap": false, '
                             '"text_too_faint_to_read": false}')
    result = run(docs_vision.vision_check([png()] * 8, lambda i: f"slide {i + 1}"))
    chats = [p for path, p in calls if path == "/api/chat"]
    assert len(chats) == 6 and result["done"] == 6 and result["model"] == "gemma3:4b"
    first = chats[0]
    assert first["stream"] is False and first["keep_alive"] == "5m"
    assert first["options"] == {"num_ctx": 4096, "temperature": 0} and len(first["messages"][0]["images"]) == 1
    assert first["format"] == docs_vision.SCHEMA
    assert all(i["severity"] == "warn" and i["message"] == "Looks like text runs off the edge of the page or out of its box"
               for i in result["issues"])
    assert result["issues"][0]["where"] == "slide 1"


def test_a_failing_page_is_skipped_and_the_rest_still_count(monkeypatch):
    seen = {"n": 0}

    def chat(_payload):
        seen["n"] += 1
        if seen["n"] == 2:
            raise OSError("timed out")
        return {"message": {"content": "I cannot tell"}}

    fake_ollama(monkeypatch, installed=["gemma3:4b"], loaded=["gemma3:4b"], chat=chat)
    result = run(docs_vision.vision_check([png()] * 3, lambda i: f"page {i + 1}"))
    assert result["done"] == 2 and result["skipped"] == 1 and result["issues"] == []


def test_a_blank_page_is_reported_from_its_pixels_without_asking_the_model(monkeypatch):
    calls = fake_ollama(monkeypatch, installed=["gemma3:4b"], loaded=["gemma3:4b"], chat="[]")
    result = run(docs_vision.vision_check([png(blank=True), png()], lambda i: f"slide {i + 1}"))
    assert [c for c, _ in calls].count("/api/chat") == 1 and result["done"] == 2
    assert result["issues"] == [{"severity": "warn", "where": "slide 1", "message": "Looks like the page is blank"}]


@pytest.mark.parametrize("pages", [[], None])
def test_no_pages_means_no_model_call(monkeypatch, pages):
    calls = fake_ollama(monkeypatch, installed=["gemma3:4b"])
    result = run(docs_vision.vision_check(pages or [], lambda i: "page"))
    assert result["done"] == 0 and not [c for c in calls if c[0] == "/api/chat"]


def test_a_deck_that_forgot_the_slide_lines_still_gets_its_slides():
    from plugins.hermes_ui.docs_parse import parse_harvis_doc
    src = "---\nkind: deck\ntitle: T\n---\n# T\nsub\n## One\n- a\n```timeline\n## not a slide | Oct 1\n```\n## Two\n- b"
    doc = parse_harvis_doc(src)
    assert [s.title for s in doc.sections] == ["T", "One", "Two"]


def test_content_on_the_title_slide_is_flagged():
    from plugins.hermes_ui.docs_check import check_source
    src = "---\nkind: deck\ntitle: T\n---\n# T\nsub\n| a | b |\n|---|---|\n| 1 | 2 |\n---\n## One\n- a"
    issues, _files, _doc = check_source(src)
    assert any(i["where"] == "slide 1" and "title slide" in i["message"] for i in issues)
    clean = "---\nkind: deck\ntitle: T\n---\n# T\nsub\n---\n## One\n- a"
    assert not any(i["where"] == "slide 1" for i in check_source(clean)[0])


def test_a_slide_listed_in_the_outline_but_never_written_is_flagged():
    from plugins.hermes_ui.docs_check import check_source
    src = ("---\nkind: deck\ntitle: Backups: Why You Need Them\noutline:\n  - The Risk of Data Loss\n"
           "  - Types of Backups\n  - Regular Testing is Key\n---\n# Backups: Why You Need Them\nsub\n---\n"
           "## The Risk of Data Loss\n- a\n---\n## Regular Testing is Key\n- b")
    issues = check_source(src)[0]
    outline = [i for i in issues if i["where"] == "outline"]
    assert len(outline) == 1 and "Types of Backups" in outline[0]["message"]
    assert "Risk" not in outline[0]["message"]


def test_an_outline_that_matches_the_body_or_names_the_title_slide_is_clean():
    from plugins.hermes_ui.docs_check import check_source
    with_title = ("---\nkind: deck\ntitle: Status\noutline:\n  - Title\n  - Plan\n---\n# Status\nsub\n---\n"
                  "## Plan\n- Ship it")
    lengthened = ("---\nkind: deck\ntitle: Status\noutline:\n  - Risk\n---\n# Status\nsub\n---\n"
                  "## The Risk of Data Loss\n- a")
    assert not [i for i in check_source(with_title)[0] if i["where"] == "outline"]
    assert not [i for i in check_source(lengthened)[0] if i["where"] == "outline"]
    assert not [i for i in check_source(DECK)[0] if i["where"] == "outline"]
    assert not [i for i in check_source(SHEET)[0] if i["where"] == "outline"]


# ---- placeholder text -------------------------------------------------------------------------

def _deck_with(bullet: str, *, title: str = "Plan", notes: str = "") -> str:
    tail = f"\nNotes: {notes}" if notes else ""
    return (f"---\nkind: deck\ntitle: Status\n---\n# Status\nA short deck\n---\n## {title}\n- {bullet}\n"
            f"- A second point that is written out in full{tail}")


def _placeholder_issues(source: str) -> list:
    return [i for i in docs_check.check_source(source)[0] if "Placeholder" in i["message"]]


@pytest.mark.parametrize("text", [
    "Lorem ipsum dolor sit amet", "Owner: TBD", "Date: tba", "TODO: write this", "Insert image here",
    "Click to add text", "[Company Name] Inc", "[insert date]", "Call xxx-xxx-xxxx", "Your name here",
    "Description goes here", "**Insert logo here**", "(insert date here)", "Insert chart here.",
    "Insert your company name here"])
def test_each_placeholder_is_an_error_with_its_location(text):
    hits = _placeholder_issues(_deck_with(text))
    assert len(hits) == 1, messages(hits)
    assert hits[0]["severity"] == "error" and hits[0]["where"] == "slide 2"


def test_a_placeholder_is_found_in_titles_notes_cells_and_the_outline():
    in_title = _placeholder_issues(_deck_with("Real content here for the plan", title="TBD"))
    assert [i["where"] for i in in_title] == ["slide 2"] and "title" in in_title[0]["message"]
    in_notes = _placeholder_issues(_deck_with("Real content for the plan", notes="TBD, say something"))
    assert in_notes and "speaker notes" in in_notes[0]["message"]
    in_cell = _placeholder_issues(SHEET.replace("| Cable | 10 | 50 |", "| Cable | TBD | 50 |"))
    assert in_cell[0]["where"] == "sheet Items" and "cell B3" in in_cell[0]["message"]
    in_outline = _placeholder_issues(DECK.replace("  - Plan\n", "  - Plan\n  - Lorem ipsum\n"))
    assert [i["where"] for i in in_outline] == ["outline"]
    in_doc = _placeholder_issues(DOC.replace("Plain words about the week.", "Lorem ipsum dolor sit amet."))
    assert in_doc[0]["where"] == "section 'Summary'"


def test_a_header_title_that_is_a_placeholder_is_reported_once():
    assert len(_placeholder_issues(DECK.replace("title: Status", "title: TBD").replace("# Status", "# TBD"))) == 1
    assert _placeholder_issues(DECK.replace("title: Status", "title: TBD"))[0]["where"] == "document"


@pytest.mark.parametrize("source", [DECK, SHEET, DOC])
def test_the_clean_fixtures_have_no_placeholders(source):
    assert _placeholder_issues(source) == []


def test_ordinary_words_that_look_like_placeholders_are_left_alone():
    fine = [
        "A TODO list app for teams", "Build a TODO app", "Read the [docs](https://example.com/x) first",
        "- [ ] ship it", "See note [1] for sources", "Set `api_key = insert key here` in the file",
        "The tbdx module", "Revenue grew XXL sizes", "An outbound link", "Use the ![chart](x.png) image",
        "Insert the SD card into the slot here", "Click Insert > Table and the new table appears here",
        "To add a row, insert it here and press Enter", "Then insert the second card here, next to the first",
    ]
    for text in fine:
        assert _placeholder_issues(_deck_with(text)) == [], text
    assert _placeholder_issues(_deck_with("A real point", title="TODO list design")) == []
    code = "```bash\nexport TOKEN=xxx   # TBD\n```"
    assert _placeholder_issues(DOC.replace("Plain words about the week.", code)) == []
    cards = "```cards\nTitle\nLorem ipsum | one\n```"
    assert _placeholder_issues(DOC.replace("Plain words about the week.", cards))


def test_the_placeholder_report_is_capped():
    many = "\n".join(f"- TBD item {i}" for i in range(40))
    assert len(docs_check.check_placeholders(parse_harvis_doc(f"---\nkind: doc\ntitle: T\n---\n## S\n{many}"))) == 1
    sections = "\n".join(f"## S{i}\nTBD" for i in range(40))
    assert len(docs_check.check_placeholders(parse_harvis_doc(f"---\nkind: doc\ntitle: T\n---\n{sections}"))) == 20


def _deck_with_table(table: str) -> str:
    return f"---\nkind: deck\ntitle: Status\n---\n# Status\nA short deck\n---\n## Plan\n{table}\n"


TRACKER = """---
kind: sheet
title: Tasks
---
## Tasks
| Task | Owner | Status |
| --- | --- | --- |
| Wire the rack | Sam | TODO |
| Label cables | Ana | TODO |
| Order switch | Lee | Done |
"""


def test_a_todo_status_in_a_task_sheet_is_not_a_placeholder():
    from owui_compat.docs_tools_checks import validate_source

    assert _placeholder_issues(TRACKER) == []
    out = validate_source(TRACKER)
    assert out["ok"], out["problems"]
    table = "| Task | Status |\n| --- | --- |\n| Wire the rack | TODO |\n| Label cables |TODO|"
    assert _placeholder_issues(_deck_with_table(table)) == []


def test_a_todo_note_in_a_sheet_cell_or_table_is_still_a_placeholder():
    note = TRACKER.replace("| TODO |\n| Label", "| TODO: ask Sam |\n| Label")
    assert [i["where"] for i in _placeholder_issues(note)] == ["sheet Tasks"]
    prose = TRACKER.replace("| Done |", "| TODO wire it up |")
    assert "cell C4" in _placeholder_issues(prose)[0]["message"]
    in_table = "| Task | Status |\n| --- | --- |\n| Wire the rack | TODO: later |"
    assert _placeholder_issues(_deck_with_table(in_table))


def test_validate_source_reports_placeholders_too():
    from owui_compat.docs_tools_checks import validate_source

    out = validate_source(_deck_with("Lorem ipsum dolor sit amet"))
    assert not out["ok"] and any("slide 2" in p and "Placeholder" in p for p in out["problems"])
    assert validate_source(DECK)["ok"] and validate_source(SHEET)["ok"]


# ---- error values in rendered pages -----------------------------------------------------------

@pytest.mark.parametrize("shown", ["#NUM!", "#NULL!", "#SPILL!", "#CALC!", "#REF!", "#N/A"])
def test_every_spreadsheet_error_in_rendered_text_is_flagged(shown):
    issues = docs_check.check_pages(parse_harvis_doc(SHEET), [f"Item Cost\nTotal {shown}"], "xlsx")
    assert errors(issues) and shown in issues[0]["message"] and issues[0]["where"] == "page 1"


def test_a_libreoffice_error_code_in_rendered_text_is_flagged():
    issues = docs_check.check_pages(parse_harvis_doc(SHEET), ["Total Err:502", "ok", "Err:508 and Err:502"], "xlsx")
    assert [i["where"] for i in issues] == ["page 1", "page 3"]
    assert "Err:502" in issues[0]["message"] and "Err:508" in issues[1]["message"]
    assert docs_check.check_pages(parse_harvis_doc(SHEET), ["Total 2150", "Error 5025"], "xlsx") == []
