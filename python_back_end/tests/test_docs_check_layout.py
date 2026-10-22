"""Layout checks on decks (text that does not fit) and glyph coverage, with fonts mocked where it matters."""

import io
import logging
import os
import sys

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from plugins.hermes_ui import docs_check, docs_check_layout as layout  # noqa: E402
from plugins.hermes_ui.docs_parse import parse_harvis_doc  # noqa: E402
from tests.test_docs_check import DECK as CHECK_DECK  # noqa: E402
from tests.test_docs_export import DECK as EXPORT_DECK  # noqa: E402
from tests.test_docs_rich import ALL_SLIDES, deck as rich_deck  # noqa: E402

LAYOUT_WORDS = ("too long for its box", "runs off the slide", "too tall for the slide", "hold more text",
                "No installed font")


def deck(*slides: str, title: str = "Status") -> str:
    return f"---\nkind: deck\ntitle: {title}\n---\n# {title}\nA subtitle\n---\n" + "\n---\n".join(slides) + "\n"


def layout_issues(source: str) -> list:
    _issues, files, doc = docs_check.check_source(source)
    return layout.check_layout(doc, files["pptx"])


def blank_deck(build) -> bytes:
    """A one-slide .pptx whose shapes ``build(slide)`` adds, 13.333 x 7.5 in like the writer's."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    build(prs.slides.add_slide(prs.slide_layouts[6]))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def textbox(slide, left, top, width, height, text, size=18, anchor=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = tb.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.paragraphs[0].add_run().text = text
    frame.paragraphs[0].runs[0].font.size = Pt(size)
    return tb


@pytest.fixture
def average_widths(monkeypatch):
    """No font files: text is measured with the average-width model, so results do not depend on the host's fonts."""
    monkeypatch.setattr(layout, "_face", lambda bold, mono: None)


# ---- measuring --------------------------------------------------------------------------------

def test_wide_characters_count_one_em_each(average_widths):
    assert layout.text_width("中文", 20) == 40
    assert layout.text_width("ab", 20) == pytest.approx(2 * 20 * 0.50)


def test_words_wrap_greedily_and_a_word_wider_than_the_line_is_split(average_widths):
    assert layout.count_lines("aaaa bbbb", 100, 10) == 1             # 9 chars x 5pt = 45pt
    assert layout.count_lines("aaaa bbbb", 30, 10) == 2
    assert layout.count_lines("x" * 100, 100, 10) == 5               # 500pt over a 100pt line
    assert layout.count_lines("one\ntwo", 500, 10) == 2


# ---- boxes ------------------------------------------------------------------------------------

def test_an_overflowing_text_box_is_flagged(average_widths):
    data = blank_deck(lambda s: textbox(s, 1, 1, 3, 1, "word " * 60))
    out = layout.deck_issues(data)
    assert len(out) == 1
    assert out[0]["severity"] == "warn" and out[0]["where"] == "slide 1"
    assert "too long for its box" in out[0]["message"] and "split the slide" in out[0]["message"]


def test_text_that_fits_its_box_is_clean(average_widths):
    data = blank_deck(lambda s: textbox(s, 1, 1, 6, 1, "A short line of text that fits easily"))
    assert layout.deck_issues(data) == []


def test_text_running_past_the_bottom_of_the_slide_is_flagged(average_widths):
    data = blank_deck(lambda s: textbox(s, 1, 6.6, 3, 0.8, "word " * 80))
    out = layout.deck_issues(data)
    assert len(out) == 1 and "runs off the slide" in out[0]["message"]


def test_a_shape_the_other_check_already_named_is_not_reported_again(average_widths):
    data = blank_deck(lambda s: textbox(s, 1, 1, 3, 1, "word " * 60))
    name = Presentation(io.BytesIO(data)).slides[0].shapes[0].name
    known = [{"severity": "warn", "where": "slide 1", "message": f"The text in '{name}' looks too long for its box"}]
    assert layout.deck_issues(data, known) == []
    assert len(layout.deck_issues(data, [dict(known[0], where="slide 2")])) == 1


def test_several_overflowing_boxes_on_a_slide_are_one_issue(average_widths):
    def build(slide):
        for i in range(4):
            textbox(slide, 0.5 + 3 * i, 1, 2.8, 0.6, "word " * 40)

    out = layout.deck_issues(blank_deck(build))
    assert len(out) == 1 and "4 text boxes" in out[0]["message"]


def test_no_wrap_and_shrink_to_fit_shapes_are_left_alone(average_widths):
    def build(slide):
        a = textbox(slide, 1, 1, 1, 0.3, "word " * 60)
        a.text_frame.word_wrap = False
        b = textbox(slide, 1, 3, 1, 0.3, "word " * 60)
        b.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    assert layout.deck_issues(blank_deck(build)) == []


def test_a_table_taller_than_the_slide_once_its_text_wraps_is_flagged(average_widths):
    def build(slide):
        table = slide.shapes.add_table(9, 2, Inches(0.7), Inches(1.5), Inches(3), Inches(2.7)).table
        for r in range(9):
            for c in range(2):
                run = table.cell(r, c).text_frame.paragraphs[0].add_run()
                run.text = "a long cell that wraps over several lines " * 3
                run.font.size = Pt(14)

    out = layout.deck_issues(blank_deck(build))
    assert len(out) == 1 and "too tall for the slide" in out[0]["message"]


def test_a_small_table_is_clean(average_widths):
    def build(slide):
        table = slide.shapes.add_table(3, 2, Inches(0.7), Inches(1.5), Inches(8), Inches(1.2)).table
        for r in range(3):
            table.cell(r, 0).text_frame.paragraphs[0].add_run().text = "Owner"

    assert layout.deck_issues(blank_deck(build)) == []


# ---- decks from the writer --------------------------------------------------------------------

def test_a_title_far_too_long_for_its_box_is_flagged():
    out = layout_issues(deck("## " + "这是一个非常长的标题" * 9 + "\n- point"))
    hit = [i for i in out if "title is too long" in i["message"]]
    assert hit and hit[0]["where"] == "slide 2" and "Shorten it" in hit[0]["message"]


def test_the_writers_own_example_decks_have_no_layout_issues():
    rich = rich_deck(*ALL_SLIDES)
    dense = deck("## Plan\n" + "\n".join("- Revenue grew 14% year over year, driven by enterprise renewals (Q3 2026)"
                                         for _ in range(8)),
                 "## QUARTERLY REVENUE REVIEW AND OUTLOOK\n- ONE\n- TWO",
                 "## A moderately long slide title that still fits on one line of the slide\n- point")
    for source in (CHECK_DECK, EXPORT_DECK, rich, dense, deck("## Plan\n- Ship it\n- Test it")):
        assert layout_issues(source) == [], source[:60]


def test_normal_slides_do_not_trip_the_real_font_metrics():
    out = layout.deck_issues(blank_deck(lambda s: textbox(s, 0.7, 1.75, 11.9, 5.2, "Revenue grew 14% year over year. " * 24)))
    assert out == []


# ---- glyphs -----------------------------------------------------------------------------------

LATIN = "---\nkind: deck\ntitle: Café\n---\n# Café “quoted” — done… · 5 × 3 ≥ 2 → ok ✓ ☑ ☐ €\nsub\n---\n## Ünïcode\n- naïve\n"


def covering(*ranges):
    return lambda cp: any(lo <= cp <= hi for lo, hi in ranges)


def test_ordinary_text_is_not_flagged_by_the_real_fonts():
    assert layout.glyph_issues(parse_harvis_doc(LATIN)) == []


def test_cjk_is_flagged_only_when_no_installed_font_has_it(monkeypatch):
    doc = parse_harvis_doc(deck("## 中文标题\n- 内容"))
    monkeypatch.setattr(layout, "_coverage", lambda: covering((0, 0x2FFF)))
    out = layout.glyph_issues(doc)
    assert len(out) == 1 and out[0]["where"] == "document" and out[0]["severity"] == "warn"
    assert "中 (U+4E2D)" in out[0]["message"] and "slide 2" in out[0]["message"]
    monkeypatch.setattr(layout, "_coverage", lambda: covering((0, 0x2FFF), (0x4E00, 0x9FFF)))
    assert layout.glyph_issues(doc) == []


def test_many_missing_characters_are_one_issue_with_a_few_samples(monkeypatch):
    doc = parse_harvis_doc(deck("## " + "一二三四五六七八九十" + "\n- a"))
    monkeypatch.setattr(layout, "_coverage", lambda: covering((0, 0x2FFF)))
    out = layout.glyph_issues(doc)
    assert len(out) == 1
    assert out[0]["message"].count("(U+") == layout.SAMPLE_CHARS and "and 5 more" in out[0]["message"]


def test_joiners_variation_selectors_and_controls_need_no_glyph(monkeypatch):
    doc = parse_harvis_doc(deck("## a‍️​ b\n- c\tc"))
    monkeypatch.setattr(layout, "_coverage", lambda: covering((0x20, 0x7E)))
    assert layout.glyph_issues(doc) == []


def test_notes_are_not_checked_because_they_are_not_drawn(monkeypatch):
    doc = parse_harvis_doc(deck("## Plan\n- ok\nNotes: 中文"))
    monkeypatch.setattr(layout, "_coverage", lambda: covering((0, 0x2FFF)))
    assert layout.glyph_issues(doc) == []


def test_no_font_list_means_no_claim(monkeypatch):
    monkeypatch.setattr(layout, "_coverage", lambda: None)
    assert layout.glyph_issues(parse_harvis_doc(deck("## 中文\n- a"))) == []


def test_coverage_comes_from_the_installed_fonts():
    covered = layout._coverage()
    if covered is None:
        pytest.skip("no way to list installed fonts here")
    assert covered(ord("a")) and covered(ord("é")) and not covered(0x10FFFE)


# ---- never breaks Check -----------------------------------------------------------------------

def test_a_crash_in_the_layout_check_is_logged_and_skipped(monkeypatch, caplog):
    def boom(*_a, **_k):
        raise RuntimeError("no font")

    monkeypatch.setattr(layout, "deck_issues", boom)
    monkeypatch.setattr(layout, "glyph_issues", boom)
    with caplog.at_level(logging.WARNING, logger=layout.log.name):
        issues, files, _doc = docs_check.check_source(CHECK_DECK)
    assert "pptx" in files and not any(i["message"].startswith("No installed font") for i in issues)
    assert len([r for r in caplog.records if "check failed" in r.getMessage()]) == 2


def test_a_corrupt_deck_does_not_raise(monkeypatch):
    doc = parse_harvis_doc(CHECK_DECK)
    assert [i for i in layout.check_layout(doc, b"not a zip") if "slide" in i["where"]] == []


def test_check_source_reports_layout_issues(monkeypatch):
    monkeypatch.setattr(layout, "_coverage", lambda: covering((0, 0x2FFF)))
    issues, _files, _doc = docs_check.check_source(deck("## 中文\n- a"))
    assert any(i["message"].startswith("No installed font") for i in docs_check.sort_issues(issues))


def test_a_layout_module_that_will_not_import_is_logged_and_skipped(monkeypatch, caplog):
    monkeypatch.setitem(sys.modules, layout.__name__, None)
    with caplog.at_level(logging.WARNING):
        issues, files, _doc = docs_check.check_source(CHECK_DECK)
    assert "pptx" in files and issues is not None
    assert any("layout check unavailable" in r.getMessage() for r in caplog.records)


# ---- placeholder roles on the writer's own decks ----------------------------------------------

def _shape(files_pptx: bytes, slide: int, name: str):
    return next(s for s in Presentation(io.BytesIO(files_pptx)).slides[slide].shapes if s.name == name)


def test_the_writers_subtitle_and_body_placeholders_get_their_own_roles():
    _issues, files, _doc = docs_check.check_source(deck("## Plan\n- one\n- two"))
    prs = Presentation(io.BytesIO(files["pptx"]))
    roles = {s.placeholder_format.type: layout._role(s) for sl in prs.slides for s in sl.placeholders}
    assert {"SUBTITLE": "subtitle", "OBJECT": "body"}.items() <= {str(k).split(" ")[0]: v for k, v in roles.items()}.items()
    assert layout._role(_shape(files["pptx"], 0, "Subtitle 2")) == "subtitle"


def test_a_long_subtitle_is_blamed_on_the_subtitle_not_the_title():
    source = f"---\nkind: deck\ntitle: Short\n---\n# Short\n{'中文字幕' * 30}\n---\n## Next\n- a\n"
    msgs = [i["message"] for i in layout_issues(source) if i["where"] == "slide 1"]
    assert not any(m.startswith("The title") for m in msgs)


def test_an_overflowing_body_placeholder_is_named_the_body_text():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.placeholders[0].text_frame.text = "Plan"
    body = slide.placeholders[1].text_frame
    body.text = "A bullet that is fairly long so the list needs room"
    for _ in range(29):
        body.add_paragraph().text = "Another bullet that is fairly long so the list needs room"
    out = io.BytesIO()
    prs.save(out)
    msgs = [i["message"] for i in layout.deck_issues(out.getvalue())]
    assert msgs and msgs[0].startswith("The body text")
