"""Text in exports: file names in any script, CJK glyphs, italic in PDF, the title line of a cards block, bold italic."""

import io
import os
import sys
from urllib.parse import unquote

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from plugins.hermes_ui import docs_fonts, docs_pdf  # noqa: E402
from plugins.hermes_ui.docs_blocks import inline_runs, parse_blocks  # noqa: E402
from plugins.hermes_ui.docs_check_rich import dropped_lines, late_card_titles  # noqa: E402
from plugins.hermes_ui.docs_export import content_disposition, export_doc, slugify  # noqa: E402
from plugins.hermes_ui.docs_parse import parse_harvis_doc  # noqa: E402
from plugins.hermes_ui.docs_rich import caption_of, parse_cards, rich_as_rows  # noqa: E402
from plugins.hermes_ui.docs_svg import svg_markup  # noqa: E402


def _doc(title="T", body="plain text"):
    return f"---\nkind: doc\ntitle: {title}\n---\n\n## S\n\n{body}\n"


# ---- (a) file names ----

@pytest.mark.parametrize("title,expected", [
    ("数据报告", "数据报告"),
    ("Café plan", "café-plan"),
    ("Q3 Security Review!", "q3-security-review"),
    ("Отчёт за квартал", "отчёт-за-квартал"),
    ("Board / Deck", "board-deck"),
    ("..\\..\\etc/passwd", "etc-passwd"),
    ('say "hi": <a|b>?*', "say-hi-a-b"),
    ("tab\there\nnew\x00line‮", "tab-here-new-line"),
    ("Ｆｕｌｌ ｗｉｄｔｈ", "full-width"),
])
def test_slugify_keeps_letters_and_digits_of_any_script(title, expected):
    assert slugify(title) == expected


def test_slugify_falls_back_only_when_nothing_is_left():
    assert slugify("") == "document" and slugify(None) == "document"
    assert slugify("///???") == "document"


def test_slugify_is_capped_and_never_a_windows_device_name():
    assert len(slugify("a" * 500)) == 60
    assert len(slugify("数" * 500).encode()) <= 150 and len(slugify("数" * 500)) == 50
    assert slugify("CON") == "con-file" and slugify("lpt1") == "lpt1-file"


def test_content_disposition_carries_an_ascii_fallback_and_the_real_name():
    header = content_disposition("数据报告-café.pdf")
    assert header.startswith('attachment; filename="cafe.pdf"; ')
    assert content_disposition("数据报告.pdf").startswith('attachment; filename="document.pdf"; ')
    assert header.endswith("filename*=UTF-8''%E6%95%B0%E6%8D%AE%E6%8A%A5%E5%91%8A-caf%C3%A9.pdf")
    assert unquote(header.rsplit("''", 1)[1]) == "数据报告-café.pdf"
    assert content_disposition("café-plan.pdf").startswith('attachment; filename="cafe-plan.pdf"; ')
    assert content_disposition("report.pdf") == "attachment; filename=\"report.pdf\"; filename*=UTF-8''report.pdf"
    header.encode("ascii")


def test_content_disposition_survives_hostile_names():
    header = content_disposition('a"b\r\nSet-Cookie: x.pdf')
    assert "\r" not in header and "\n" not in header
    assert header.count('"') == 2


def test_export_names_the_file_after_the_title_in_its_own_script():
    out = export_doc(_doc(), "md", "数据报告 Café")
    assert out.filename == "数据报告-café.md"
    assert export_doc(_doc("Café plan"), "md").filename == "café-plan.md"


def test_the_route_header_has_both_names():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from auth_optimized import get_current_user_optimized
    from plugins.hermes_ui.rest_docs import router

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_current_user_optimized] = lambda: {"id": 1}
    r = TestClient(app).post("/hermes-api/api/harvis/docs/export",
                             json={"source": _doc(), "format": "md", "title": "数据报告 Café"})
    assert r.status_code == 200
    header = r.headers["content-disposition"]
    assert 'filename="cafe.md"' in header
    assert unquote(header.split("filename*=UTF-8''")[1]) == "数据报告-café.md"


# ---- (b) CJK fonts ----

def _fonts_of(pdf: bytes) -> set[str]:
    pypdf = pytest.importorskip("pypdf")
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    return {str(f.get_object()["/BaseFont"]) for p in reader.pages for f in (p["/Resources"].get("/Font") or {}).values()}


def test_a_cjk_font_is_added_to_the_stack_when_one_is_installed(monkeypatch):
    monkeypatch.setattr(docs_fonts, "cjk_family", lambda: "Noto Sans CJK SC")
    assert docs_fonts.font_stack("DejaVu Sans", "Arial", "sans-serif") == '"DejaVu Sans", "Arial", "Noto Sans CJK SC", sans-serif'
    assert docs_fonts.font_stack("DejaVu Sans Mono", "monospace").endswith('"Noto Sans CJK SC", monospace')
    assert "Noto Sans CJK SC" in docs_pdf._fonts()[0]


def test_nothing_is_added_and_nothing_breaks_without_a_cjk_font(monkeypatch):
    monkeypatch.setattr(docs_fonts, "cjk_family", lambda: None)
    assert docs_fonts.font_stack("DejaVu Sans", "sans-serif") == '"DejaVu Sans", sans-serif'
    out = export_doc(_doc("数据报告", "这是一个测试 text"), "pdf")
    assert out.data.startswith(b"%PDF")
    assert export_doc(_doc("数据报告", "这是一个测试"), "docx").data[:2] == b"PK"


def test_the_cjk_family_is_found_through_fontconfig(monkeypatch):
    listing = ("Noto Serif CJK JP,Noto Serif CJK JP Medium\nNoto Sans Mono CJK SC\n"
               "Noto Sans CJK JP,Noto Sans CJK JP Medium\nNoto Sans CJK SC\nWenQuanYi Micro Hei\n")
    monkeypatch.setattr(docs_fonts, "_fc", lambda *a: listing)
    docs_fonts.cjk_family.cache_clear()
    try:
        assert docs_fonts.cjk_family() == "Noto Sans CJK SC"
        monkeypatch.setattr(docs_fonts, "_fc", lambda *a: "")
        docs_fonts.cjk_family.cache_clear()
        assert docs_fonts.cjk_family() is None
    finally:
        docs_fonts.cjk_family.cache_clear()


def test_a_missing_fontconfig_is_not_an_error(monkeypatch):
    def boom(*a, **k):
        raise FileNotFoundError("fc-list")
    monkeypatch.setattr(docs_fonts.subprocess, "run", boom)
    docs_fonts.cjk_family.cache_clear()
    docs_fonts.has_real_italic.cache_clear()
    docs_fonts.has_real_mono_italic.cache_clear()
    try:
        assert docs_fonts.cjk_family() is None and docs_fonts.has_real_italic() is False
        assert docs_fonts.has_real_mono_italic() is False
    finally:
        docs_fonts.cjk_family.cache_clear()
        docs_fonts.has_real_italic.cache_clear()
        docs_fonts.has_real_mono_italic.cache_clear()


@pytest.mark.skipif(docs_fonts.cjk_family() is None, reason="no CJK font installed here")
def test_cjk_text_is_drawn_with_a_cjk_font():
    fonts = _fonts_of(export_doc(_doc("数据报告", "这是一个测试"), "pdf").data)
    assert any("CJK" in f or "Han" in f or "WenQuanYi" in f or "Fallback" in f for f in fonts), fonts


# ---- (c) italic ----

def _ink_lean(pdf: bytes) -> float:
    """How far the top half of the first body line sits right of its bottom half, in pixels."""
    pdfium = pytest.importorskip("pypdfium2")
    img = pdfium.PdfDocument(pdf)[0].render(scale=4).to_pil().convert("L")
    w, h = img.size
    px = img.load()
    blocks, start = [], None
    for y in range(h + 1):
        ink = y < h and any(px[x, y] < 128 for x in range(w))
        if ink and start is None:
            start = y
        elif not ink and start is not None:
            blocks.append((start, y - 1))
            start = None
    top, bottom = blocks[1]  # the first block is the document title, the last the page number
    mid = (top + bottom) // 2

    def centre(lo, hi):
        xs = [x for y in range(lo, hi) for x in range(w) if px[x, y] < 128]
        return sum(xs) / len(xs)
    return centre(top, mid) - centre(mid + 1, bottom + 1)


@pytest.mark.parametrize("has_italic", [True, False])
def test_an_italic_run_is_marked_italic_in_the_html(monkeypatch, has_italic):
    monkeypatch.setattr(docs_pdf, "has_real_italic", lambda: has_italic)
    html = docs_pdf.build_html(parse_harvis_doc(_doc(body="plain *slanted words* and ***both***")), "T")
    assert "<em>" in html
    assert ('class="fi"' in html) is (not has_italic)
    assert ".fi {" in html if not has_italic else ".fi {" not in html


def test_italic_markup_keeps_bold_code_and_links_inside_the_fallback(monkeypatch):
    monkeypatch.setattr(docs_pdf, "has_real_italic", lambda: False)
    html = docs_pdf.build_html(parse_harvis_doc(_doc(body="*one [two](https://example.com) three*")), "T")
    assert '<a href="https://example.com"><em><span class="fi">two</span></em></a>' in html
    assert '<span class="fi">one</span> <span class="fi">' in html


def test_the_pdf_names_an_italic_face_for_an_italic_run():
    fonts = _fonts_of(export_doc(_doc(body="plain *slanted*"), "pdf").data)
    assert any("Oblique" in f or "Italic" in f for f in fonts), fonts
    assert not any("Oblique" in f or "Italic" in f for f in _fonts_of(export_doc(_doc(body="plain upright"), "pdf").data))


@pytest.mark.parametrize("has_italic", [True, False])
def test_italic_text_is_actually_slanted_on_the_page(monkeypatch, has_italic):
    """Both paths slant: a real face when the image has one, the skew when it does not."""
    if has_italic and not docs_fonts.has_real_italic():
        pytest.skip("no real italic face installed here")
    monkeypatch.setattr(docs_pdf, "has_real_italic", lambda: has_italic)
    flat = _ink_lean(export_doc("---\nkind: pdf\n---\n\nIIIIIIIIIIIIIIII IIIIIIIIIIIIIIII\n", "pdf").data)
    leaning = _ink_lean(export_doc("---\nkind: pdf\n---\n\n*IIIIIIIIIIIIIIII IIIIIIIIIIIIIIII*\n", "pdf").data)
    assert leaning - flat > 1.0, (flat, leaning)


@pytest.mark.parametrize("body_italic,mono_italic,skewed", [
    (False, True, "<code>grep</code>"),
    (True, False, '<span class="fi"><code>grep</code></span>'),
])
def test_italic_code_is_skewed_only_when_the_mono_face_has_no_italic(monkeypatch, body_italic, mono_italic, skewed):
    monkeypatch.setattr(docs_pdf, "has_real_italic", lambda: body_italic)
    monkeypatch.setattr(docs_pdf, "has_real_mono_italic", lambda: mono_italic)
    html = docs_pdf.build_html(parse_harvis_doc(_doc(body="*use `grep` now*")), "T")
    assert (skewed in html) is True
    if mono_italic:
        assert "<span class=\"fi\"><code>" not in html and '<span class="fi">use</span>' in html
    else:
        assert "<em><code>" not in html


@pytest.mark.skipif(not docs_fonts.has_real_mono_italic(), reason="no mono oblique face installed here")
def test_italic_code_leans_about_as_much_as_italic_text_not_twice_as_much():
    def lean(body):
        return _ink_lean(export_doc(f"---\nkind: pdf\n---\n\n{body}\n", "pdf").data)
    text = "IIIIIIIIIIIIIIII IIIIIIIIIIIIIIII"
    flat = lean(f"`{text}`")
    italic_code = lean(f"*`{text}`*") - flat
    assert 1.0 < italic_code < 1.5 * (lean(f"*{text}*") - lean(text)), italic_code


# ---- (d) the title line of a cards block ----

CARDS = "title: Three risks\nPhishing | 40% | Lures imitate payroll\nMalware | 12% | Mostly macros\nInsider | 3%"


def test_a_leading_title_line_is_the_block_title_not_a_card():
    cards = parse_cards(CARDS)
    assert cards.title == "Three risks"
    assert [c.title for c in cards] == ["Phishing", "Malware", "Insider"]
    assert caption_of("cards", cards) == "Three risks"
    assert [r[1] for r in rich_as_rows("cards", cards)] == ["Title", "Phishing", "Malware", "Insider"]


def test_title_is_read_case_insensitively_after_blank_lines():
    cards = parse_cards("\n  Title : Risks \nOne | 1")
    assert cards.title == "Risks" and [c.title for c in cards] == ["One"]


@pytest.mark.parametrize("body", ["Phishing | 40%\ntitle: late", "title: A | B | C", "subtitle: not a title\nOne | 1"])
def test_only_a_first_title_line_without_a_pipe_is_the_title(body):
    cards = parse_cards(body)
    assert cards.title == ""
    assert any(c.title.lower().startswith(("title", "subtitle")) for c in cards)


def test_a_cards_block_without_title_is_unchanged():
    cards = parse_cards("One | 1 | a\nTwo")
    assert cards.title == "" and [c.title for c in cards] == ["One", "Two"]
    assert caption_of("cards", cards) == "Cards"


def test_the_check_agrees_with_the_renderer_about_the_title_line():
    assert dropped_lines("cards", CARDS) == []
    assert dropped_lines("cards", "One | 1\ntitle: late") == []          # drawn as a card, so not skipped
    assert late_card_titles("One | 1\ntitle: late") == ["title: late"]   # but the check still points it out


def test_the_drawn_cards_carry_the_title_and_no_card_called_title():
    block = parse_blocks("```cards\n" + CARDS + "\n```")[0]
    svg = svg_markup(block.rich, block.data, 640, 880, 640)
    assert "Three risks" in svg and "title: Three risks" not in svg
    assert svg.count(">Phishing<") == 1 or "Phishing" in svg


def test_exports_do_not_draw_the_title_line_as_a_card():
    src = _doc(body="```cards\n" + CARDS + "\n```")
    for fmt in ("pdf", "docx"):
        assert export_doc(src, fmt).data
    deck = "---\nkind: deck\ntitle: D\n---\n# D\n---\n## Risks\n```cards\n" + CARDS + "\n```\n"
    assert export_doc(deck, "pptx").data[:2] == b"PK"
    from pptx import Presentation
    texts = [sh.text_frame.text for s in Presentation(io.BytesIO(export_doc(deck, "pptx").data)).slides
             for sh in s.shapes if sh.has_text_frame]
    assert not any(t.lower().startswith("title:") for t in texts)
    assert "Phishing" in texts


def test_an_empty_leading_title_line_is_consumed_and_a_later_one_is_a_card():
    body = "title:\ntitle: Late\nOne | 1"
    cards = parse_cards(body)
    assert cards.title == "" and [c.title for c in cards] == ["title: Late", "One"]
    assert dropped_lines("cards", body) == [] and late_card_titles(body) == ["title: Late"]
    assert dropped_lines("cards", "\n  title:  \nOne | 1") == []
    assert parse_cards("\n title:\nOne").title == ""


def _cards_slide(body):
    from pptx import Presentation
    deck = "---\nkind: deck\ntitle: D\n---\n# D\n---\n## Risks\n```cards\n" + body + "\n```\n"
    return Presentation(io.BytesIO(export_doc(deck, "pptx").data)).slides[1]


def test_the_pptx_draws_the_cards_title_above_the_cards():
    from pptx.enum.shapes import MSO_SHAPE
    slide = _cards_slide(CARDS)
    titles = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "Three risks"]
    cards = [s for s in slide.shapes if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
    assert len(titles) == 1 and len(cards) == 3
    assert titles[0].top + titles[0].height <= min(c.top for c in cards)


def test_the_pptx_cards_without_a_title_draw_no_extra_text():
    slide = _cards_slide("One | 1\nTwo | 2")
    assert {"One", "Two"} <= {s.text_frame.text for s in slide.shapes if s.has_text_frame}


@pytest.mark.parametrize("title", ["\u3164", "\u115f\u1160", "\uffa0", "\u034f", "\ufe0f", "\u3164 \u200b \u115f"])
def test_slugify_ignores_invisible_default_ignorable_characters(title):
    assert slugify(title) == "document"
    assert slugify(title + "Plan" + title) == "plan"


def test_slugify_does_not_split_a_word_on_an_invisible_mark():
    assert slugify("ab\u3164cd") == "abcd" and slugify("字\ufe0f字") == "字字"


def test_triple_asterisks_are_one_bold_italic_run():
    runs = inline_runs("a ***both*** b ___also___ c")
    both = [r for r in runs if r.bold and r.italic]
    assert [r.text for r in both] == ["both", "also"]
    assert "".join(r.text for r in runs) == "a both b also c"


def test_bold_italic_reaches_every_exporter():
    src = _doc(body="Plain ***loud slant*** end")
    from docx import Document
    para = next(p for p in Document(io.BytesIO(export_doc(src, "docx").data)).paragraphs if "loud" in p.text)
    run = next(r for r in para.runs if "loud slant" in r.text)
    assert run.bold and run.italic and "*" not in para.text
    deck = "---\nkind: deck\ntitle: D\n---\n# D\n---\n## S\nPlain ***loud slant*** end\n"
    from pptx import Presentation
    runs = [r for s in Presentation(io.BytesIO(export_doc(deck, "pptx").data)).slides for sh in s.shapes
            if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    hit = next(r for r in runs if "loud slant" in r.text)
    assert hit.font.bold and hit.font.italic and not any("*" in r.text for r in runs)
    html = docs_pdf.build_html(parse_harvis_doc(src), "T")
    para = html.split("<p>")[-1]
    assert "<em>" in para and "<strong>loud</strong>" in para and "<strong>slant</strong>" in para
    assert "*" not in html.split("<body")[-1]


def test_the_check_warns_that_a_late_title_line_is_drawn_as_a_card():
    from plugins.hermes_ui.docs_check import check_source
    issues, _, _ = check_source(_doc(body="```cards\nRisks\nOne | 1\ntitle: Late\n```"), None)
    said = " ".join(i["message"] for i in issues)
    assert "'title: Late' is drawn as a card" in said and "skipped" not in said


def test_triple_underscores_inside_a_word_are_not_bold_italic():
    runs = inline_runs("snake___case___x")
    assert not any(r.bold and r.italic for r in runs)
    assert [(r.text, r.bold, r.italic) for r in inline_runs("a ___b___ c")][1] == ("b", True, True)
