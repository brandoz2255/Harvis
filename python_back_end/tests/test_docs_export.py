"""A harvis-doc block becomes a real file that its own library opens back."""

import csv
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from plugins.hermes_ui import docs_export, docs_pdf  # noqa: E402
from plugins.hermes_ui.docs_blocks import inline_runs, parse_blocks, split_notes  # noqa: E402
from plugins.hermes_ui.docs_export import ExportError, export_doc, slugify  # noqa: E402
from plugins.hermes_ui.docs_formula import FormulaError, evaluate, is_safe_formula, parse_number  # noqa: E402
from plugins.hermes_ui.docs_parse import parse_harvis_doc  # noqa: E402

DECK = """---
kind: deck
title: Q3 Security Review
outline:
  - Title
  - Threat landscape
  - Next steps
---
# Q3 Security Review
Quarterly briefing
Notes: Welcome everyone.
---
## Threat landscape
- Phishing grew **40%** over Q2
  - Most lures imitated payroll notices
- Two *credential-stuffing* waves hit the [VPN](https://example.com/vpn)
Notes: Mention the payroll lure.
Second line of notes.
---
## Next steps
1. Enforce passkeys
2. Run a drill

| Owner | Due |
| --- | --- |
| Ana | Nov 15 |
| Raj | Dec 1 |
"""

DOC = """---
kind: doc
title: Network Outage Memo
outline:
  - Summary
  - Cause
---
# Network Outage Memo
Opening line before any section.
## Summary
On **Tuesday** the lab network was down for `40` minutes.

- First point
  - Nested point
- Second point

1. One
2. Two

Then a second list:

1. Again one
2. Again two
## Cause
### Detail
A firmware update restarted the *core* switch. See [the log](https://example.com/log).

| Item | Count |
| --- | ---: |
| Switches | 3 |

```python
## not a section
print("hi")
```
"""

SHEET = """---
kind: sheet
title: Household Budget
outline:
  - Monthly
  - "Q&A: [notes]"
---
## Monthly
| Category | Planned | Actual | Share |
| --- | --- | --- | --- |
| Rent | $1,200 | 1,250 | 12% |
| Food | 450 | 520.5 | 7.5% |
| **Total** | =SUM(B2:B3) | =SUM(C2:C3) | =C4/B4 |
| Phone | 007 | n/a | =WEBSERVICE("http://evil.example/x") |
## Q&A: [notes]
| Question | Answer |
| --- | --- |
| Why? | Because |
"""


# ---- parsing ----

def test_a_truncated_header_reports_what_has_arrived():
    doc = parse_harvis_doc("---\nkind: deck\ntitle: Q3 Sec")
    assert (doc.header_complete, doc.kind, doc.sections) == (False, "deck", [])
    assert doc.title == "Q3 Sec"


def test_a_half_written_outline_item_is_not_counted():
    doc = parse_harvis_doc("---\nkind: doc\noutline:\n  - Intro\n  - Thre")
    assert doc.outline == ["Intro"]
    assert parse_harvis_doc("---\nkind: doc\noutline:\n  - Intro\n  - Thre\n").outline == ["Intro", "Thre"]


def test_the_full_deck_reads_into_slides_with_a_title():
    doc = parse_harvis_doc(DECK)
    assert (doc.header_complete, doc.kind, doc.title) == (True, "deck", "Q3 Security Review")
    assert [s.title for s in doc.sections] == ["Q3 Security Review", "Threat landscape", "Next steps"]
    assert doc.outline == ["Title", "Threat landscape", "Next steps"]


def test_a_kind_line_may_carry_the_trailing_comment_the_prompt_example_has():
    assert parse_harvis_doc("---\nkind: deck   # deck | doc | sheet | pdf\n---\n# A").kind == "deck"


def test_the_outer_fence_is_tolerated_and_a_missing_header_means_a_doc():
    fenced = "````harvis-doc\n---\nkind: sheet\n---\n## A\n| x |\n| - |\n| 1 |\n````"
    assert parse_harvis_doc(fenced).kind == "sheet"
    bare = parse_harvis_doc("# Hello\n## One\ntext")
    assert (bare.kind, bare.title, [s.title for s in bare.sections]) == ("doc", "Hello", ["One"])


def test_a_hash_hash_line_inside_a_code_fence_is_code():
    doc = parse_harvis_doc(DOC)
    assert [s.title for s in doc.sections] == ["", "Summary", "Cause"]
    assert "## not a section" in doc.sections[2].body


def test_sheets_take_their_names_from_headings_and_cells_stay_raw():
    doc = parse_harvis_doc(SHEET)
    assert [s.name for s in doc.sheets] == ["Monthly", "Q&A: [notes]"]
    assert doc.sheets[0].rows[3] == ["**Total**", "=SUM(B2:B3)", "=SUM(C2:C3)", "=C4/B4"]


def test_speaker_notes_are_split_from_the_slide_body():
    body, notes = split_notes("- a\n- b\nNotes: first\ncontinues here\n\n- c")
    assert body == "- a\n- b\n\n- c" and notes == "first continues here"


def test_inline_markdown_becomes_runs_and_never_html():
    runs = inline_runs("a **b *c* d** `d` [e](https://x.test) [f](javascript:alert(1)) <b>g</b>")
    assert [r.text for r in runs if r.bold] == ["b ", "c", " d"]
    assert [r.text for r in runs if r.italic] == ["c"]
    assert [r.href for r in runs if r.href] == ["https://x.test"]
    assert any(r.code and r.text == "d" for r in runs)
    assert "<b>g</b>" in "".join(r.text for r in runs)


def test_blocks_nest_lists_by_indent_and_keep_ordered_counts():
    items = parse_blocks("- a\n    - deep\n  - mid\n1. one\n2. two")[0].items
    assert [(i.level, i.ordered) for i in items] == [(0, False), (1, False), (1, False), (0, True), (0, True)]
    assert [i.number for i in items if i.ordered] == [1, 2]


# ---- formulas and numbers ----

@pytest.mark.parametrize("raw,value,fmt", [
    ("1200", 1200, None), ("1,200", 1200, "#,##0"), ("$40", 40, "$#,##0"), ("$1,200.50", 1200.5, "$#,##0.00"),
    ("12%", 0.12, "0%"), ("12.5%", 0.125, "0.0%"), ("(300)", -300, None), ("-4.5", -4.5, None),
])
def test_numbers_are_read_with_their_format(raw, value, fmt):
    assert parse_number(raw) == (value, fmt)


@pytest.mark.parametrize("raw", ["007", "abc", "1,2", "12345678901234567", "", "5 apples"])
def test_identifiers_and_words_stay_text(raw):
    assert parse_number(raw) is None


def test_formulas_outside_the_safe_set_are_refused():
    assert is_safe_formula("=SUM(B2:B4)/COUNT(B2:B4)")
    assert is_safe_formula('=IF(A1>1,"SUM(",0)')
    assert not is_safe_formula('=WEBSERVICE("http://x.test")')
    assert not is_safe_formula("=cmd|' /C calc'!A0")
    assert not is_safe_formula("=[1]Sheet1!A1")


def test_the_evaluator_does_arithmetic_and_refuses_the_rest():
    grid = {(0, 0): 1.0, (1, 0): 2.0, (2, 0): 3.0}

    def cell(r, c):
        return grid.get((r, c))

    assert evaluate("=SUM(A1:A3)/COUNT(A1:A3)", cell) == 2.0
    assert evaluate("=10%*50+2^3-A1", cell) == 12.0
    for bad in ("=A1/0", "=HYPERLINK(1)", "=SUM(", "=A1 A2"):
        with pytest.raises(FormulaError):
            evaluate(bad, cell)


# ---- pptx ----

def test_pptx_has_one_slide_per_section_in_16_by_9_with_notes_and_a_table():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation(io.BytesIO(export_doc(DECK, "pptx").data))
    assert len(prs.slides) == 3
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    assert prs.slide_width == Inches(13.333)
    first, middle, last = prs.slides
    assert first.shapes.title.text_frame.text == "Q3 Security Review"
    assert first.notes_slide.notes_text_frame.text == "Welcome everyone."
    assert middle.shapes.title.text_frame.text == "Threat landscape"
    assert middle.notes_slide.notes_text_frame.text == "Mention the payroll lure. Second line of notes."
    body = [p for p in middle.placeholders if p.placeholder_format.idx == 1][0].text_frame
    assert [p.level for p in body.paragraphs] == [0, 1, 0]
    assert body.paragraphs[0].runs[1].font.bold is True               # **40%**
    assert any(r.hyperlink.address == "https://example.com/vpn" for r in body.paragraphs[2].runs)
    tables = [s for s in last.shapes if s.has_table]
    assert len(tables) == 1 and tables[0].table.cell(1, 0).text_frame.text == "Ana"


def test_a_slide_without_a_heading_still_exports():
    from pptx import Presentation

    src = "---\nkind: deck\ntitle: T\n---\n# T\n---\n- only bullets"
    assert len(Presentation(io.BytesIO(export_doc(src, "pptx").data)).slides) == 2


# ---- docx ----

def test_docx_keeps_headings_lists_tables_links_and_code():
    from docx import Document

    d = Document(io.BytesIO(export_doc(DOC, "docx").data))
    by_style = [(p.style.name, p.text) for p in d.paragraphs]
    assert ("Title", "Network Outage Memo") in by_style
    assert ("Heading 1", "Summary") in by_style and ("Heading 1", "Cause") in by_style
    assert ("Heading 2", "Detail") in by_style
    assert ("List Bullet 2", "Nested point") in by_style
    assert ("List Number", "Two") in by_style
    assert len(d.tables) == 1 and d.tables[0].cell(0, 0).text == "Item" and d.tables[0].cell(1, 1).text == "3"
    assert d.tables[0].cell(0, 0).paragraphs[0].runs[0].bold is True
    summary = next(p for p in d.paragraphs if p.text.startswith("On Tuesday"))
    assert [r.bold for r in summary.runs][:2] == [None, True]
    assert any(r.font.name == "Consolas" and r.text == "40" for r in summary.runs)
    code = next(p for p in d.paragraphs if "print(" in p.text)
    assert "## not a section" in code.text and code.runs[0].font.name == "Consolas"
    assert any("example.com/log" in rel.target_ref for rel in d.part.rels.values() if rel.is_external)


def test_a_second_numbered_list_restarts_at_one():
    from docx import Document
    from docx.oxml.ns import qn

    d = Document(io.BytesIO(export_doc(DOC, "docx").data))
    firsts = [p for p in d.paragraphs if p.text in ("One", "Again one")]
    ids = [p._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val")) for p in firsts]
    assert len(ids) == 2 and ids[0] != ids[1]


# ---- xlsx / csv ----

def test_xlsx_has_real_numbers_live_formulas_and_a_frozen_bold_header():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(export_doc(SHEET, "xlsx").data))
    assert wb.sheetnames == ["Monthly", "Q&A notes"]
    ws = wb["Monthly"]
    assert ws["B2"].value == 1200 and ws["B2"].data_type == "n" and ws["B2"].number_format == "$#,##0"
    assert ws["C2"].value == 1250 and ws["C2"].number_format == "#,##0"
    assert ws["D2"].value == pytest.approx(0.12) and ws["D2"].number_format == "0%"
    assert ws["D3"].value == pytest.approx(0.075) and ws["D3"].number_format == "0.0%"
    assert ws["B4"].value == "=SUM(B2:B3)" and ws["B4"].data_type == "f"
    assert ws["D4"].value == "=C4/B4" and ws["D4"].data_type == "f"
    assert ws["A4"].value == "Total" and ws["A4"].font.bold
    assert ws["B5"].value == "007" and ws["B5"].data_type == "s"
    assert ws["D5"].data_type == "s" and ws["D5"].value.startswith("=WEBSERVICE")   # not a live formula
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws.column_dimensions["A"].width >= 8


def test_sheet_names_follow_excels_rules_and_stay_unique():
    from plugins.hermes_ui.docs_xlsx import _sheet_name

    used = set()
    assert _sheet_name("Budget/2026: [draft]?", 1, used) == "Budget 2026 draft"
    assert _sheet_name("budget/2026: [draft]?", 2, used) == "budget 2026 draft (2)"
    long = _sheet_name("x" * 50, 3, used)
    assert len(long) == 31
    assert _sheet_name("'''", 4, used) == "Sheet4"


def test_csv_carries_computed_values_not_formula_text():
    raw = export_doc(SHEET, "csv").data
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0] == ["Category", "Planned", "Actual", "Share"]
    assert rows[3] == ["Total", "1650", "1770.5", repr(round(1770.5 / 1650, 10))]
    assert rows[4][3] == "#VALUE!"                                       # WEBSERVICE is not computable
    assert not any(c.startswith("=") for r in rows for c in r)


def test_csv_quotes_text_that_opens_like_a_formula():
    src = "---\nkind: sheet\n---\n## S\n| a | b |\n| --- | --- |\n| @cmd | - item |\n| 5 | -5 |"
    rows = list(csv.reader(io.StringIO(export_doc(src, "csv").data.decode("utf-8-sig"))))
    assert rows[1] == ["'@cmd", "'- item"] and rows[2] == ["5", "-5"]


# ---- pdf / md ----

def test_a_document_pdf_opens_and_carries_its_text():
    from pypdf import PdfReader

    data = export_doc(DOC, "pdf").data
    assert data.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "Network Outage Memo" in text and "Nested point" in text and "1 / 1" in text.replace("\n", " ")
    assert abs(float(reader.pages[0].mediabox.width) - 595) < 3         # A4


def test_a_deck_pdf_is_one_landscape_page_per_slide():
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(export_doc(DECK, "pdf").data))
    assert len(reader.pages) == 3
    box = reader.pages[0].mediabox
    assert abs(float(box.width) / float(box.height) - 16 / 9) < 0.01
    assert "Threat landscape" in reader.pages[1].extract_text()


def test_markup_in_a_document_is_text_in_the_pdf_html():
    src = "---\nkind: doc\ntitle: <b>x</b>\n---\n## <script>alert(1)</script>\n<img src=\"http://169.254.169.254/\">\n"
    html = docs_pdf.build_html(parse_harvis_doc(src), "<b>x</b>")
    assert "<script>" not in html and "<img" not in html and "&lt;img" in html
    assert export_doc(src, "pdf").data.startswith(b"%PDF")


def _fetch(fetcher, url):
    return fetcher.fetch(url) if hasattr(fetcher, "fetch") else fetcher(url)


def test_the_pdf_renderer_refuses_every_url_but_data():
    fetcher = docs_pdf._fetcher()
    for url in ("http://127.0.0.1:8000/secret", "file:///etc/passwd", "ftp://example.test/x"):
        with pytest.raises(ValueError):
            _fetch(fetcher, url)
    assert _fetch(fetcher, "data:text/plain;base64,aGk=") is not None


def test_weasyprint_never_calls_out_even_when_the_html_asks_it_to():
    import http.server
    import threading

    hits = []

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path)
            self.send_response(404)
            self.end_headers()

        def log_message(self, *args):
            pass

    server = http.server.HTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        url = f"http://127.0.0.1:{server.server_port}"
        html = (f'<html><head><link rel="stylesheet" href="{url}/a.css"></head>'
                f'<body><img src="{url}/b.png"><p style="background:url({url}/c.png)">hi</p></body></html>')
        data = docs_pdf.HTML(string=html, url_fetcher=docs_pdf._fetcher()).write_pdf()
    finally:
        server.shutdown()
    assert data.startswith(b"%PDF") and hits == []


def test_markdown_export_restores_title_and_sections():
    out = export_doc(DOC, "md")
    text = out.data.decode()
    assert out.media_type.startswith("text/markdown") and out.filename == "network-outage-memo.md"
    assert text.startswith("# Network Outage Memo\n\nOpening line") and "## Summary\n" in text


# ---- dispatch ----

@pytest.mark.parametrize("kind_src,fmt,ok", [
    (DECK, "pptx", True), (DECK, "pdf", True), (DECK, "docx", False), (DECK, "csv", False),
    (DOC, "docx", True), (DOC, "pdf", True), (DOC, "md", True), (DOC, "xlsx", False),
    (SHEET, "xlsx", True), (SHEET, "csv", True), (SHEET, "pdf", False), (SHEET, "md", False),
    (DOC.replace("kind: doc", "kind: pdf"), "pdf", True), (DOC.replace("kind: doc", "kind: pdf"), "docx", True),
    (DOC.replace("kind: doc", "kind: pdf"), "md", False),
])
def test_each_kind_only_exports_to_the_formats_that_fit(kind_src, fmt, ok):
    if ok:
        assert export_doc(kind_src, fmt).data
    else:
        with pytest.raises(ExportError):
            export_doc(kind_src, fmt)


@pytest.mark.parametrize("source,fmt", [
    ("", "pdf"), ("   \n", "pdf"), (DOC, "exe"), (DOC, ""), ("---\nkind: doc\ntitle: T\n", "pdf"),
    ("---\nkind: doc\n---\n", "pdf"), ("---\nkind: sheet\n---\n## S\nno table here", "xlsx"),
    ("x" * (docs_export.MAX_SOURCE_BYTES + 1), "pdf"),
])
def test_bad_requests_are_export_errors_with_a_message(source, fmt):
    with pytest.raises(ExportError) as err:
        export_doc(source, fmt)
    assert str(err.value)


def test_filenames_are_a_slug_of_the_title():
    assert slugify("Q3 Security Review!") == "q3-security-review"
    assert slugify("日本語") == "日本語"
    assert slugify("!!!") == "document"
    assert export_doc(DECK, "pptx").filename == "q3-security-review.pptx"
    assert export_doc(DECK, "pptx", "Board / Deck").filename == "board-deck.pptx"
    assert export_doc("---\nkind: doc\n---\nplain text", "md").filename == "untitled.md"


def test_the_to_functions_return_bytes():
    assert docs_export.to_pptx(DECK)[:2] == b"PK"
    assert docs_export.to_docx(DOC)[:2] == b"PK"
    assert docs_export.to_xlsx(SHEET)[:2] == b"PK"
    assert docs_export.to_pdf(DOC)[:4] == b"%PDF"
    assert docs_export.to_csv(SHEET).startswith(b"\xef\xbb\xbf")
    assert docs_export.to_md(DOC).startswith(b"# ")


def test_small_model_habits_body_tag_and_bold_formula_cells():
    from plugins.hermes_ui.docs_parse import parse_harvis_doc

    deck = parse_harvis_doc("---\nkind: deck\ntitle: T\n---\n<body>\n# First\n- a\n---\n## Second\n</body>\n")
    assert [s.title for s in deck.sections] == ["First", "Second"]
    sheet = parse_harvis_doc("---\nkind: sheet\ntitle: B\n---\n## S\n| Item | Cost |\n| --- | --- |\n"
                             "| **Total** | **=SUM(B2:B3)** |\n")
    assert sheet.sheets[0].rows[1] == ["**Total**", "**=SUM(B2:B3)**"]
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(export_doc("---\nkind: sheet\ntitle: B\n---\n## S\n| Item | Cost |\n| --- | --- |\n"
                                            "| A | 2 |\n| B | 3 |\n| **Total** | **=SUM(B2:B3)** |\n", "xlsx").data)).active
    assert ws["B4"].data_type == "f" and ws["B4"].value == "=SUM(B2:B3)" and ws["B4"].font.bold
    assert export_doc("---\nkind: sheet\ntitle: B\n---\n## S\n| Item | Cost |\n| --- | --- |\n"
                      "| A | 2 |\n| **Total** | **=SUM(B2:B2)** |\n", "csv").data.decode("utf-8-sig").splitlines()[-1] == "Total,2"


@pytest.mark.parametrize("divider", ["| - | - |", "|-|-|", "| :-: | -: |", "| --- | --- |"])
def test_a_table_divider_may_have_one_hyphen_per_cell(divider):
    from docx import Document

    src = f"---\nkind: doc\ntitle: T\n---\n## S\n| a | b |\n{divider}\n| 1 | 2 |\n"
    tables = Document(io.BytesIO(export_doc(src, "docx").data)).tables

    assert len(tables) == 1 and tables[0].cell(1, 0).text == "1"


def test_a_single_hyphen_line_is_not_a_table_divider():
    from plugins.hermes_ui.docs_parse import is_divider

    assert is_divider("| - | - |") and is_divider("|---|---|")
    assert not is_divider("- item") and not is_divider("a | b")


def test_pdf_deck_titles_render_inline_markdown_like_the_pptx():
    src = ("---\nkind: deck\ntitle: T\n---\n# Cover\nsub\n\n---\n# Results for **Q3**\nSome text\n\n---\n# **Thank** you!")
    html = docs_pdf.build_html(parse_harvis_doc(src), "T")
    assert "**" not in html
    assert "<h1>Results for <strong>Q3</strong></h1>" in html
    assert "<h1><strong>Thank</strong> you!</h1>" in html
