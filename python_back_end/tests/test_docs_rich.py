"""Rich harvis-doc blocks (timeline, chart, quadrant, cards, flow, csv), checklists, math and chips in every export."""

import io
import os
import sys
from datetime import date

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from plugins.hermes_ui import docs_pdf  # noqa: E402
from plugins.hermes_ui.docs_blocks import inline_runs, parse_blocks  # noqa: E402
from plugins.hermes_ui.docs_chips import chip_tone  # noqa: E402
from plugins.hermes_ui.docs_export import export_doc  # noqa: E402
from plugins.hermes_ui.docs_parse import parse_harvis_doc  # noqa: E402
from plugins.hermes_ui.docs_rich import (  # noqa: E402
    parse_chart, parse_date, parse_flow, parse_number, parse_quadrant, parse_rich, parse_timeline,
    parse_cards, parse_csv, rich_as_rows, timeline_layout)
from plugins.hermes_ui.docs_svg import svg_markup  # noqa: E402

TIMELINE = """title: Go live lands four weeks after planning starts
Planning | Oct 26 2026 | Nov 1 2026
Build and configure | Nov 2 2026 | Nov 15 2026
Config freeze | Nov 16 2026
*Go live | Nov 23 2026
Review | Nov 30 2026"""
CHART = "title: Open findings by host\ntype: bar\nseries: Open, Fixed\nfw01: 3, 5\nweb01: 7, 2\n"
QUADRANT = ("title: Impact vs effort\nx: Effort\ny: Impact\nQuick wins: MFA rollout, Patch web01\n"
            "Big bets: SIEM migration\nFill-ins: Rename shares\nMoney pits: Rewrite firewall rules by hand")
CARDS = "fw01 | 10.0.0.1 | Firewall · pfSense\nweb01 | 10.0.0.10 | Web server · Ubuntu 24.04\ndb01"
FLOW = "Internet -> fw01 -> web01 -> db01\nweb01 -> mon01"
CSV = "Host,IP\nfw01,10.0.0.1\nweb01,10.0.0.10"
TABLE = "| Task | Status |\n| --- | --- |\n| Network | **Done** |\n| Cutover | Blocked |\n| Notes | plain words |"


def deck(*slides: str) -> str:
    body = "\n---\n".join(slides)
    return f"---\nkind: deck\ntitle: Rich deck\n---\n# Rich deck\nA subtitle\n\nJan 5 · Ana\n---\n{body}\n"


def fence(kind: str, body: str) -> str:
    return f"```{kind}\n{body}\n```"


ALL_SLIDES = [
    "## Cards\n" + fence("cards", CARDS),
    "## Timeline\n" + fence("timeline", TIMELINE),
    "## Even timeline\n" + fence("timeline", "Discovery | week one | week two\n*Launch | week six"),
    "## Quadrant\n" + fence("quadrant", QUADRANT),
    "## Chart\n" + fence("chart", CHART),
    "## Pie\n" + fence("chart", "type: pie\nA: 3\nB: 5"),
    "## Line\n" + fence("chart", "type: line\nseries: X, Y\nJan: 1, 2\nFeb: 3, 1"),
    "## Flow\n" + fence("flow", FLOW + "\ndb01 -> web01"),
    "## CSV\n" + fence("csv", CSV),
    "## Table\n" + TABLE,
    "## Bullets\n- one\n- [x] two\n- [ ] three\n$$x^2$$",
]


# ---- parsing ----

def test_dates_read_in_the_formats_the_spec_lists():
    today = date(2026, 3, 1)
    assert parse_date("2026-10-26", today) == (date(2026, 10, 26), True)
    assert parse_date("Oct 26", today) == (date(2026, 10, 26), False)
    assert parse_date("Oct 26 2027", today) == (date(2027, 10, 26), True)
    assert parse_date("26 Oct", today) == (date(2026, 10, 26), False)
    assert parse_date("Sept 3rd, 2026", today) == (date(2026, 9, 3), True)
    for bad in ("week one", "Feb 30", "", "Octember 4", "13/45"):
        assert parse_date(bad, today) is None


def test_a_timeline_of_real_dates_is_drawn_to_scale():
    t = parse_timeline(TIMELINE)
    assert t.title.startswith("Go live") and t.scaled and len(t.items) == 5
    assert [i.is_span for i in t.items] == [True, True, False, False, False]
    assert t.items[3].key and t.items[3].label == "Go live" and not t.items[0].key
    assert t.items[0].when == "Oct 26 2026 to Nov 1 2026"
    rows, ticks = timeline_layout(t)
    starts = [r[0] for r in rows]
    assert starts == sorted(starts) and 0 <= starts[0] < 0.01 and rows[2][1] is None      # milestone has no end
    assert abs(rows[1][0] - 7 / 36) < 0.01                    # Nov 2 is 7 days into a 36-day range
    assert ticks[0][0] == 0 and len(ticks) >= 5


def test_a_timeline_with_one_unreadable_date_is_spaced_evenly_with_its_text_kept():
    t = parse_timeline("Plan | Oct 26 | Nov 1\nBuild | next sprint | later\n*Ship | Dec 1")
    assert not t.scaled and t.items[1].when == "next sprint to later"
    rows, ticks = timeline_layout(t)
    assert ticks == []
    assert [round(r[0], 3) for r in rows] == [0.017, 0.35, 0.833]            # one third of the width per item
    assert rows[0][1] is not None and rows[1][1] is not None and rows[2][1] is None


def test_malformed_timeline_lines_are_skipped_or_kept_never_raised():
    t = parse_timeline("subtitle: sub\n | Oct 1\njust a label\n*\nOk | Oct 1 |\nA | Dec 20 | Jan 5")
    assert t.subtitle == "sub"
    assert [i.label for i in t.items] == ["Ok", "A"]           # "just a label" has no date, like the panel
    assert t.scaled
    assert t.items[1].end.year == t.items[1].start.year + 1   # Dec to Jan rolls the year
    assert parse_timeline("").items == [] and parse_rich("timeline", None).items == []


def test_numbers_carry_money_percent_and_k():
    for raw, value in (("$1,200", 1200), ("45%", 45), ("1.2k", 1200), ("-3", -3), ("€5", 5), (".5", 0.5), (" 7 ", 7)):
        assert parse_number(raw) == value
    for bad in ("", "abc", "nan", "inf", "1.2.3", "2M", "1e3", "1_000", "(5)"):
        assert parse_number(bad) is None


def test_a_chart_reads_series_rows_and_survives_junk():
    c = parse_chart(CHART)
    assert (c.type, c.series, c.rows) == ("bar", ["Open", "Fixed"], [("fw01", [3, 5]), ("web01", [7, 2])])
    c = parse_chart("type: line\nA: $1,200, 3\nB: 4\nno colon here\nC: nothing\nD: 2k, 1%")
    assert c.type == "line" and c.series == ["Series 1", "Series 2"]
    assert c.rows == [("A", [1200, 3]), ("B", [4, 0]), ("D", [2000, 1])]
    assert parse_chart("type: pie\nA: 1").type == "pie" and parse_chart("type: radar\nA: 1").type == "bar"
    assert parse_chart("title: Count\nA: 1").series == ["Count"]


def test_a_quadrant_fills_missing_cells_with_empty_ones():
    q = parse_quadrant(QUADRANT)
    assert (q.title, q.x, q.y) == ("Impact vs effort", "Effort", "Impact")
    assert q.cells[0] == ("Quick wins", ["MFA rollout", "Patch web01"]) and q.cells[3][0] == "Money pits"
    short = parse_quadrant("Only one: a, b")
    assert short.cells[0] == ("Only one", ["a", "b"]) and short.cells[3] == ("", [])


def test_cards_flow_and_csv_read_their_lines():
    cards = parse_cards(CARDS + "\n\n | stray")
    assert [c.title for c in cards] == ["fw01", "web01", "db01"] and cards[1].highlight == "10.0.0.10"
    assert cards[2].highlight == "" and cards[0].detail == "Firewall · pfSense"
    f = parse_flow(FLOW)
    assert f.layers == [["Internet"], ["fw01"], ["web01"], ["db01", "mon01"]] and f.back_edges == []
    cyc = parse_flow("a -> b -> c\nc -> a\nd -> d\nx ->")
    assert cyc.back_edges == [("c", "a")] and cyc.layers[0][0] == "a" and cyc.layers[2] == ["c"]
    assert [n for layer in cyc.layers for n in layer].count("a") == 1
    assert parse_csv('a,b\n"x, y",2\n\n3').rows == [["a", "b"], ["x, y", "2"], ["3", ""]]


def test_every_rich_kind_has_a_table_fallback():
    for kind, body in (("timeline", TIMELINE), ("chart", CHART), ("quadrant", QUADRANT), ("cards", CARDS),
                       ("flow", FLOW), ("csv", CSV)):
        rows = rich_as_rows(kind, parse_rich(kind, body))
        assert len(rows) >= 2 and all(isinstance(c, str) for r in rows for c in r)


# ---- parse_blocks ----

def test_a_rich_fence_is_a_rich_block_and_a_code_fence_stays_code():
    blocks = parse_blocks(fence("timeline", TIMELINE) + "\n\n" + fence("python", "print(1)") + "\n\n" + fence("csv", CSV))
    assert [(b.kind, b.rich or b.lang) for b in blocks] == [("rich", "timeline"), ("code", "python"), ("rich", "csv")]
    assert blocks[0].data.scaled and blocks[2].data.rows[0] == ["Host", "IP"]


def test_checklists_math_and_chips():
    blocks = parse_blocks("- [ ] open\n- [x] done **bold**\n- plain\n\n$$\nE = mc^2\n$$\n\n$$ a+b $$\n\n" + TABLE)
    items = blocks[0].items
    assert [i.checked for i in items] == [False, True, None] and items[1].runs[1].bold
    assert [(b.kind, b.text) for b in blocks[1:3]] == [("math", "E = mc^2"), ("math", "a+b")]
    table = blocks[3]
    assert [t[1] for t in table.tones] == [None, "green", "red", None]
    assert table.tones[1][0] is None


def test_an_unclosed_display_math_is_just_a_paragraph():
    assert [b.kind for b in parse_blocks("$$ never closed\nmore text")] == ["para"]


def test_inline_math_is_monospace_tex_and_prices_are_not_math():
    runs = inline_runs("Area is $x^2 + y$ here")
    assert [(r.text, r.code) for r in runs] == [("Area is ", False), ("x^2 + y", True), (" here", False)]
    price = inline_runs("It costs $5 and then $10 more")
    assert len(price) == 1 and price[0].text == "It costs $5 and then $10 more" and not price[0].code


@pytest.mark.parametrize("cell,tone", [
    ("Done", "green"), ("**Complete**", "green"), ("In progress", "blue"), ("doing", "blue"),
    ("Not started", "gray"), ("To do", "gray"), ("TODO", "gray"), ("Blocked", "red"), ("At risk", "red"),
    ("Critical", "red"), ("High", "red"), ("Medium", "amber"), ("Low", "gray"),
    ("Done early", None), ("", None), ("Highly", None)])
def test_chip_tones(cell, tone):
    assert chip_tone(cell) == tone


# ---- svg ----

@pytest.mark.parametrize("kind,body", [("timeline", TIMELINE), ("chart", CHART), ("quadrant", QUADRANT),
                                        ("cards", CARDS), ("flow", FLOW + "\nmon01 -> web01"),
                                        ("chart", "type: pie\nA: 3\nB: 1"), ("chart", "type: line\nA: 1\nB: 5")])
def test_each_figure_is_an_svg_with_escaped_text(kind, body):
    svg = svg_markup(kind, parse_rich(kind, body.replace("fw01", "<b>fw01</b>")))
    assert svg.startswith("<svg") and svg.endswith("</svg>") and "<b>" not in svg


def test_a_figure_is_scaled_down_to_the_height_it_is_given():
    data = parse_rich("timeline", TIMELINE)
    full = svg_markup("timeline", data, 800)
    small = svg_markup("timeline", data, 800, max_height=100)
    assert 'height="100.0"' in small and 'width="800.0"' in full and 'width="800.0"' not in small


# ---- pptx ----

def _prs(source: str):
    from pptx import Presentation
    return Presentation(io.BytesIO(export_doc(source, "pptx").data))


def _inside(prs, slide):
    for shape in slide.shapes:
        assert shape.width >= 0 and shape.height >= 0, (shape.name, shape.width, shape.height)
        assert shape.left >= 0 and shape.top >= 0, shape.name
        assert shape.left + shape.width <= prs.slide_width, (shape.name, shape.left, shape.width)
        assert shape.top + shape.height <= prs.slide_height, (shape.name, shape.top, shape.height)


def test_every_layout_exports_a_deck_that_reopens_with_everything_on_the_slide():
    prs = _prs(deck(*ALL_SLIDES))
    assert len(prs.slides) == len(ALL_SLIDES) + 1
    for slide in prs.slides:
        _inside(prs, slide)


def test_the_title_slide_is_dark_with_the_accent_squares():
    from pptx.util import Inches

    slide = _prs(deck("## Next\n- x")).slides[0]
    assert slide.shapes.title.text_frame.text == "Rich deck"
    assert str(slide.background.fill.fore_color.rgb) == "0E1E33"
    squares = [s for s in slide.shapes if abs(s.width - Inches(1.45)) < 10 and s.left > Inches(9)]
    assert len(squares) == 4
    assert slide.shapes.title.text_frame.paragraphs[0].alignment == 1             # left
    assert any(s.has_text_frame and "Jan 5" in s.text_frame.text for s in slide.shapes)


def _texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def test_cards_become_rounded_rectangles_with_number_badges():
    from pptx.enum.shapes import MSO_SHAPE

    slide = _prs(deck(ALL_SLIDES[0])).slides[1]
    rounded = [s for s in slide.shapes if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
    assert len(rounded) == 3
    texts = _texts(slide)
    assert {"1", "2", "3", "fw01", "10.0.0.1"} <= set(texts)
    mono = [r for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs
            if r.text == "10.0.0.1"]
    assert mono and mono[0].font.name == "Consolas"


def test_a_timeline_slide_has_bars_diamonds_and_a_key_milestone_in_accent_two():
    from pptx.enum.shapes import MSO_SHAPE

    slide = _prs(deck(ALL_SLIDES[1])).slides[1]
    shapes = [s for s in slide.shapes if s.shape_type == 1]
    diamonds = [s for s in shapes if s.auto_shape_type == MSO_SHAPE.DIAMOND]
    bars = [s for s in shapes if s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
    assert len(diamonds) == 3 and len(bars) == 2
    assert [str(d.fill.fore_color.rgb) for d in diamonds].count("2F6FDE") == 1
    assert any("Oct 26 2026 to Nov 1 2026" in t for t in _texts(slide))


def test_a_quadrant_is_four_rectangles_with_axis_captions():
    slide = _prs(deck(ALL_SLIDES[3])).slides[1]
    texts = _texts(slide)
    assert "Quick wins" in texts and any("MFA rollout" in t for t in texts)
    assert any(t.startswith("↑ Impact") for t in texts) and any(t.startswith("Effort") for t in texts)


@pytest.mark.parametrize("index,expected", [(4, "COLUMN_CLUSTERED"), (5, "PIE"), (6, "LINE_MARKERS")])
def test_a_chart_slide_holds_a_native_chart(index, expected):
    slide = _prs(deck(ALL_SLIDES[index])).slides[1]
    charts = [s.chart for s in slide.shapes if s.has_chart]
    assert len(charts) == 1 and charts[0].chart_type.name == expected
    if expected == "PIE":
        assert len(list(charts[0].plots[0].series)) == 1
    else:
        assert len(list(charts[0].plots[0].series)) >= 1


def test_the_pie_uses_only_the_first_series():
    src = deck("## P\n" + fence("chart", "type: pie\nseries: A, B\nx: 1, 9\ny: 2, 8"))
    chart = [s.chart for s in _prs(src).slides[1].shapes if s.has_chart][0]
    assert list(chart.plots[0].series[0].values) == [1.0, 2.0]


def test_a_flow_slide_has_boxes_and_connected_connectors():
    slide = _prs(deck(ALL_SLIDES[7])).slides[1]
    boxes = [s for s in slide.shapes if s.shape_type == 1 and s.has_text_frame and s.text_frame.text]
    assert {s.text_frame.text for s in boxes} >= {"Internet", "fw01", "web01", "db01", "mon01"}
    connectors = [s for s in slide.shapes if s.shape_type == 9 or s.__class__.__name__ == "Connector"]
    assert len(connectors) == 4 + 3                                              # four forward edges + a three-segment return edge


def test_table_and_csv_slides_are_native_tables_with_colored_status_cells():
    prs = _prs(deck(ALL_SLIDES[8], ALL_SLIDES[9]))
    csv_slide, table_slide = prs.slides[1], prs.slides[2]
    t = [s for s in csv_slide.shapes if s.has_table][0].table
    assert t.cell(0, 0).text_frame.text == "Host" and t.cell(2, 1).text_frame.text == "10.0.0.10"
    t = [s for s in table_slide.shapes if s.has_table][0].table
    colors = {t.cell(r, 1).text_frame.text: t.cell(r, 1).text_frame.paragraphs[0].runs[0].font.color.rgb
              for r in range(1, 4)}
    assert str(colors["Done"]) == "1E8E5A" and str(colors["Blocked"]) == "C8372D"
    assert str(t.cell(3, 1).text_frame.paragraphs[0].runs[0].font.color.rgb) == "10243B"     # plain words: ink


def test_checklist_and_math_lines_become_glyph_and_monospace_text():
    slide = _prs(deck(ALL_SLIDES[10])).slides[1]
    body = [p for p in slide.placeholders if p.placeholder_format.idx == 1][0].text_frame
    assert [p.text for p in body.paragraphs] == ["one", "☑ two", "☐ three", "x^2"]
    assert body.paragraphs[3].runs[0].font.name == "Consolas"


def test_notes_stay_speaker_notes_on_a_figure_slide():
    src = deck("## Cards\n" + fence("cards", CARDS) + "\nNotes: say the ports out loud.")
    assert _prs(src).slides[1].notes_slide.notes_text_frame.text == "say the ports out loud."


def test_a_caption_line_sits_above_the_figure_and_bullets_with_a_figure_share_the_slide():
    slide = _prs(deck("## Cap\nWhat we found\n" + fence("chart", CHART))).slides[1]
    assert "What we found" in _texts(slide) and any(s.has_chart for s in slide.shapes)
    mixed = _prs(deck("## Mixed\n- one\n- two\n- three\n" + TABLE)).slides[1]
    assert any(s.has_table for s in mixed.shapes)
    _inside(_prs(deck("## Mixed\n- one\n- two\n- three\n" + TABLE)), mixed)


def test_big_inputs_still_stay_on_the_slide():
    many_cards = "\n".join(f"card {i} | 10.0.0.{i} | " + "a long description " * 8 for i in range(11))
    many_items = "\n".join(f"Item number {i} with a rather long label | Oct {i + 1} | Oct {i + 3}" for i in range(24))
    wide_flow = " -> ".join(f"step{i}" for i in range(12)) + "\n" + "\n".join(f"step0 -> side{i}" for i in range(9))
    many_rows = "| A | B | C |\n| --- | --- | --- |\n" + "\n".join(
        f"| row {i} | " + "very long cell text " * 12 + f" | Done |" for i in range(60))
    many_cats = "type: line\n" + "\n".join(f"category number {i}: {i}, {i * 2}" for i in range(40))
    title = "## " + "An extremely long slide title that keeps on going and going " * 4 + "\n"
    prs = _prs(deck(title + fence("cards", many_cards), title + fence("timeline", many_items),
                    "## Flow\n" + fence("flow", wide_flow), "## Rows\n" + many_rows,
                    "## Cats\n" + fence("chart", many_cats), title + "\n".join(f"- bullet {i} " + "x " * 60 for i in range(30))))
    for slide in prs.slides:
        _inside(prs, slide)
    table = [s for s in prs.slides[4].shapes if s.has_table][0]
    assert table.top + table.height <= prs.slide_height
    assert any("more rows" in c.text_frame.text for c in [table.table.cell(len(table.table.rows) - 1, 0)])


def test_empty_and_broken_rich_blocks_do_not_break_the_deck():
    prs = _prs(deck("## A\n" + fence("timeline", ""), "## B\n" + fence("chart", "junk"), "## C\n" + fence("flow", ""),
                    "## D\n" + fence("cards", ""), "## E\n" + fence("csv", ""), "## F\n" + fence("quadrant", "")))
    assert len(prs.slides) == 7
    for slide in prs.slides:
        _inside(prs, slide)


# ---- pdf ----

def test_a_deck_pdf_draws_the_figures_inline_and_keeps_one_page_per_slide():
    import pypdfium2

    source = deck(*ALL_SLIDES)
    html = docs_pdf.build_html(parse_harvis_doc(source), "Rich deck")
    assert html.count("<svg") == 8 and "chip-green" in html and "☑" in html and 'class="math"' in html
    assert "Config freeze" in html and "<table>" in html
    data = export_doc(source, "pdf").data
    assert len(pypdfium2.PdfDocument(data)) == len(ALL_SLIDES) + 1


def test_a_document_pdf_carries_figures_checklists_and_math():
    from pypdf import PdfReader

    src = ("---\nkind: doc\ntitle: Report\n---\n## Plan\n" + fence("timeline", TIMELINE) + "\n\n- [x] shipped\n- [ ] later\n\n"
           "$$\nE = mc^2\n$$\n\n" + TABLE + "\n\n" + fence("csv", CSV))
    data = export_doc(src, "pdf").data
    assert data.startswith(b"%PDF") and len(PdfReader(io.BytesIO(data)).pages) >= 1
    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert "Go live" in text and "E = mc^2" in text and "fw01" in text


def test_svg_in_the_pdf_html_is_escaped():
    src = "---\nkind: doc\n---\n## T\n" + fence("cards", "<script>alert(1)</script> | <img src=x> | a")
    html = docs_pdf.build_html(parse_harvis_doc(src), "T")
    assert "<script>" not in html and "<img" not in html


# ---- docx ----

def _docx(src: str):
    from docx import Document
    return Document(io.BytesIO(export_doc(src, "docx").data))


DOC = ("---\nkind: doc\ntitle: Report\n---\n## Plan\n" + fence("timeline", TIMELINE) + "\n\n- [x] shipped\n- [ ] later\n\n"
       "$$\nE = mc^2\n$$\n\n" + TABLE + "\n\n" + fence("csv", CSV) + "\n\n" + fence("chart", CHART))


def test_a_timeline_in_a_docx_is_an_inline_picture_with_alt_text():
    d = _docx(DOC)
    assert len(d.inline_shapes) == 2
    assert d.inline_shapes[0]._inline.docPr.get("descr").startswith("Go live lands")
    assert d.inline_shapes[0].width > 0 and d.inline_shapes[0].height > 0


def test_docx_checklist_math_chips_and_csv():
    d = _docx(DOC)
    paragraphs = [p.text for p in d.paragraphs]
    assert "☑ shipped" in paragraphs and "☐ later" in paragraphs
    math = next(p for p in d.paragraphs if p.text == "E = mc^2")
    assert math.runs[0].font.name == "Consolas"
    status = next(t for t in d.tables if t.cell(0, 0).text == "Task")
    done = status.cell(1, 1).paragraphs[0].runs[0]
    assert str(done.font.color.rgb) == "1E8E5A" and done.bold
    assert status.cell(3, 1).paragraphs[0].runs[0].font.color.type is None           # plain words stay uncolored
    assert any(t.cell(0, 0).text == "Host" and t.cell(2, 1).text == "10.0.0.10" for t in d.tables)


def test_a_docx_falls_back_to_a_table_when_the_picture_cannot_be_made(monkeypatch):
    def boom(*_a, **_k):
        raise RuntimeError("no renderer")

    monkeypatch.setattr(docs_pdf, "svg_to_png", boom)
    d = _docx(DOC)
    assert len(d.inline_shapes) == 0
    assert any(t.cell(0, 0).text == "Item" and t.cell(4, 0).text == "* Go live" for t in d.tables)
    assert any(t.cell(0, 1).text == "Open" for t in d.tables)                       # the chart's table


def test_svg_to_png_makes_a_real_png():
    png = docs_pdf.svg_to_png(svg_markup("cards", parse_rich("cards", CARDS), 600), scale=1)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_parsers_match_the_panel_on_dots_spaces_and_arrows():
    assert parse_date("Oct. 26")[0].month == 10
    assert parse_csv('Name, Note, N\nA, "b, c", 1').rows[1] == ["A", "b, c", "1"]
    assert [e for e in parse_flow("A => B\nB ---> C\nC ⟶ D").edges] == [("A", "B"), ("B", "C"), ("C", "D")]


def test_timeline_date_labels_beside_a_bar_or_diamond_never_wrap():
    tl = "Build | Oct 26 | Nov 1\n*Launch | Nov 24\nKickoff | Oct 28"
    slide = _prs(deck("## Plan\n" + fence("timeline", tl))).slides[1]
    labels = {"Oct 26 to Nov 1", "Nov 24", "Oct 28"}               # none of these is a gridline label
    found = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip() in labels
             and s.shape_type == 17]
    assert {s.text_frame.text.strip() for s in found} == labels
    assert all(s.text_frame.word_wrap is False for s in found)
