"""A document block to a .docx: headings, lists, tables and code, from the shared block reader.

The theme (docs_theme) sets the fonts and colours of the styles, and the file gets a page-number footer, a table of
contents field when it has four or more sections, and no thumbnail of the stock template."""
from __future__ import annotations

import io
from contextvars import ContextVar

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.table import _Cell

from .docs_blocks import Block, ListItem, Run, csv_table, inline_runs, parse_blocks, plain_text
from .docs_chips import tone_color
from .docs_parse import ParsedDoc
from .docs_rich import caption_of, rich_as_rows
from .docs_theme import DEFAULT_THEME, Theme, get_theme, restyle_theme_xml

MONO = "Consolas"
TOC_MIN_SECTIONS = 4
_theme: ContextVar[Theme] = ContextVar("docx_theme", default=DEFAULT_THEME)     # the theme of the document being built


def _style_run(run, r: Run) -> None:
    run.bold = r.bold or None
    run.italic = r.italic or None
    if r.code:
        run.font.name = MONO
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
        run.font.size = Pt(9.5)


def _write_runs(paragraph, runs: list[Run]) -> None:
    for r in runs:
        if not r.href:
            _style_run(paragraph.add_run(r.text), r)
            continue
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(r.href, RT.HYPERLINK, is_external=True))
        wr = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color, under = OxmlElement("w:color"), OxmlElement("w:u")
        color.set(qn("w:val"), _theme.get().accent2)
        under.set(qn("w:val"), "single")
        rpr.extend([color, under])
        if r.bold:
            rpr.append(OxmlElement("w:b"))
        if r.italic:
            rpr.append(OxmlElement("w:i"))
        text = OxmlElement("w:t")
        text.text = r.text
        text.set(qn("xml:space"), "preserve")
        wr.extend([rpr, text])
        link.append(wr)
        paragraph._p.append(link)


def _shade(element, fill: str, parent_tag: str) -> None:
    props = element.find(qn(parent_tag))
    if props is None:
        props = OxmlElement(parent_tag)
        element.insert(0, props)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def _restart_numbering(doc, paragraph) -> None:
    """Point the paragraph at a fresh copy of its list style's numbering, counting from 1.

    Every 'List Number' paragraph shares one numbering id, so without this a second
    numbered list in the document would carry on from the first."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_num = paragraph.style.element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_numId().val = num.numId
    num_pr.get_or_add_ilvl().val = 0


def _list_item(doc, it: ListItem) -> None:
    if it.checked is not None:                          # a checklist item: the box is the marker
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3 * (it.level + 1))
        p.paragraph_format.first_line_indent = Inches(-0.25)
        _write_runs(p, [Run("\u2611 " if it.checked else "\u2610 "), *it.runs])
        return
    suffix = "" if it.level == 0 else f" {min(it.level, 2) + 1}"
    p = doc.add_paragraph(style=("List Number" if it.ordered else "List Bullet") + suffix)
    _write_runs(p, it.runs)
    if it.ordered and it.level == 0 and it.number == 1:
        _restart_numbering(doc, p)


def _table(doc, block: Block) -> None:
    cols = max(len(r) for r in block.rows)
    table = doc.add_table(rows=len(block.rows), cols=cols)
    table.style = "Table Grid"
    trs = table._tbl.tr_lst
    for ri, row in enumerate(block.rows):
        for ci, tc in enumerate(trs[ri].tc_lst):          # table.cell(r, c) rescans the whole grid: quadratic
            cell = _Cell(tc, table)
            p = cell.paragraphs[0]
            runs = row[ci] if ci < len(row) else []
            if ri == 0:
                runs = [Run(r.text, True, r.italic, r.code, r.href) for r in runs]
                _shade(cell._tc, _theme.get().tint, "w:tcPr")
            _write_runs(p, runs)
            tone = block.tones[ri][ci] if ri < len(block.tones) and ci < len(block.tones[ri]) else None
            if tone:                                    # a status chip is colored bold text
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor.from_string(tone_color(tone))
    header = table.rows[0]._tr.get_or_add_trPr()
    header.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph()


def _code(doc, block: Block) -> None:
    p = doc.add_paragraph()
    _shade(p._p, _theme.get().card, "w:pPr")
    lines = block.text.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = MONO
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
        run.font.size = Pt(9)
        if i < len(lines) - 1:
            run.add_break()


def _math(doc, block: Block) -> None:
    p = doc.add_paragraph()
    _shade(p._p, _theme.get().card, "w:pPr")
    for i, line in enumerate(block.text.split("\n")):
        run = p.add_run(line)
        run.font.name = MONO
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
        run.font.size = Pt(10)
        if i:
            run.add_break()


_MAX_FIGURE_PX = 1000                          # about a page of picture at 6.3in wide


def _rich(doc, block: Block) -> None:
    """A figure as a picture (svg -> weasyprint pdf -> pypdfium2 png); if that cannot run, as a table."""
    if block.rich == "csv":
        table = csv_table(block.data)
        if table.rows:
            _table(doc, table)
        return
    try:
        from .docs_pdf import svg_to_png
        from .docs_svg import svg_markup, svg_size

        markup = svg_markup(block.rich, block.data, 720, max_height=_MAX_FIGURE_PX)
        picture = doc.add_picture(io.BytesIO(svg_to_png(markup)), width=Inches(6.3 * svg_size(markup)[0] / 720))
        picture._inline.docPr.set("descr", caption_of(block.rich, block.data))
        return
    except Exception:                                   # noqa: BLE001 - a missing renderer must not sink the file
        pass
    rows = rich_as_rows(block.rich, block.data)
    if rows:
        _table(doc, Block("table", rows=[[[Run(c)] for c in r] for r in rows]))


def _rule(doc) -> None:
    """A horizontal rule: an empty paragraph with a thin bottom border (it is the same line at any page width)."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", _theme.get().line)):
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    pPr.append(border)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def _blocks(doc, blocks: list[Block]) -> None:
    for b in blocks:
        if b.kind == "heading":
            # Section titles are Heading 1, so a `###` inside a section is Heading 2.
            _write_runs(doc.add_heading(level=min(max(b.level - 1, 1), 4)), b.runs)
        elif b.kind == "para":
            _write_runs(doc.add_paragraph(), b.runs)
        elif b.kind == "list":
            for it in b.items:
                _list_item(doc, it)
        elif b.kind == "table" and b.rows:
            _table(doc, b)
        elif b.kind == "code":
            _code(doc, b)
        elif b.kind == "quote":
            _write_runs(doc.add_paragraph(style="Quote"), b.runs)
        elif b.kind == "rule":
            _rule(doc)
        elif b.kind == "math":
            _math(doc, b)
        elif b.kind == "rich":
            _rich(doc, b)


def _style_font(style, face: str) -> None:
    """Name the face on the style itself: the stock styles point at the theme's fonts, and a theme name beats ``w:ascii``."""
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)
    fonts.set(qn("w:ascii"), face)
    fonts.set(qn("w:hAnsi"), face)


def _apply_theme(document, theme: Theme) -> None:
    """The theme's colour scheme and fonts in the theme part, and its fonts and colours on the styles Word shows."""
    part = document.part.part_related_by(RT.THEME)
    part._blob = restyle_theme_xml(part.blob, theme)
    styles = document.styles
    _style_font(styles["Normal"], theme.body_font)
    for name, color in (("Title", theme.ink), ("Heading 1", theme.ink), ("Heading 2", theme.accent),
                        ("Heading 3", theme.ink), ("Heading 4", theme.muted)):
        try:
            style = styles[name]
        except KeyError:
            continue
        _style_font(style, theme.heading_font)
        style.font.color.rgb = RGBColor.from_string(color)
    border = styles["Title"].element.find(f".//{qn('w:bottom')}")
    if border is not None:
        border.set(qn("w:color"), theme.accent)
        border.attrib.pop(qn("w:themeColor"), None)
    styles["Quote"].font.color.rgb = RGBColor.from_string(theme.muted)


def _field_runs(paragraph, instruction: str, shown: str = "", size: float | None = None, color: str | None = None):
    """A complex field (begin, instruction, separate, the text shown until Word updates it, end) at the end of the paragraph."""
    def run(child):
        r = paragraph.add_run()
        r.font.size = Pt(size) if size else None
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        r._r.append(child)
        return r

    def mark(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run(mark("begin"))
    run(instr)
    run(mark("separate"))
    paragraph.add_run(shown)
    if size or color:
        shown_run = paragraph.runs[-1]
        shown_run.font.size = Pt(size) if size else None
        if color:
            shown_run.font.color.rgb = RGBColor.from_string(color)
    run(mark("end"))


def _footer(document, theme: Theme) -> None:
    """``Page N of M`` centred in the footer, as fields Word and LibreOffice fill in."""
    p = document.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style = dict(size=9, color=theme.muted)
    r = p.add_run("Page ")
    r.font.size, r.font.color.rgb = Pt(9), RGBColor.from_string(theme.muted)
    _field_runs(p, "PAGE", "1", **style)
    r = p.add_run(" of ")
    r.font.size, r.font.color.rgb = Pt(9), RGBColor.from_string(theme.muted)
    _field_runs(p, "NUMPAGES", "1", **style)


def _contents(document, titles: list[str]) -> None:
    """A table-of-contents field for the section headings, with the headings listed as what it shows until it is updated."""
    document.add_paragraph("Contents", style="TOC Heading")
    for i, name in enumerate(titles):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if i == 0:
            for kind in ("begin", "instr", "separate"):
                r = p.add_run()
                if kind == "instr":
                    el = OxmlElement("w:instrText")
                    el.set(qn("xml:space"), "preserve")
                    el.text = ' TOC \\o "1-2" \\h \\z \\u '
                else:
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), kind)
                r._r.append(el)
        p.add_run(name)
        if i == len(titles) - 1:
            r = p.add_run()
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            r._r.append(end)


def build_docx(doc: ParsedDoc, title: str) -> bytes:
    theme = get_theme(getattr(doc, "theme", ""))
    token = _theme.set(theme)
    try:
        document = Document()
        document.core_properties.title = title
        document.core_properties.author = "Harvis"
        _apply_theme(document, theme)
        _footer(document, theme)
        package_rels = document.part.package.rels
        for rid in [rid for rid, rel in package_rels.items() if rel.reltype == RT.THUMBNAIL]:
            package_rels.pop(rid)                      # the stock template's picture of an empty page
        document.add_paragraph(title, style="Title")
        titles = [plain_text(inline_runs(s.title)) for s in doc.sections if s.title]
        if len(titles) >= TOC_MIN_SECTIONS:
            _contents(document, titles)
        for section in doc.sections:
            if section.title:
                _write_runs(document.add_heading(level=1), inline_runs(section.title))
            _blocks(document, parse_blocks(section.body))
        out = io.BytesIO()
        document.save(out)
        return out.getvalue()
    finally:
        _theme.reset(token)
