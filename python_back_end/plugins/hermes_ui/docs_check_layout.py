"""Layout checks for a harvis-doc that need no model and no sandbox: text that does not fit the slide, and
characters no installed font can draw.

The pptx writer fits text with a rough width guess (``docs_metrics``). That guess runs wide for ordinary
lower-case prose and narrow for capitals, wide scripts and symbols, so a slide can still overflow. This module
reads the finished deck back and measures every text box with the real font's advance widths. It is built to
stay quiet: it measures with a face at least as narrow as what a viewer will draw and allows a margin before
it speaks, so a slide the writer fitted normally never trips it. ``check_layout`` never raises.
"""
from __future__ import annotations

import io
import logging
import math
import re
import subprocess
import unicodedata
from bisect import bisect_right
from functools import lru_cache
from pathlib import Path
from typing import Callable, Optional, Sequence

from .docs_blocks import split_notes
from .docs_check import issue, page_label
from .docs_parse import ParsedDoc

log = logging.getLogger(__name__)

LEADING = 1.2                     # line height over font size: what the writer plans with
SLACK = 1.15                      # text may need this much more than the box before it is a finding
SLACK_PT = 2.0
EDGE_PT = 7.2                     # 0.1 in past the slide edge is rounding, not a finding
BULLET_INDENT_PT = 0.45 * 72      # per level, as docs_pptx._bullet_plan reserves it
BODY_SPACE_BEFORE = 0.2           # the layout's gap above a body paragraph, in font sizes
DEFAULT_PT = 18.0
SAMPLE_CHARS = 5
# Families a viewer is most likely to draw the writer's default (Calibri) and mono (Consolas) faces with.
_SANS = ("Carlito", "Calibri", "Liberation Sans", "Arimo")
_MONO = ("Consolas", "Liberation Mono", "Cousine")
# Without one of those installed, DejaVu is measured and scaled down to Arial-like widths (measured 0.80-1.0).
_DEJAVU_SANS, _DEJAVU_MONO = 0.85, 0.92
_AVG_EM = {(False, False): 0.50, (True, False): 0.54, (False, True): 0.55, (True, True): 0.55}


# ---- measuring text ---------------------------------------------------------------------------

def _fc_match(pattern: str) -> tuple[str, str]:
    out = subprocess.run(["fc-match", "-f", "%{family}\t%{file}", pattern], capture_output=True, text=True,
                         timeout=5).stdout
    family, _, path = out.partition("\t")
    return family.split(",")[0].strip(), path.strip()


@lru_cache(maxsize=4)
def _face(bold: bool, mono: bool) -> Optional[tuple[str, float]]:
    """(font file, width scale) to measure with, or None when no font file can be found."""
    style = ":bold" if bold else ""
    try:
        for family in _MONO if mono else _SANS:
            found, path = _fc_match(family + style)
            if found.lower() == family.lower() and Path(path).is_file():
                return path, 1.0
        found, path = _fc_match(("DejaVu Sans Mono" if mono else "DejaVu Sans") + style)
        if Path(path).is_file():
            return path, _DEJAVU_MONO if mono else _DEJAVU_SANS
    except (OSError, subprocess.SubprocessError):
        pass
    return None


@lru_cache(maxsize=8)
def _pil_font(path: str):
    from PIL import ImageFont

    return ImageFont.truetype(path, 100)


def _is_wide(ch: str) -> bool:
    return unicodedata.east_asian_width(ch) in ("W", "F") or ord(ch) >= 0x1F000


def text_width(text: str, size: float, *, bold: bool = False, mono: bool = False) -> float:
    """Points a line of text takes at ``size``. Wide characters (CJK, emoji) count one em each."""
    wide = sum(1 for ch in text if _is_wide(ch))
    rest = "".join(ch for ch in text if not _is_wide(ch))
    face = _face(bold, mono)
    if face:
        try:
            return _pil_font(face[0]).getlength(rest) * size / 100 * face[1] + wide * size
        except Exception:                                           # noqa: BLE001 - fall back to the average
            log.debug("font metrics unavailable for %s", face[0], exc_info=True)
    return len(rest) * size * _AVG_EM[(bold, mono)] + wide * size


def count_lines(text: str, width: float, size: float, *, bold: bool = False, mono: bool = False) -> int:
    """Lines a paragraph takes when wrapped greedily at ``width`` points; a word wider than a line is split."""
    if width <= 0:
        return 10_000
    space = text_width(" ", size, bold=bold, mono=mono)
    total = 0
    for para in re.split(r"[\n\v]", text):
        lines, cur = 1, 0.0
        for word in para.split():
            w = text_width(word, size, bold=bold, mono=mono)
            if cur and cur + space + w > width:
                lines, cur = lines + 1, 0.0
            if w > width:
                extra = math.ceil(w / width) - 1
                lines, cur = lines + extra, w - extra * width
            else:
                cur = cur + (space if cur else 0.0) + w
        total += lines
    return total


# ---- text in the deck -------------------------------------------------------------------------

def _paragraph_need(paragraph, width: float, *, indent: float = 0.0, before: float = 0.0) -> float:
    """Points of height a paragraph needs in a box ``width`` points wide (0 when it has no text)."""
    runs = [r for r in paragraph.runs if r.text]
    text = "".join(r.text for r in runs).strip()
    if not text:
        return 0.0
    sizes = [r.font.size.pt for r in runs if r.font.size]
    size = max(sizes) if sizes else DEFAULT_PT
    bold = all(r.font.bold for r in runs)
    mono = all((r.font.name or "").lower() == "consolas" for r in runs)
    lines = count_lines("".join(r.text for r in runs), width - indent, size, bold=bold, mono=mono)
    return lines * size * LEADING + before * size


def _indent(paragraph) -> float:
    marl = paragraph._p.pPr.get("marL") if paragraph._p.pPr is not None else None
    return int(marl) / 12700 if marl is not None else BULLET_INDENT_PT * (paragraph.level + 1)


def _role(shape) -> str:
    if shape.is_placeholder:
        kind = str(shape.placeholder_format.type)
        # SUBTITLE contains TITLE, so it is tested first; the writer's body is an OBJECT placeholder.
        for needle, role in (("SUBTITLE", "subtitle"), ("TITLE", "title"), ("BODY", "body"), ("OBJECT", "body")):
            if needle in kind:
                return role
    return "text"


def _text_need(shape) -> float:
    frame = shape.text_frame
    width = (shape.width - (frame.margin_left or 0) - (frame.margin_right or 0)) / 12700
    body = _role(shape) == "body"
    return sum(_paragraph_need(p, width, indent=_indent(p) if body else 0.0, before=BODY_SPACE_BEFORE if body else 0.0)
               for p in frame.paragraphs)


def _vertical_extent(shape, need: float) -> tuple[float, float]:
    """Top and bottom, in points, of the text once laid out in its box."""
    from pptx.enum.text import MSO_ANCHOR

    frame = shape.text_frame
    top, height = shape.top / 12700, shape.height / 12700
    anchor = frame.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        return top + (height - need) / 2, top + (height + need) / 2
    if anchor == MSO_ANCHOR.BOTTOM:
        return top + height - need, top + height
    return top + (frame.margin_top or 0) / 12700, top + (frame.margin_top or 0) / 12700 + need


def _measurable(shape) -> bool:
    from pptx.enum.text import MSO_AUTO_SIZE

    if not shape.has_text_frame or not shape.text_frame.text.strip():
        return False
    if None in (shape.left, shape.top, shape.width, shape.height) or getattr(shape, "rotation", 0):
        return False
    frame = shape.text_frame
    return frame.word_wrap is not False and frame.auto_size != MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


_FIX = "Shorten it, or split the slide with a line of ---"
_NAMES = {"title": "The title", "subtitle": "The subtitle", "body": "The body text"}


def _shape_findings(shape, slide_h: float) -> Optional[tuple[str, str]]:
    """(role, message) when the text in one shape overflows its box or the slide, else None."""
    need = _text_need(shape)
    if not need:
        return None
    box_pt = shape.height / 12700
    top, bottom = _vertical_extent(shape, need)
    role = _role(shape)
    who = _NAMES.get(role, "A text box")
    if bottom > slide_h + EDGE_PT or top < -EDGE_PT:
        return role, f"{who} runs off the slide: it needs about {need / 72:.1f} in of height. {_FIX}"
    if need > box_pt * SLACK + SLACK_PT:
        return role, (f"{who} is too long for its box: it needs about {need / 72:.1f} in and the box is "
                      f"{box_pt / 72:.1f} in. {_FIX}")
    return None


def _table_findings(shape, slide_h: float) -> Optional[str]:
    """A message when the rows of a table, grown to fit their text, run off the bottom of the slide."""
    table = shape.table
    widths = [c.width / 12700 for c in table.columns]
    total = 0.0
    for row in table.rows:
        need = row.height / 12700
        for c, cell in enumerate(row.cells):
            width = widths[c] - ((cell.margin_left or 0) + (cell.margin_right or 0)) / 12700
            text_h = sum(_paragraph_need(p, width) for p in cell.text_frame.paragraphs)
            need = max(need, text_h + ((cell.margin_top or 0) + (cell.margin_bottom or 0)) / 12700)
        total += need
    if shape.top / 12700 + total > slide_h + EDGE_PT:
        return f"The table is too tall for the slide once its text wraps (about {total / 72:.1f} in). Use fewer rows or shorter cells"
    return None


def _known(issues: Sequence[dict], where: str, name: str) -> bool:
    return any(i["where"] == where and f"'{name}'" in i["message"] for i in issues)


def deck_issues(data: bytes, known: Sequence[dict] = ()) -> list[dict]:
    """Overflow findings for every slide of a .pptx; ``known`` are issues already found, so a shape is not reported twice."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slide_h = prs.slide_height / 12700
    label, out = page_label("pptx"), []
    for n, slide in enumerate(prs.slides):
        where, boxes = label(n), 0
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                if shape.top is not None and shape.height is not None:
                    msg = _table_findings(shape, slide_h)
                    if msg:
                        out.append(issue("warn", where, msg))
                continue
            if not _measurable(shape) or _known(known, where, shape.name):
                continue
            found = _shape_findings(shape, slide_h)
            if not found:
                continue
            if found[0] == "text":
                boxes += 1
                if boxes == 1:
                    first = found[1]
                continue
            out.append(issue("warn", where, found[1]))
        if boxes:
            out.append(issue("warn", where, first if boxes == 1 else
                             f"{boxes} text boxes on this slide hold more text than fits; the first: {first[0].lower()}{first[1:]}"))
    return out


# ---- characters no installed font has ---------------------------------------------------------

def _merge(ranges: list[tuple[int, int]]) -> list[tuple[int, int]]:
    merged: list[tuple[int, int]] = []
    for lo, hi in sorted(ranges):
        if merged and lo <= merged[-1][1] + 1:
            merged[-1] = (merged[-1][0], max(merged[-1][1], hi))
        else:
            merged.append((lo, hi))
    return merged


def _fontconfig_ranges() -> list[tuple[int, int]]:
    out = subprocess.run(["fc-list", "--format", "%{charset}\n"], capture_output=True, text=True, timeout=10).stdout
    ranges = []
    for tok in out.split():
        lo, _, hi = tok.partition("-")
        ranges.append((int(lo, 16), int(hi or lo, 16)))
    return ranges


def _fonttools_ranges() -> list[tuple[int, int]]:
    from fontTools.ttLib import TTFont

    ranges = []
    for path in sorted(Path("/usr/share/fonts").rglob("*.[to]tf"))[:300]:
        try:
            ranges += [(cp, cp) for cp in TTFont(str(path), lazy=True).getBestCmap()]
        except Exception:                                           # noqa: BLE001 - one broken font is not the check
            continue
    return ranges


@lru_cache(maxsize=1)
def _coverage() -> Optional[Callable[[int], bool]]:
    """A predicate 'does any installed font have this code point', or None when the fonts cannot be listed."""
    ranges: list[tuple[int, int]] = []
    for source in (_fontconfig_ranges, _fonttools_ranges):
        try:
            ranges = source()
        except Exception:                                           # noqa: BLE001 - try the next way of asking
            continue
        if ranges:
            break
    if not ranges:
        return None
    merged = _merge(ranges)
    starts = [lo for lo, _ in merged]

    def covered(cp: int) -> bool:
        i = bisect_right(starts, cp) - 1
        return i >= 0 and cp <= merged[i][1]

    return covered


def _undrawn(ch: str) -> bool:
    """Characters that are not drawn on purpose (controls, joiners, variation selectors) need no glyph."""
    return unicodedata.category(ch) in ("Cc", "Cf", "Cs", "Co", "Cn", "Zl", "Zp") or 0xFE00 <= ord(ch) <= 0xFE0F


def _sections_text(doc: ParsedDoc) -> list[tuple[str, str]]:
    parts = [("the title", doc.title)]
    for n, s in enumerate(doc.sections, 1):
        body = split_notes(s.body)[0] if doc.kind == "deck" else s.body
        parts.append((f"slide {n}" if doc.kind == "deck" else f"'{s.title}'" if s.title else "the document",
                      f"{s.title}\n{body}"))
    for sheet in doc.sheets:
        parts.append((f"sheet {sheet.name}", "\n".join([sheet.name, *(c for row in sheet.rows for c in row)])))
    return parts


def glyph_issues(doc: ParsedDoc) -> list[dict]:
    """One warning naming a few characters the installed fonts cannot draw; empty when fonts cannot be listed."""
    covered = _coverage()
    if covered is None:
        return []
    missing: dict[str, str] = {}
    for where, text in _sections_text(doc):
        for ch in text:
            if ch not in missing and not _undrawn(ch) and not covered(ord(ch)):
                missing[ch] = where
    if not missing:
        return []
    shown = list(missing)[:SAMPLE_CHARS]
    more = len(missing) - len(shown)
    chars = " ".join(f"{ch} (U+{ord(ch):04X})" for ch in shown) + (f" and {more} more" if more else "")
    return [issue("warn", "document", f"No installed font can draw {chars}; first seen in {missing[shown[0]]}. "
                  "They show as empty boxes in the PDF and the preview. Use plain words or a symbol the fonts have")]


# ---- the hook ---------------------------------------------------------------------------------

def check_layout(doc: ParsedDoc, pptx: Optional[bytes], known: Sequence[dict] = ()) -> list[dict]:
    """Layout and glyph findings; any failure inside is logged and skipped, so Check always completes."""
    out: list[dict] = []
    if doc.kind == "deck" and pptx:
        try:
            out += deck_issues(pptx, known)
        except Exception:                                           # noqa: BLE001
            log.warning("deck layout check failed", exc_info=True)
    try:
        out += glyph_issues(doc)
    except Exception:                                               # noqa: BLE001
        log.warning("glyph coverage check failed", exc_info=True)
    return out
