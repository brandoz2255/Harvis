"""Native-shape timeline and flow slide bodies (bars, diamonds, boxes and connectors)."""
from __future__ import annotations

from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from .docs_chips import ACCENT, ACCENT2, CARD, INK, LINE, MUTED
from .docs_metrics import side_label, text_width
from .docs_blocks import Run
from .docs_pptx_kit import CENTER, MIDDLE, RIGHT, box, flat, muted, rgb, text
from .docs_pptx_layouts import Area, draw_table
from .docs_rich import Flow, Timeline, rich_as_rows, timeline_layout

MAX_ROWS = 14
MIN_ROW_H = 0.24                                    # the smallest row that still holds a label
_BACK = "C8372D"


def draw_timeline(slide, t: Timeline, area: Area) -> None:
    left, top, width, height = area
    y = top
    if t.title:
        text(slide, left, y, width, 0.4, t.title, 18, bold=True, min_size=12)
        y += 0.42
    if t.subtitle:
        muted(slide, left, y, width, 0.3, t.subtitle, 13, min_size=9)
        y += 0.32
    head = 0.34 if t.scaled else 0.05                # week ticks only exist on a to-scale timeline
    room = top + height - y - head
    count = min(len(t.items), MAX_ROWS, int(room / MIN_ROW_H))
    if count < len(t.items):                         # leave room for the "… N more" line
        count = max(min(MAX_ROWS, int((room - 0.3) / MIN_ROW_H)), 0)
    items = t.items[:count]
    if not items:
        muted(slide, left, y, width, 0.3, f"… {len(t.items)} items not shown", 11, min_size=8)
        return
    shown = Timeline(t.title, t.subtitle, items, t.scaled)
    rows, ticks = timeline_layout(shown)
    label_w = min(width * 0.28, max(text_width(i.label, 13, bold=True) for i in items) / 72 + 0.3)
    x0, x1 = left + label_w + 0.1, left + width
    bx0, bx1 = x0 + 0.2, x1 - 0.2
    row_h = min(0.62, (top + height - y - head - (0.3 if count < len(t.items) else 0)) / len(items))
    size = 13 if row_h >= 0.4 else 10
    grid_top, grid_bottom = y + head, y + head + row_h * len(items)
    for i in range(0, len(items), 2):
        box(slide, left, grid_top + i * row_h, width, row_h, fill=CARD)
    for at, label in ticks:
        x = bx0 + at * (bx1 - bx0)
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _emu(x), _emu(y + head - 0.05), _emu(x),
                                        _emu(grid_bottom))
        ln.line.color.rgb = rgb(LINE)
        ln.line.width = Pt(0.75)
        flat(ln)
        wlab = 0.9
        text(slide, min(x, left + width - wlab), y, wlab, 0.3, label, 11, color=MUTED, min_size=8)
    for i, (item, (a, b)) in enumerate(zip(items, rows)):
        mid = grid_top + (i + 0.5) * row_h
        color = ACCENT2 if item.key else ACCENT
        text(slide, left + 0.05, mid - row_h / 2, label_w, row_h, item.label, size, bold=item.key,
             anchor=MIDDLE, min_size=8)
        xa = bx0 + a * (bx1 - bx0)
        bar_h = min(row_h * 0.52, 0.3)
        if b is None:
            d = min(row_h * 0.6, 0.3)
            box(slide, xa - d / 2, mid - d / 2, d, d, fill=color, shape=MSO_SHAPE.DIAMOND)
            lo, hi = xa - d / 2, xa + d / 2
        else:
            xb = max(bx0 + b * (bx1 - bx0), xa + 0.08)
            box(slide, xa, mid - bar_h / 2, xb - xa, bar_h, fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
            lo, hi = xa, xb
        tw = text_width(item.when, 11, bold=item.key) / 72 + 0.12
        x, side = side_label(lo, hi, tw, x0, left + width, 0.06)
        if side == "inside":
            text(slide, lo, mid - row_h / 2, hi - lo, row_h, item.when, 11, color="FFFFFF", bold=item.key,
                 anchor=MIDDLE, min_size=8)
        elif side == "start":
            text(slide, x, mid - row_h / 2, tw, row_h, item.when, 11, color=color if item.key else MUTED,
                 bold=item.key, anchor=MIDDLE, min_size=8)
        else:
            text(slide, x - tw, mid - row_h / 2, tw, row_h, item.when, 11, color=color if item.key else MUTED,
                 bold=item.key, anchor=MIDDLE, align=RIGHT, min_size=8)
    if len(t.items) > len(items):
        muted(slide, left, grid_bottom + 0.02, width, 0.28, f"… {len(t.items) - len(items)} more", 11, min_size=8)


def _emu(inches: float) -> int:
    return int(Inches(inches))


def _arrow(connector) -> None:
    ln = connector.line._get_or_add_ln()
    ln.append(ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                             {"type": "triangle", "w": "med", "len": "med"}))


def _segment(slide, x1, y1, x2, y2, color: str, dash: bool = False, arrow: bool = False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _emu(x1), _emu(y1), _emu(x2), _emu(y2))
    flat(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.5)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        _arrow(c)
    return c


def draw_flow(slide, flow: Flow, area: Area) -> None:
    left, top, width, height = area
    layers = flow.layers
    if not layers:
        return
    lane_room = 0.4 if flow.back_edges else 0.0            # return edges run in a lane above or below the boxes
    top, height = top + lane_room, height - 2 * lane_room
    n = len(layers)
    gap_x = 0.7 if n <= 4 else 0.45 if n <= 6 else 0.3
    box_w = min(2.3, (width - gap_x * (n - 1)) / n)
    tallest = max(len(layer) for layer in layers)
    gap_y = 0.3
    box_h = min(0.85, (height - gap_y * (tallest - 1)) / tallest)
    if box_w < 0.9 or box_h < 0.3:                          # too many boxes to read: the same edges as a table
        rows = [[[Run(c)] for c in r] for r in rich_as_rows("flow", flow)]
        draw_table(slide, rows, [], (left, top - lane_room, width, height + 2 * lane_room))
        return
    used_w = n * box_w + (n - 1) * gap_x
    x_start = left + (width - used_w) / 2
    shapes, where = {}, {}
    for li, layer in enumerate(layers):
        col_h = len(layer) * box_h + (len(layer) - 1) * gap_y
        for ni, label in enumerate(layer):
            x, y = x_start + li * (box_w + gap_x), top + (height - col_h) / 2 + ni * (box_h + gap_y)
            source = li == 0
            s = box(slide, x, y, box_w, box_h, fill=ACCENT if source else "FFFFFF", line=ACCENT,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, line_w=1.5)
            text(slide, 0, 0, box_w, box_h, label, 16, color="FFFFFF" if source else INK, bold=True,
                 align=CENTER, anchor=MIDDLE, shape=s, min_size=9)
            shapes[label], where[label] = s, (x, y, ni == len(layer) - 1)
    for a, b in flow.edges:
        if (a, b) in flow.back_edges:
            continue
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
        flat(c)
        c.begin_connect(shapes[a], 3)
        c.end_connect(shapes[b], 1)
        c.line.color.rgb = rgb(MUTED)
        c.line.width = Pt(1.5)
        _arrow(c)
    for a, b in flow.back_edges:
        (ax, ay, last), (bx, by, _l) = where[a], where[b]
        below = last
        lane = top + height + lane_room / 2 if below else top - lane_room / 2
        x1, x2 = ax + box_w / 2, bx + box_w / 2
        y1, y2 = (ay + box_h, by + box_h) if below else (ay, by)
        _segment(slide, x1, y1, x1, lane, _BACK, True)
        _segment(slide, x1, lane, x2, lane, _BACK, True)
        _segment(slide, x2, lane, x2, y2, _BACK, True, True)
