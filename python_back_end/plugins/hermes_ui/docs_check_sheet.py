"""Code checks on an exported .xlsx: formulas that point at the wrong cells, and a stated limit nothing checks.

Reads the workbook the exporter wrote with openpyxl and evaluates formulas with the same
small evaluator the CSV export uses, so no spreadsheet program is needed. Where the
evaluator does not know a function the formula is left alone rather than guessed at, unless the workbook
went through the sandbox's recalculation: then it carries every formula's value, those values stand in for
the evaluator's, and any cell that came out as an error (#DIV/0!, #REF!, Err:...) is reported with its address
and formula.
"""
from __future__ import annotations

import io
import re
from typing import Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .docs_formula import SUPPORTED, FormulaError, Unknowable, evaluate_any

# Functions whose arguments are meant to be numbers: a range of text there is a mistake. IF, SUMIF, lookups and
# text functions legitimately read text, so formulas using them are only checked for reaching outside the table.
_NUMERIC = frozenset("SUM AVERAGE MIN MAX PRODUCT MEDIAN SUMPRODUCT COUNT ROUND ROUNDUP ROUNDDOWN ABS INT "
                     "MOD POWER SQRT".split())
_FUNC = re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
_REF = re.compile(r"(?<![A-Za-z0-9_.!'$])\$?([A-Za-z]{1,3})\$?(\d+)(?::\$?([A-Za-z]{1,3})\$?(\d+))?(?![\w(!])")
_COMPARES = re.compile(r"[-<>]|\bIF\b", re.I)
_AMOUNT = re.compile(r"\$\s?(\d[\d,]*(?:\.\d+)?)\s*([kKmM])?|(?:budget|limit|cap)\s*(?:of|at|under|:|-)?\s*"
                     r"\$?\s?(\d[\d,]*(?:\.\d+)?)\s*([kKmM])?", re.I)
_WORDS = re.compile(r"\b(budget|limit|cap|spend|spending|under|max|maximum|ceiling)\b", re.I)


def _col(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def _issue(severity: str, where: str, message: str) -> dict:
    return {"severity": severity, "where": where, "message": message}


_UNKNOWN = object()


class _Grid:
    """One worksheet read into numbers and text, formulas evaluated lazily with cycle detection."""

    def __init__(self, ws, cached=None) -> None:
        self.ws, self.cached = ws, cached        # cached: the same sheet read with data_only=True, or None
        self.rows, self.cols = ws.max_row, ws.max_column
        self._cache: dict[tuple[int, int], object] = {}
        self._busy: set[tuple[int, int]] = set()

    def raw(self, r: int, c: int):
        if r < 0 or c < 0 or r >= self.rows or c >= self.cols:
            return None
        return self.ws.cell(row=r + 1, column=c + 1).value

    def saved(self, r: int, c: int):
        """The value the workbook has stored for a formula cell (None when it was never calculated)."""
        if self.cached is None:
            return None
        return self.cached.cell(row=r + 1, column=c + 1).value

    def is_formula(self, r: int, c: int) -> bool:
        cell = self.ws.cell(row=r + 1, column=c + 1)
        return isinstance(cell.value, str) and cell.value.startswith("=") and cell.data_type == "f"

    def cell(self, r: int, c: int):
        """What a formula reads from the cell: a number, text, a bool, or None when it is empty. Raises
        FormulaError on a cycle, a bad formula, or a value we cannot know."""
        key = (r, c)
        if key in self._cache:
            if self._cache[key] is _UNKNOWN:
                raise Unknowable("unknown")
            return self._cache[key]
        if key in self._busy:
            raise Unknowable("circular")
        raw = self.raw(r, c)
        self._busy.add(key)
        try:
            if self.is_formula(r, c):
                if _evaluable(raw):
                    result = evaluate_any(raw, self.cell)
                else:
                    result = self.saved(r, c)
                    if result is None or (isinstance(result, str) and _error_name(result)):
                        # A function the evaluator lacks and no stored value: the cell has a value we cannot
                        # know, and so does every formula that reads it. Raising keeps that from being
                        # mistaken for text or zero.
                        self._cache[key] = _UNKNOWN
                        raise Unknowable("unknown")
            elif isinstance(raw, (str, bool)):
                result = raw
            else:
                result = float(raw) if isinstance(raw, (int, float)) else None
        finally:
            self._busy.discard(key)
        self._cache[key] = result
        return result

    def value(self, r: int, c: int) -> Optional[float]:
        """The cell as a number, or None when it is empty or text. Raises FormulaError on a cycle or a bad formula."""
        v = self.cell(r, c)
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

    def maybe_number(self, r: int, c: int) -> bool:
        """True for a number, and for a formula, whose result we may not be able to work out."""
        return self.is_formula(r, c) or self.safe(r, c) is not None

    def safe(self, r: int, c: int) -> Optional[float]:
        try:
            return self.value(r, c)
        except (FormulaError, ZeroDivisionError, TypeError, OverflowError):
            return None


_ERRORS = {
    "#DIV/0!": "divides by zero",
    "#REF!": "points at a cell that does not exist",
    "#NAME?": "uses a name or function the spreadsheet does not know",
    "#VALUE!": "does arithmetic on text",
    "#N/A": "found no match",
    "#NUM!": "does not give a valid number",
    "#NULL!": "uses a range that does not exist",
}
_ERROR_SHAPE = re.compile(r"^#[A-Z0-9/]+[!?]$|^#N/A$")


def _error_name(value: str) -> Optional[str]:
    """The error a stored cell value stands for (``#DIV/0!``, ``Err:502``), or None for ordinary text."""
    if _ERROR_SHAPE.match(value) or value.startswith("Err:"):
        return value
    return None


def _stored_error(grid: _Grid, r: int, c: int, where: str) -> Optional[dict]:
    """The error the recalculated workbook stored in a formula cell, with the cell and its formula."""
    value = grid.saved(r, c)
    error = _error_name(value) if isinstance(value, str) else None
    if not error:
        return None
    name = f"{get_column_letter(c + 1)}{r + 1}"
    return _issue("error", where, f"{name} {grid.raw(r, c)} shows {error}: it {_ERRORS.get(error, 'could not be calculated')}")


def _functions(formula: str) -> set[str]:
    return {m.upper() for m in _FUNC.findall(re.sub(r'"[^"]*"', '""', formula))}


def _evaluable(formula: str) -> bool:
    return _functions(formula) <= SUPPORTED


def _formula_issues(grid: _Grid, r: int, c: int, where: str) -> list[dict]:
    formula = grid.raw(r, c)
    name = f"{get_column_letter(c + 1)}{r + 1}"
    code = re.sub(r'"[^"]*"', '""', formula)
    out: list[dict] = []
    numeric_context = _functions(formula) <= _NUMERIC
    for c1, r1, c2, r2 in ((m.group(1), m.group(2), m.group(3), m.group(4)) for m in _REF.finditer(code)):
        top, left = int(r1) - 1, _col(c1)
        bottom, right = (int(r2) - 1, _col(c2)) if c2 else (top, left)
        top, bottom, left, right = min(top, bottom), max(top, bottom), min(left, right), max(left, right)
        shown = f"{c1.upper()}{r1}" + (f":{c2.upper()}{r2}" if c2 else "")
        if top >= grid.rows or left >= grid.cols:
            out.append(_issue("error", where, f"{name} {formula} points at {shown}, which is outside the table "
                                              f"(it ends at {get_column_letter(grid.cols)}{grid.rows})"))
            continue
        if (bottom >= grid.rows or right >= grid.cols) and c2:
            out.append(_issue("warn", where, f"{name} {formula} reaches {shown}, past the end of the table "
                                             f"({get_column_letter(grid.cols)}{grid.rows})"))
        if not numeric_context:
            continue
        cells = [(rr, cc) for rr in range(top, min(bottom, grid.rows - 1) + 1)
                 for cc in range(left, min(right, grid.cols - 1) + 1) if (rr, cc) != (r, c)]
        if c2 and cells and not any(grid.maybe_number(rr, cc) for rr, cc in cells):
            out.append(_issue("error", where, f"{name} {formula} points at {shown}, which holds no numbers "
                                              "(wrong column or row?)"))
        elif not c2 and cells:
            cell = grid.raw(top, left)
            if cell is None:
                out.append(_issue("warn", where, f"{name} {formula} uses {shown}, which is empty"))
            elif not grid.maybe_number(top, left):
                out.append(_issue("error", where, f"{name} {formula} uses {shown}, which holds text, not a number"))
    return out


def _evaluation_issue(grid: _Grid, r: int, c: int, where: str) -> Optional[dict]:
    name = f"{get_column_letter(c + 1)}{r + 1}"
    try:
        grid.value(r, c)
    except ZeroDivisionError:
        return _issue("error", where, f"{name} {grid.raw(r, c)} divides by zero (#DIV/0!)")
    except FormulaError as exc:
        if str(exc) == "div0":
            return _issue("error", where, f"{name} {grid.raw(r, c)} divides by zero (#DIV/0!)")
        if str(exc) == "circular":
            return _issue("error", where, f"{name} {grid.raw(r, c)} depends on its own result (circular reference)")
    return None


def limits_in(texts: list[str]) -> list[float]:
    """Amounts the title or outline names as a budget or limit ($2000, $2k, "budget 2000")."""
    found: list[float] = []
    for text in texts:
        if not _WORDS.search(text or ""):
            continue
        for m in _AMOUNT.finditer(text):
            num, unit = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
            value = float(num.replace(",", "")) * {"k": 1e3, "m": 1e6}.get((unit or "").lower(), 1)
            if value > 0 and value not in found:
                found.append(value)
    return found


def _shows_gap(grids: list[_Grid], limit: float) -> bool:
    """Some formula compares or subtracts the limit, written in or pointing at a cell that holds it."""
    figure = re.compile(rf"(?<![\d.]){int(limit)}(?![\d.])")
    for g in grids:
        holders = {(r, c) for r in range(g.rows) for c in range(g.cols)
                   if isinstance(g.raw(r, c), (int, float)) and g.raw(r, c) == limit}
        for r in range(g.rows):
            for c in range(g.cols):
                if not g.is_formula(r, c):
                    continue
                formula = g.raw(r, c)
                if not _COMPARES.search(formula):
                    continue
                if figure.search(formula.replace(",", "")):
                    return True
                for m in _REF.finditer(formula):
                    if (int(m.group(2)) - 1, _col(m.group(1))) in holders:
                        return True
    return False


def _totals(grids: list[_Grid]) -> list[float]:
    out: list[float] = []
    for g in grids:
        for r in range(1, g.rows):
            label = g.raw(r, 0)
            if isinstance(label, str) and "total" in label.lower():
                out += [v for c in range(1, g.cols) if (v := g.safe(r, c)) is not None]
    return out


def _limit_issues(grids: list[_Grid], limits: list[float]) -> list[dict]:
    out = []
    totals = _totals(grids)
    for limit in limits:
        if _shows_gap(grids, limit):
            continue
        amount = f"{limit:,.0f}" if limit == int(limit) else f"{limit:,.2f}"
        if totals and max(totals) > limit:
            out.append(_issue("error", "document", f"The total ({max(totals):,.2f}) is over the {amount} limit "
                                                   "in the title, and no formula shows the difference"))
        elif not totals:
            out.append(_issue("warn", "document", f"The title mentions a limit of {amount}, but no formula "
                                                  "checks the total against it"))
    return out


def check_xlsx(data: bytes, title: str = "", outline: Optional[list[str]] = None) -> list[dict]:
    """Issues in the workbook's formulas, plus a limit named in the title or outline that nothing checks."""
    try:
        wb = load_workbook(io.BytesIO(data))
    except Exception as exc:                                            # noqa: BLE001 - any unreadable file is a finding
        return [_issue("error", "document", f"The spreadsheet could not be opened ({type(exc).__name__})")]
    try:
        stored = load_workbook(io.BytesIO(data), data_only=True)      # the values a recalculation saved, if any
    except Exception:                                                   # noqa: BLE001 - only the formulas are then checked
        stored = None
    issues: list[dict] = []
    grids: list[_Grid] = []
    for ws in wb.worksheets:
        where = f"sheet {ws.title}"
        grid = _Grid(ws, stored[ws.title] if stored is not None and ws.title in stored.sheetnames else None)
        grids.append(grid)
        for r in range(1, grid.rows):
            for c in range(grid.cols):
                raw = grid.raw(r, c)
                if isinstance(raw, str) and raw.startswith("=") and not grid.is_formula(r, c):
                    issues.append(_issue("warn", where, f"{get_column_letter(c + 1)}{r + 1} {raw} is shown as text "
                                                        "because it uses a function that is not allowed"))
                elif grid.is_formula(r, c):
                    found = _evaluation_issue(grid, r, c, where)
                    issues += [found] if found else _formula_issues(grid, r, c, where)
                    stuck = None if found else _stored_error(grid, r, c, where)
                    issues += [stuck] if stuck else []
    return issues + _limit_issues(grids, limits_in([title or "", *(outline or [])]))
