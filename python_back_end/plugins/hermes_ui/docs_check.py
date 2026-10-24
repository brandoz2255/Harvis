"""Code checks on a harvis-doc and the files it exports to. Pure Python: no sandbox, no model.

``check_source`` exports every format the kind allows (a failure is an issue), then
reads the primary file back: a deck's shapes and text boxes, a sheet's formulas, the
rich blocks of a document. ``check_pages`` looks at the text of the rendered pages, which
only the sandbox can produce, so the route calls it when a render is available.
Every issue is ``{severity: 'error'|'warn', where, message}``.
"""
from __future__ import annotations

import io
import logging
import re
from typing import Optional

from .docs_blocks import parse_blocks, plain_text, split_notes
from .docs_check_rich import dropped_lines, is_empty, late_card_titles, scan_rich
from .docs_check_sheet import check_xlsx
from .docs_export import FORMATS, ExportError, export_doc
from .docs_metrics import wrap
from .docs_parse import ParsedDoc, parse_harvis_doc
from .docs_rich import RICH_KINDS

log = logging.getLogger(__name__)

MAX_BULLETS = 8
EDGE_TOLERANCE_EMU = 91_440                      # 0.1 in: rounding in the writer, not a shape off the slide
SHEET_ERRORS = ("#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
_LO_ERROR = re.compile(r"\bErr:\d{3}\b")                      # what LibreOffice draws for a formula it cannot compute
_FIGURES = ("timeline", "chart", "quadrant", "cards", "flow")
_NORMAL_TEXT_PT = 18.0


def issue(severity: str, where: str, message: str) -> dict:
    return {"severity": severity, "where": where, "message": message}


def page_label(fmt: str):
    """Names page ``i`` (0-based) the way a person reads it: 'slide 3' in a deck, 'page 3' elsewhere."""
    word = "slide" if fmt == "pptx" else "page"
    return lambda i: f"{word} {i + 1}"


# ---- exports ----------------------------------------------------------------------------------

def export_all(source: str, title: Optional[str]) -> tuple[list[dict], dict[str, bytes], ParsedDoc]:
    """(issues, {format: bytes} for each format that exported, the parsed document)."""
    doc = parse_harvis_doc(source)
    files: dict[str, bytes] = {}
    failures: dict[str, list[str]] = {}
    for fmt in FORMATS[doc.kind]:
        try:
            files[fmt] = export_doc(source, fmt, title).data
        except ExportError as exc:
            failures.setdefault(str(exc), []).append(fmt)
        except Exception as exc:                                    # noqa: BLE001 - a crash in a writer is the finding
            failures.setdefault(f"the writer failed ({type(exc).__name__})", []).append(fmt)
    issues = [issue("error", "document", f"Cannot export as {' or '.join(fmts)}: {why}")
              for why, fmts in failures.items()]
    return issues, files, doc


# ---- rich blocks ------------------------------------------------------------------------------

def _sections(doc: ParsedDoc):
    for n, s in enumerate(doc.sections, 1):
        yield (f"slide {n}" if doc.kind == "deck" else "document"), s


def check_rich(doc: ParsedDoc) -> list[dict]:
    """Blocks left unclosed, empty, or with lines the reader dropped; and, outside decks, blocks that cannot be drawn."""
    out: list[dict] = []
    for where, section in _sections(doc):
        blocks, problems = scan_rich(section.body)
        out += [issue("error", where, p.capitalize()) for p in problems]
        for kind, body in blocks:
            if is_empty(kind, body):
                out.append(issue("error", where, f"The {kind} block is empty or unreadable, so nothing is drawn"))
                continue
            lost = dropped_lines(kind, body)
            if lost:
                shown = "; ".join(repr(ln[:40]) for ln in lost[:3])
                out.append(issue("warn", where, f"The {kind} block has {len(lost)} line(s) the reader skipped: {shown}"))
            if kind == "cards":
                for ln in late_card_titles(body)[:3]:
                    out.append(issue("warn", where, f"{ln[:40]!r} is drawn as a card: only the first line of a cards "
                                                    "block can be its title"))
            if doc.kind in ("doc", "pdf") and kind in _FIGURES:
                out += _undrawable(where, section.title, kind, body)
    return out


def _undrawable(where: str, title: str, kind: str, body: str) -> list[dict]:
    from .docs_pdf import DOC_FIT
    from .docs_rich import parse_rich
    from .docs_svg import svg_markup

    try:
        svg_markup(kind, parse_rich(kind, body), *DOC_FIT)
        return []
    except Exception:                                               # noqa: BLE001 - the exporter falls back to a table
        return [issue("warn", where, f"The {kind} block in '{title}' could not be drawn and appears as a table")]


# ---- decks ------------------------------------------------------------------------------------

def _points(emu: int) -> float:
    return emu / 12700


def _is_body(shape) -> bool:
    return bool(shape.is_placeholder and "BODY" in str(shape.placeholder_format.type))


def _overflows(shape) -> bool:
    """True when the text in a shape needs clearly more height than the shape has. The estimate runs wide."""
    from pptx.enum.text import MSO_AUTO_SIZE

    frame = shape.text_frame
    if frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE or shape.width is None or shape.height is None:
        return False
    margin_x = _points((frame.margin_left or 0) + (frame.margin_right or 0))
    margin_y = _points((frame.margin_top or 0) + (frame.margin_bottom or 0))
    body = _is_body(shape)
    need = 0.0
    for p in frame.paragraphs:
        text = "".join(r.text for r in p.runs).strip()
        if not text:
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        size = max(sizes) if sizes else _NORMAL_TEXT_PT
        width = _points(shape.width) - margin_x - (0.45 * 72 * (p.level + 1) if body else 0)
        need += len(wrap(text, max(width, size), size, bold=any(r.font.bold for r in p.runs))) * size * 1.2
    return need > (_points(shape.height) - margin_y) * 1.15 + 2


def check_pptx(data: bytes, doc: ParsedDoc) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:                                        # noqa: BLE001
        return [issue("error", "document", f"The presentation could not be opened ({type(exc).__name__})")]
    out: list[dict] = []
    width, height = prs.slide_width, prs.slide_height
    for n, slide in enumerate(prs.slides, 1):
        where = f"slide {n}"
        title = slide.shapes.title
        if title is None or not title.text_frame.text.strip():
            out.append(issue("warn", where, "The slide has no title"))
        tables = 0
        for shape in slide.shapes:
            name = shape.name or shape.shape_type
            if getattr(shape, "has_table", False):
                tables += 1
            if shape.width is not None and shape.height is not None and (shape.width < 0 or shape.height < 0):
                out.append(issue("error", where, f"'{name}' has a negative size"))
                continue
            if None not in (shape.left, shape.top, shape.width, shape.height) and shape.shape_type != MSO_SHAPE_TYPE.LINE:
                if (shape.left < -EDGE_TOLERANCE_EMU or shape.top < -EDGE_TOLERANCE_EMU
                        or shape.left + shape.width > width + EDGE_TOLERANCE_EMU
                        or shape.top + shape.height > height + EDGE_TOLERANCE_EMU):
                    out.append(issue("error", where, f"'{name}' runs past the edge of the slide"))
            if shape.has_text_frame and shape.text_frame.text.strip():
                if _overflows(shape):
                    out.append(issue("warn", where, f"The text in '{name}' looks too long for its box"))
                if any(("".join(r.text for r in p.runs)).rstrip().endswith("…") for p in shape.text_frame.paragraphs):
                    out.append(issue("warn", where, f"Text in '{name}' was cut short with … to fit the slide"))
        out += _table_fallbacks(doc, n, where, tables)
    return out + _bullet_issues(doc) + _title_slide_issues(doc)


def _title_slide_issues(doc: ParsedDoc) -> list[dict]:
    """Slide 1 is drawn as a title and one subtitle line (plus a short byline); anything else on it is not drawn."""
    if not doc.sections:
        return []
    blocks = parse_blocks(split_notes(doc.sections[0].body)[0])
    hidden = [b for b in blocks if b.kind not in ("para", "list", "heading")]
    lines = sum(len(b.items) if b.kind == "list" else 1 for b in blocks if b.kind in ("para", "list", "heading"))
    if hidden or lines > 4:
        return [issue("warn", "slide 1", "The title slide only shows its title and a subtitle line; the rest of "
                      "slide 1 is not drawn. Put it on its own slide after a line with ---")]
    return []


def _bullet_issues(doc: ParsedDoc) -> list[dict]:
    """Slides whose source lists more than MAX_BULLETS items. Counted from the source: the writer drops what will not fit."""
    out = []
    for n, section in enumerate(doc.sections, 1):
        count = sum(len(b.items) for b in parse_blocks(section.body) if b.kind == "list")
        if count > MAX_BULLETS:
            out.append(issue("warn", f"slide {n}", f"The slide has {count} bullets; about {MAX_BULLETS} is the most that reads"))
    return out


def _table_fallbacks(doc: ParsedDoc, n: int, where: str, tables: int) -> list[dict]:
    """A diagram the source asked for that came out as a table, because the slide had no room to draw it."""
    if not tables or n > len(doc.sections):
        return []
    blocks = parse_blocks(doc.sections[n - 1].body)
    wanted = [b.rich for b in blocks if b.kind == "rich" and b.rich in _FIGURES]
    wants_table = any(b.kind == "table" or (b.kind == "rich" and b.rich == "csv") for b in blocks)
    if wanted and not wants_table:
        return [issue("warn", where, f"The {wanted[0]} did not fit and was drawn as a table; fewer items would draw")]
    return []


# ---- the whole code pass ----------------------------------------------------------------------

def check_source(source: str, title: Optional[str] = None) -> tuple[list[dict], dict[str, bytes], ParsedDoc]:
    """(issues, exported files, the parsed document). Slow (writes every format): run it in a thread."""
    issues, files, doc = export_all(source, title)
    name = (title or doc.title or "").strip()
    if doc.kind == "deck" and "pptx" in files:
        issues += check_pptx(files["pptx"], doc)
    if doc.kind == "sheet" and "xlsx" in files:
        issues += check_xlsx(files["xlsx"], plain_text_title(name), doc.outline)
    if doc.kind != "sheet":
        issues += check_rich(doc)
        issues += _outline_issues(doc)
    issues += check_placeholders(doc)
    try:
        from .docs_check_layout import check_layout      # imports this module, so it loads here

        issues += check_layout(doc, files.get("pptx"), issues)
    except Exception:                                    # noqa: BLE001 - the layout check is optional
        log.warning("layout check unavailable", exc_info=True)
    return issues, files, doc


# ---- placeholder text -------------------------------------------------------------------------

_PLACEHOLDERS = (
    re.compile(r"\blorem\s+ipsum\b|\bdolor\s+sit\s+amet\b", re.I),
    re.compile(r"\b(?:TBD|TBA)\b", re.I),
    re.compile(r"\bTODO\b(?!\s+(?:lists?|apps?|items?|trackers?|managers?)\b)"),     # 'TODO list' is a topic, 'TODO:' is a note
    # 'Insert logo here' as a short line of its own, or '(insert date here)'. A how-to sentence that happens to
    # contain 'insert ... here' ('Insert the SD card into the slot here') is prose, not a gap.
    re.compile(r"^[\s>*#\-+\d.)]*insert(?:\s+[\w'\u2019-]+){1,4}\s+here\b[\s.!:*_]*$", re.I | re.M),
    re.compile(r"\(\s*insert\b[^()\n]{0,40}?\bhere\s*\)", re.I),
    re.compile(r"\bclick to (?:add|edit)\b", re.I),
    re.compile(r"\bx{3,}\b", re.I),
    re.compile(r"\byour (?:text|name|company|title|logo|content|tagline) here\b", re.I),
    re.compile(r"\b(?:text|title|subtitle|content|description|heading|name) goes here\b", re.I),
)
_BRACKET = re.compile(r"\[([^\[\]\n]{2,60})\](?![(\[:])")                    # not a link, a reference or a checkbox
_PROMPT_WORDS = frozenset(
    "insert your name title date company here text description placeholder add enter topic image logo client "
    "project address number subtitle heading content point bullet details summary author email phone location "
    "product amount value xxx".split())
_FENCE = re.compile(r"^\s*(`{3,}|~{3,})\s*([A-Za-z0-9_-]*)")
_INLINE_CODE = re.compile(r"`[^`\n]*`")
_TABLE_TODO = re.compile(r"(?<=\|)[ \t]*TODO[ \t]*(?=\||$)")          # a table cell that says only TODO is a status
_MAX_PLACEHOLDER_ISSUES = 20


def _prose(text: str) -> str:
    """``text`` without code: a placeholder shown inside a code block or in backticks is an example, not a gap.
    The blocks the document format draws itself (chart, cards ...) stay, since their text is read by people."""
    out: list[str] = []
    fence: Optional[tuple[str, int, bool]] = None                 # (fence character, length, keep the lines)
    for line in text.split("\n"):
        m = _FENCE.match(line)
        if fence is None:
            if m:
                fence = (m.group(1)[0], len(m.group(1)), m.group(2).lower() in RICH_KINDS)
            else:
                out.append(line)
        elif m and m.group(1)[0] == fence[0] and len(m.group(1)) >= fence[1] and not m.group(2):
            fence = None
        elif fence[2]:
            out.append(line)
    out = [_TABLE_TODO.sub(" ", ln) if ln.lstrip().startswith("|") else ln for ln in out]
    return _INLINE_CODE.sub(" ", "\n".join(out))


def placeholder_matches(text: str, *, title: bool = False) -> list[str]:
    """The stand-in phrases left in ``text`` (at most three, in the order found)."""
    text = _prose(text)
    found: list[str] = []
    for pattern in _PLACEHOLDERS:
        m = pattern.search(text)
        if m and not (title and "list" in text.lower() and m.group(0) == "TODO"):
            found.append(m.group(0).strip(" \t>*#-+.:!_"))
    for m in _BRACKET.finditer(text):
        if _PROMPT_WORDS & set(re.findall(r"[a-z]+", m.group(1).lower())):
            found.append(m.group(0))
            break
    return found[:3]


def _cell_name(row: int, col: int) -> str:
    letters = ""
    col += 1
    while col:
        col, rem = divmod(col - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row + 1}"


def _text_parts(doc: ParsedDoc):
    """(where, what it is, text, is a title) for every piece of text a reader will see."""
    if doc.title not in (s.title for s in doc.sections):          # else it is reported where it is drawn
        yield "document", "document title", doc.title, True
    for i, entry in enumerate(doc.outline, 1):
        yield "outline", f"outline entry {i}", entry, False
    for n, s in enumerate([] if doc.kind == "sheet" else doc.sections, 1):    # a sheet's text is read cell by cell below
        where = f"slide {n}" if doc.kind == "deck" else f"section '{s.title}'" if s.title else "document"
        body, notes = split_notes(s.body) if doc.kind == "deck" else (s.body, "")
        yield where, "title" if doc.kind == "deck" else "heading", s.title, True
        yield where, "text", body, False
        yield where, "speaker notes", notes, False
    for sheet in doc.sheets:
        yield f"sheet {sheet.name}", "name", sheet.name, True
        for r, row in enumerate(sheet.rows):
            for c, cell in enumerate(row):
                if not cell.lstrip().startswith("="):
                    yield f"sheet {sheet.name}", f"cell {_cell_name(r, c)}", cell, False


def check_placeholders(doc: ParsedDoc) -> list[dict]:
    """Errors for stand-in text left in a document: lorem ipsum, TBD, 'click to add', [Company Name] and the like."""
    out: list[dict] = []
    for where, what, text, is_title in _text_parts(doc):
        if not text or not text.strip():
            continue
        if what.startswith("cell ") and text.strip() == "TODO":      # a status in a tracker, not a gap
            continue
        for hit in placeholder_matches(text, title=is_title):
            out.append(issue("error", where, f"Placeholder text is left in the {what}: {hit[:40]!r}. Write the real content"))
    return sort_issues(out)[:_MAX_PLACEHOLDER_ISSUES]


_TITLE_SLIDE_NAMES = {"intro", "introduction", "cover", "opening"}    # a first entry that names the title slide


def _outline_issues(doc: ParsedDoc) -> list[dict]:
    """Outline entries the body never wrote (a small model can list a slide and skip it)."""
    if doc.kind == "sheet" or not doc.outline:
        return []
    written = [w for w in (_norm(s.title) for s in doc.sections) if w]
    skip = {"title", "title slide", _norm(doc.title)}
    missing = [o for i, o in enumerate(doc.outline)
               if (n := _norm(o)) and n not in skip
               and not (doc.kind == "deck" and i == 0 and n in _TITLE_SLIDE_NAMES)
               and not any(n in w or w in n for w in written)]
    if not missing:
        return []
    unit = "slide" if doc.kind == "deck" else "section"
    return [issue("warn", "outline", f"The outline lists {unit}s that were never written: {', '.join(missing[:5])}")]


def plain_text_title(title: str) -> str:
    return re.sub(r"[*_`]", "", title)


# ---- rendered pages ---------------------------------------------------------------------------

def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[*_`#]", "", text)).strip().lower()


def _headings(doc: ParsedDoc) -> list[str]:
    found = [_norm(s.title) for s in doc.sections if s.title]
    for s in doc.sections:
        found += [_norm(plain_text(b.runs)) for b in parse_blocks(s.body) if b.kind == "heading"]
    return [h for h in dict.fromkeys(found) if len(h) >= 3]


def check_pages(doc: ParsedDoc, texts: list[str], fmt: str) -> list[dict]:
    """Problems visible in the text of the rendered pages. ``texts[i]`` is page ``i``."""
    label, out = page_label(fmt), []
    headings = _headings(doc) if doc.kind in ("doc", "pdf") else []
    for i, text in enumerate(texts):
        where = label(i)
        if doc.kind == "sheet":
            shown = [e for e in SHEET_ERRORS if e in text] + list(dict.fromkeys(_LO_ERROR.findall(text)))
            if shown:
                out.append(issue("error", where, f"The sheet shows {', '.join(shown)} in a cell"))
        if fmt == "pptx" and len(text.strip()) < 3:
            out.append(issue("warn", where, "No text was drawn on this slide"))
        if i == len(texts) - 1 or not headings:
            continue
        tail = _norm(" ".join([ln for ln in text.splitlines() if ln.strip()][-4:]))
        for h in headings:
            at = tail.rfind(h)
            if at >= 0 and len(tail[at + len(h):].split()) < 6:
                out.append(issue("warn", where, f"The heading '{h[:50]}' is the last thing on this page; "
                                                "its text starts on the next"))
                break
        else:
            if len(text.split()) < 10:
                out.append(issue("warn", where, "This page is almost empty"))
    return out


def sort_issues(issues: list[dict]) -> list[dict]:
    seen, out = set(), []
    for it in sorted(issues, key=lambda x: x["severity"] != "error"):
        key = (it["severity"], it["where"], it["message"])
        if key not in seen:
            seen.add(key)
            out.append(it)
    return out
