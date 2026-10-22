"""Shape and text helpers for the pptx writers: native shapes that fit the box they are given."""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from .docs_blocks import Run
from .docs_metrics import fit_size
from .docs_chips import INK, MUTED, ACCENT

MONO = "Consolas"
LEFT, CENTER, RIGHT = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP, MIDDLE, BOTTOM = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM
_MARGIN = 0.05                                   # inches inside every text box


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def box(slide, l: float, t: float, w: float, h: float, *, fill: str | None = None, line: str | None = None,
        shape=MSO_SHAPE.RECTANGLE, radius: float | None = None, line_w: float = 0.75):
    """A filled shape; sizes in inches. ``radius`` (0..0.5 of the short side) rounds a rounded rectangle."""
    s = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(max(w, 0.01)), Inches(max(h, 0.01)))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return flat(s)


def flat(shape):
    """Switch off the theme's shadow, which LibreOffice draws even when the shape says ``effectLst`` is empty."""
    ref = shape._element.find(".//" + qn("a:effectRef"))
    if ref is not None:
        ref.set("idx", "0")
    return shape


def write_runs(paragraph, runs: list[Run], size: float, color: RGBColor | str = INK, *, bold: bool = False,
               italic: bool = False) -> None:
    color = rgb(color) if isinstance(color, str) else color
    for r in runs or [Run("")]:
        run = paragraph.add_run()
        run.text = r.text
        font = run.font
        font.size = Pt(size - 2 if r.code else size)
        font.bold = bold or r.bold
        font.italic = italic or r.italic
        font.color.rgb = rgb(ACCENT) if r.href else color
        if r.code:
            font.name = MONO
        if r.href:
            run.hyperlink.address = r.href
            font.underline = True


def text(slide, l: float, t: float, w: float, h: float, content: str, size: float, *, color: str = INK,
         bold: bool = False, italic: bool = False, align=LEFT, anchor=TOP, mono: bool = False,
         min_size: float | None = None, shape=None, wrap: bool = True) -> float:
    """Text that fits its box: shrunk down to ``min_size``, then cut with an ellipsis. Returns the size used.

    With ``shape`` the text goes into that shape instead of a new text box."""
    w, h = max(w, 0.02), max(h, 0.02)                    # a size below zero makes the file need repair
    inner_w, inner_h = max(w - 2 * _MARGIN, 0.01) * 72, max(h - 2 * _MARGIN, 0.01) * 72
    used, lines = fit_size(content, inner_w, inner_h, size, min_size if min_size is not None else max(size * 0.6, 8),
                           bold=bold, mono=mono)
    tb = shape or slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    frame = tb.text_frame
    frame.word_wrap = wrap
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(_MARGIN)
    flat = lambda v: "".join(v).replace(" ", "").replace("\n", "")      # noqa: E731
    paras = content.split("\n") if flat(lines) == flat([content]) else lines     # cut text keeps its cut
    for i, line in enumerate(paras or [""]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(used)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
        if mono:
            run.font.name = MONO
    return used


def muted(slide, l, t, w, h, content: str, size: float = 14, **kw) -> float:
    return text(slide, l, t, w, h, content, size, color=MUTED, **kw)


def clamp_to_slide(slide, width: float, height: float) -> None:
    """Last line of defence: no shape has a negative size or reaches past the slide edge, whatever the content."""
    emu_w, emu_h = int(Inches(width)), int(Inches(height))
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.LINE or not hasattr(shape, "width"):
            continue                                     # connectors keep their ends; their geometry is already inside
        left, top = min(max(shape.left, 0), emu_w), min(max(shape.top, 0), emu_h)
        shape.left, shape.top = left, top
        shape.width, shape.height = min(max(shape.width, 0), emu_w - left), min(max(shape.height, 0), emu_h - top)
