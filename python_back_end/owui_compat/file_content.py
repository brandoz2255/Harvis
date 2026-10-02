"""Turn an uploaded attachment into text a model can be trusted with.

Three failure modes this replaces, all seen on the k8s server with gemma4:e2b:

* a 50,000-row CSV was head-cut to 24,000 characters (759 rows) with a bare
  ``…[truncated]`` the user never saw, and the model invented totals;
* PDFs went in as ``%PDF-1.4`` bytes decoded as UTF-8;
* the only disclosure of any cut lived inside the prompt, never in words the
  model would repeat to the user.

Every renderer here returns plain text. Tabular files get an overview computed
in Python over the WHOLE file (exact row count, per-column types and stats) and
then as many verbatim rows as fit, with a note that says how many of how many are
shown. Documents are text-extracted with whatever library the image has (pypdf,
python-docx, openpyxl, python-pptx) and say so plainly when one is missing.
Nothing here raises: a renderer that cannot read a file returns a visible note.
"""

from __future__ import annotations

import csv
import io
import logging
import math
import mimetypes
import os
import re
from collections import Counter, deque
from typing import Iterable, Iterator, Optional

logger = logging.getLogger(__name__)

# Cap injected text so a giant attachment can't blow the context budget.
MAX_TEXT_FILE_CHARS = int(os.getenv("HARVIS_ATTACHMENT_TEXT_CHARS", "24000"))

_TABULAR_EXTS = frozenset({".csv", ".tsv", ".tab"})
_XLSX_EXTS = frozenset({".xlsx", ".xlsm"})
_TABULAR_MIMES = frozenset({"text/csv", "text/tab-separated-values", "application/csv"})
_XLSX_MIMES = frozenset({
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.macroenabled.12",
})
_PDF_MIMES = frozenset({"application/pdf", "application/x-pdf"})
_DOCX_MIMES = frozenset({"application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
_PPTX_MIMES = frozenset({"application/vnd.openxmlformats-officedocument.presentationml.presentation"})

_NUM_RE = re.compile(r"^[+-]?\$?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?([eE][+-]?\d+)?$|^[+-]?\$?\.\d+$")
_HEAD_ROWS_MIN = 5
_TAIL_ROWS = 5
_TOP_VALUES = 3
_DISTINCT_CAP = 2000  # stop counting distinct text values past this; say "2000+"


def classify(filename: str, content_type: str) -> str:
    """One of: image, media, table, xlsx, pdf, docx, pptx, text."""
    ctype = (content_type or "").split(";")[0].strip().lower()
    ext = os.path.splitext(filename or "")[1].lower()
    if ctype.startswith("image/"):
        return "image"
    if ctype.startswith(("audio/", "video/")):
        return "media"
    if ext in _XLSX_EXTS or ctype in _XLSX_MIMES:
        return "xlsx"
    if ext in _TABULAR_EXTS or ctype in _TABULAR_MIMES:
        return "table"
    if ext == ".pdf" or ctype in _PDF_MIMES:
        return "pdf"
    if ext == ".docx" or ctype in _DOCX_MIMES:
        return "docx"
    if ext == ".pptx" or ctype in _PPTX_MIMES:
        return "pptx"
    if not ctype or ctype == "application/octet-stream":
        guessed = mimetypes.guess_type(filename or "")[0] or ""
        if guessed in _PDF_MIMES:
            return "pdf"
    return "text"


def missing_file_note(filename: str) -> str:
    return (
        f'[Harvis: the attached file "{filename}" is no longer available on the server '
        "(its stored copy is missing, usually after a restart on a host without persistent "
        "upload storage). Tell the user to re-upload it; do not guess at its contents.]"
    )


def _fmt_int(n: int) -> str:
    return f"{n:,}"


def _fmt_num(x: float) -> str:
    if math.isfinite(x) and float(x).is_integer() and abs(x) < 1e15:
        return f"{int(x):,}"
    return f"{x:,.4g}" if abs(x) < 1e-3 or abs(x) >= 1e15 else f"{x:,.2f}"


def _to_number(cell: str) -> Optional[float]:
    s = cell.strip()
    if not s or not _NUM_RE.match(s):
        return None
    try:
        return float(s.replace(",", "").replace("$", ""))
    except ValueError:
        return None


# ── Tabular profile ──────────────────────────────────────────────────────────
class _Column:
    __slots__ = ("name", "filled", "numeric", "ints", "lo", "hi", "total", "values", "capped")

    def __init__(self, name: str):
        self.name = name
        self.filled = 0
        self.numeric = 0
        self.ints = 0
        self.lo = math.inf
        self.hi = -math.inf
        self.total = 0.0
        self.values: Counter = Counter()
        self.capped = False

    def add(self, cell: str) -> None:
        if cell is None or not str(cell).strip():
            return
        cell = str(cell)
        self.filled += 1
        num = _to_number(cell)
        if num is not None:
            self.numeric += 1
            if num.is_integer():
                self.ints += 1
            self.lo = min(self.lo, num)
            self.hi = max(self.hi, num)
            self.total += num
        if not self.capped:
            self.values[cell.strip()] += 1
            if len(self.values) > _DISTINCT_CAP:
                self.capped = True
                self.values = Counter()

    def describe(self, rows: int) -> str:
        if self.filled == 0:
            return f"{self.name} — empty"
        empty = rows - self.filled
        tail = f"; {_fmt_int(empty)} empty" if empty else ""
        if self.numeric and self.numeric >= 0.9 * self.filled:
            kind = "integer" if self.ints == self.numeric else "number"
            mean = self.total / self.numeric
            stats = (f"min {_fmt_num(self.lo)}, max {_fmt_num(self.hi)}, "
                     f"mean {_fmt_num(mean)}, sum {_fmt_num(self.total)}")
            if self.numeric != self.filled:
                kind += f" ({_fmt_int(self.filled - self.numeric)} non-numeric)"
            return f"{self.name} — {kind}; {stats}{tail}"
        if self.capped:
            return f"{self.name} — text; {_fmt_int(_DISTINCT_CAP)}+ distinct values{tail}"
        distinct = len(self.values)
        top = ", ".join(f"{v!r} ({_fmt_int(c)})" for v, c in self.values.most_common(_TOP_VALUES))
        return f"{self.name} — text; {_fmt_int(distinct)} distinct; top: {top}{tail}"


class TableProfile:
    def __init__(self, header: list[str]):
        self.header = header
        self.columns = [_Column(h or f"column {i + 1}") for i, h in enumerate(header)]
        self.rows = 0
        self.ragged = 0
        self.head: list[list[str]] = []
        self.tail: deque = deque(maxlen=_TAIL_ROWS)

    def add(self, row: list[str], keep_head_up_to: int) -> None:
        self.rows += 1
        if len(row) != len(self.columns):
            self.ragged += 1
            while len(row) > len(self.columns):
                self.columns.append(_Column(f"column {len(self.columns) + 1}"))
        for col, cell in zip(self.columns, row):
            col.add(cell)
        if len(self.head) < keep_head_up_to:
            self.head.append(row)
        self.tail.append(row)


def profile_rows(rows: Iterable[list], max_head_rows: int) -> Optional[TableProfile]:
    it: Iterator = iter(rows)
    header = None
    for first in it:
        header = ["" if c is None else str(c) for c in first]
        break
    if header is None:
        return None
    prof = TableProfile(header)
    for row in it:
        prof.add(["" if c is None else str(c) for c in row], max_head_rows)
    return prof


def _serialise(rows: Iterable[list], delimiter: str) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator="\n")
    for r in rows:
        w.writerow(r)
    return buf.getvalue()


def render_profile(prof: TableProfile, *, title: str, delimiter: str, max_chars: int,
                   whole_text: Optional[str] = None, extra: str = "") -> str:
    """Overview + as many rows as fit + a visible partial-view note."""
    n = prof.rows
    width = len(prof.columns)
    lines = [title, f"Overview computed by Harvis in Python over the WHOLE file (exact figures):",
             f"- Rows: {_fmt_int(n)} data rows (+1 header row); Columns: {width}"]
    if prof.ragged:
        lines.append(f"- {_fmt_int(prof.ragged)} rows had a different number of fields than the header")
    if extra:
        lines.append(extra)
    lines.append("- Columns:")
    for i, col in enumerate(prof.columns, 1):
        lines.append(f"  {i}. {col.describe(n)}")
    overview = "\n".join(lines)

    if whole_text is not None and len(overview) + len(whole_text) + 64 <= max_chars:
        return f"{overview}\n\nFull contents ({_fmt_int(n)} rows, nothing omitted):\n{whole_text.rstrip()}"

    header_line = _serialise([prof.header], delimiter)
    tail_rows = list(prof.tail)
    tail_text = _serialise(tail_rows, delimiter) if n > len(prof.head) else ""
    budget = max_chars - len(overview) - len(header_line) - len(tail_text) - 700
    shown: list[list[str]] = []
    used = 0
    for row in prof.head:
        line = _serialise([row], delimiter)
        if used + len(line) > budget and len(shown) >= _HEAD_ROWS_MIN:
            break
        shown.append(row)
        used += len(line)
    if n <= len(shown):
        body = f"Rows 1–{_fmt_int(n)} of {_fmt_int(n)}:\n{_serialise(shown, delimiter)}"
        return f"{overview}\n\n{header_line}{body}".rstrip()

    tail_shown = [r for r in tail_rows][max(0, len(tail_rows) - (n - len(shown))):]
    shown_total = len(shown) + len(tail_shown)
    omitted = n - shown_total
    note = (
        f"**IMPORTANT — PARTIAL VIEW: only {_fmt_int(shown_total)} of {_fmt_int(n)} rows are "
        f"reproduced below (the first {_fmt_int(len(shown))} and the last {_fmt_int(len(tail_shown))}; "
        f"{_fmt_int(omitted)} rows were not sent to the model). Do NOT estimate totals, counts, "
        "averages or 'last rows' from the excerpt: use the exact figures in the overview above, "
        f"and tell the user plainly that only {_fmt_int(shown_total)} of {_fmt_int(n)} rows were shown. "
        "For anything the overview does not answer, say the full file was not read and offer a "
        "workspace analysis over the whole file.**"
    )
    parts = [overview, note, header_line + f"Rows 1–{_fmt_int(len(shown))}:\n" + _serialise(shown, delimiter)]
    if tail_shown:
        parts.append(f"…[{_fmt_int(omitted)} rows omitted]…\n"
                     f"Rows {_fmt_int(n - len(tail_shown) + 1)}–{_fmt_int(n)} (the last rows):\n"
                     + _serialise(tail_shown, delimiter))
    return "\n\n".join(p.rstrip() for p in parts)


def _decode(raw: bytes) -> str:
    return raw.decode("utf-8-sig", errors="replace")


def _sniff_delimiter(sample: str, filename: str) -> str:
    if os.path.splitext(filename or "")[1].lower() in (".tsv", ".tab"):
        return "\t"
    try:
        return csv.Sniffer().sniff(sample[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def render_table(filename: str, raw: bytes, max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    text = _decode(raw)
    delimiter = _sniff_delimiter(text, filename)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    # Every row may need to be shown when the file is small, so keep enough head
    # rows to fill the budget; a row is rarely under 2 characters.
    prof = profile_rows(reader, max_head_rows=max(_HEAD_ROWS_MIN, max_chars // 2))
    if prof is None:
        return f"### Attached file: {filename}\n[Harvis: the file is empty]"
    kind = "TSV" if delimiter == "\t" else "CSV"
    title = f"### Attached file: {filename} ({kind}, {_fmt_int(len(raw))} bytes)"
    extra = "" if delimiter in (",", "\t") else f"- Field delimiter: {delimiter!r}"
    return render_profile(prof, title=title, delimiter=delimiter, max_chars=max_chars,
                          whole_text=text, extra=extra)


def render_xlsx(filename: str, raw: bytes, max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    try:
        import openpyxl
    except ImportError:
        return (f"### Attached file: {filename}\n[Harvis: cannot read this spreadsheet — "
                "openpyxl is not installed on this server. Ask the user for a CSV export.]")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheets = wb.sheetnames
        prof, used = None, None
        for name in sheets:
            prof = profile_rows(wb[name].iter_rows(values_only=True),
                                max_head_rows=max(_HEAD_ROWS_MIN, max_chars // 2))
            if prof is not None and (prof.rows or any(prof.header)):
                used = name
                break
    finally:
        wb.close()
    if prof is None or used is None:
        return f"### Attached file: {filename}\n[Harvis: the workbook has no data]"
    others = [s for s in sheets if s != used]
    extra = f"- Sheet shown: {used!r}"
    if others:
        extra += (f"; other sheets NOT shown: {', '.join(repr(s) for s in others)} "
                  "(ask the user which one they mean if it matters)")
    title = f"### Attached file: {filename} (Excel workbook, {_fmt_int(len(raw))} bytes)"
    whole = _serialise([prof.header] + prof.head, ",") if prof.rows <= len(prof.head) else None
    return render_profile(prof, title=title, delimiter=",", max_chars=max_chars,
                          whole_text=whole, extra=extra)


# ── Documents ────────────────────────────────────────────────────────────────
def _truncate_note(text: str, filename: str, max_chars: int, unit: str = "characters") -> str:
    if len(text) <= max_chars:
        return text
    total = len(text)
    kept = text[:max_chars]
    return (
        f"{kept}\n\n[Harvis: PARTIAL VIEW — \"{filename}\" has {_fmt_int(total)} {unit} of text; only "
        f"the first {_fmt_int(max_chars)} are shown above and the rest was NOT sent to the model. "
        "Tell the user the file was only partially read; do not summarise or quote anything beyond "
        "what is shown.]"
    )


def render_pdf(filename: str, raw: bytes, max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    head = f"### Attached file: {filename} (PDF, {_fmt_int(len(raw))} bytes)"
    try:
        import pypdf
    except ImportError:
        return f"{head}\n[Harvis: cannot read this PDF — pypdf is not installed on this server.]"
    try:
        reader = pypdf.PdfReader(io.BytesIO(raw))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                pages.append(f"--- page {i} ---\n{t}")
        n_pages = len(reader.pages)
    except Exception as exc:
        logger.warning("file_content: pdf extraction failed for %s: %s", filename, exc)
        return f"{head}\n[Harvis: could not extract text from this PDF ({type(exc).__name__}).]"
    if not pages:
        return (f"{head}\n[Harvis: this PDF ({n_pages} page(s)) contains no extractable text — "
                "it is probably scanned or image-only. Tell the user; do not guess its contents.]")
    body = "\n".join(pages)
    return f"{head}, {n_pages} page(s)\n" + _truncate_note(body, filename, max_chars)


def render_docx(filename: str, raw: bytes, max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    head = f"### Attached file: {filename} (Word document, {_fmt_int(len(raw))} bytes)"
    try:
        import docx
    except ImportError:
        return f"{head}\n[Harvis: cannot read this document — python-docx is not installed on this server.]"
    try:
        doc = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(c.text.strip() for c in row.cells))
    except Exception as exc:
        logger.warning("file_content: docx extraction failed for %s: %s", filename, exc)
        return f"{head}\n[Harvis: could not extract text from this document ({type(exc).__name__}).]"
    body = "\n".join(parts).strip()
    if not body:
        return f"{head}\n[Harvis: this document contains no extractable text.]"
    return f"{head}\n" + _truncate_note(body, filename, max_chars)


def render_pptx(filename: str, raw: bytes, max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    head = f"### Attached file: {filename} (PowerPoint, {_fmt_int(len(raw))} bytes)"
    try:
        from pptx import Presentation
    except ImportError:
        return f"{head}\n[Harvis: cannot read this presentation — python-pptx is not installed on this server.]"
    try:
        prs = Presentation(io.BytesIO(raw))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [sh.text_frame.text.strip() for sh in slide.shapes
                     if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
            if texts:
                slides.append(f"--- slide {i} ---\n" + "\n".join(texts))
        n = len(prs.slides)
    except Exception as exc:
        logger.warning("file_content: pptx extraction failed for %s: %s", filename, exc)
        return f"{head}\n[Harvis: could not extract text from this presentation ({type(exc).__name__}).]"
    if not slides:
        return f"{head}\n[Harvis: this presentation ({n} slide(s)) contains no extractable text.]"
    return f"{head}, {n} slide(s)\n" + _truncate_note("\n".join(slides), filename, max_chars)


def render_text(filename: str, raw: bytes, max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    text = _decode(raw)
    if not text.strip():
        return ""
    if len(text) <= max_chars:
        return f"### Attached file: {filename}\n{text}"
    return f"### Attached file: {filename} ({_fmt_int(len(raw))} bytes)\n" + _truncate_note(text, filename, max_chars)


def render_inline_text(filename: str, text: str, max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    return f"### Attached file: {filename}\n" + _truncate_note(text, filename, max_chars)


_RENDERERS = {
    "table": render_table,
    "xlsx": render_xlsx,
    "pdf": render_pdf,
    "docx": render_docx,
    "pptx": render_pptx,
    "text": render_text,
}


def render_attachment(filename: str, content_type: str, raw: bytes,
                      max_chars: int = MAX_TEXT_FILE_CHARS) -> str:
    """Text block for one non-image, non-media upload. Never raises."""
    kind = classify(filename, content_type)
    fn = _RENDERERS.get(kind, render_text)
    try:
        return fn(filename, raw, max_chars)
    except Exception as exc:
        logger.warning("file_content: %s renderer failed for %s: %s", kind, filename, exc, exc_info=True)
        return (f"### Attached file: {filename}\n[Harvis: could not read this file "
                f"({content_type or 'unknown type'}, {type(exc).__name__}). Tell the user.]")
